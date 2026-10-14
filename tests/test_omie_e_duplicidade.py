from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.dominio import MotivoDeNaoIdentificacao
from conciliacao.encargo import ParametrosDeEncargo
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import (
    LinhaItau,
    TituloOmie,
    TituloProtheus,
    configuracao_itau,
    escrever_extrato_itau,
    escrever_omie,
    escrever_protheus,
)
from tests.parecerista import PareceristaFalso

CNPJ_PROTHEUS = "10.200.300/0001-40"
CNPJ_OMIE_FILIAL = "10200300000299"


def test_exportacao_omie_e_lida_como_vem_e_casa_por_cnpj_como_o_protheus(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_omie(
        pasta,
        [
            TituloOmie(documento="FAT-2025-0001", valor="31250.00", vencimento="2025-03-05"),
            TituloOmie(documento="FAT-2025-0002", valor="980.50", vencimento="2025-03-12"),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("12/03/2025", "PIX RECEBIDO 10200300000140", "980,50")])

    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())

    assert resultado.arquivos_recusados == ()
    [item] = resultado.lista_de_baixa
    [titulo] = item.titulos
    assert (titulo.documento, titulo.sistema_de_origem, titulo.valor_centavos) == ("FAT-2025-0002", "Omie", 98050)
    assert titulo.cnpj == "10200300000140"
    assert titulo.nosso_numero is None
    assert "Valor único" in item.explicacao
    assert "Omie" in item.explicacao


def test_atraso_em_titulo_do_omie_usa_o_vencimento_iso(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_omie(pasta, [TituloOmie(documento="FAT-2025-0003", valor="1000.00", vencimento="2025-03-05")])
    escrever_extrato_itau(pasta, [LinhaItau("14/03/2025", "PIX RECEBIDO 10200300000140", "1.023,00")])
    juros_de_1_por_cento = ParametrosDeEncargo(juros_ao_mes_percentual=Decimal("1"))

    resultado = executar_conciliacao(
        pasta, armazenamento, configuracao_itau(encargo=juros_de_1_por_cento), PareceristaFalso()
    )

    [item] = resultado.lista_de_baixa
    assert [t.documento for t in item.titulos] == ["FAT-2025-0003"]
    assert "Encargo" in item.explicacao


@pytest.mark.parametrize(
    ("titulo", "trecho"),
    [
        (TituloOmie(documento="FAT-1", valor="1.250"), "valor '1.250'"),
        (TituloOmie(documento="FAT-1", valor="31250"), "valor '31250'"),
        (TituloOmie(documento="FAT-1", valor="100.00", vencimento="05/03/2025"), "data '05/03/2025'"),
        (TituloOmie(documento="FAT-1", valor="100.00", vencimento="2025-02-30"), "data '2025-02-30'"),
        (TituloOmie(documento="FAT-1", valor="100.00", cnpj="10.200.300/0001-40"), "CNPJ '10.200.300/0001-40'"),
        (TituloOmie(documento="FAT-1", valor="100.00", pago="S"), "'S'"),
    ],
    ids=["valor com tres casas", "valor sem centavos", "data ptbr", "data inexistente", "cnpj pontuado", "pago"],
)
def test_exportacao_omie_fora_do_formato_e_recusada_com_motivo(
    pasta: Path, armazenamento: Armazenamento, titulo: TituloOmie, trecho: str
) -> None:
    escrever_omie(pasta, [titulo], nome="omie.csv")

    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())

    [recusado] = resultado.arquivos_recusados
    assert recusado.nome == "omie.csv"
    assert "linha 2" in recusado.motivo
    assert trecho in recusado.motivo


def test_valor_unico_considera_os_titulos_do_franqueado_em_todos_os_sistemas(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 920001", valor="500,00", vencimento="05/03/2025")])
    escrever_omie(
        pasta,
        [
            TituloOmie(documento="FAT-0500", valor="500.00", vencimento="2025-03-20", cnpj=CNPJ_OMIE_FILIAL),
            TituloOmie(documento="FAT-0700", valor="700.00", vencimento="2025-03-20", cnpj=CNPJ_OMIE_FILIAL),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "PIX RECEBIDO 10200300000140", "500,00"),
            LinhaItau("05/03/2025", "PIX RECEBIDO 10200300000140", "700,00"),
        ],
    )

    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())

    [ambiguo] = resultado.fila_de_excecoes
    assert ambiguo.motivo is MotivoDeNaoIdentificacao.MAIS_DE_UMA_COMBINACAO
    assert "NF 920001" in ambiguo.explicacao
    assert "FAT-0500" in ambiguo.explicacao
    [unico] = resultado.lista_de_baixa
    assert [(t.documento, t.sistema_de_origem) for t in unico.titulos] == [("FAT-0700", "Omie")]


def test_mesmo_franqueado_valor_e_vencimento_em_protheus_e_omie_vira_possivel_duplicidade(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 930001", valor="2.400,00", vencimento="10/03/2025", cnpj=CNPJ_PROTHEUS),
            TituloProtheus(documento="NF 930002", valor="800,00", vencimento="10/03/2025", cnpj=CNPJ_PROTHEUS),
        ],
    )
    escrever_omie(
        pasta,
        [
            TituloOmie(documento="FAT-9301", valor="2400.00", vencimento="2025-03-10", cnpj=CNPJ_OMIE_FILIAL),
            TituloOmie(documento="FAT-9302", valor="800.00", vencimento="2025-03-11", cnpj=CNPJ_OMIE_FILIAL),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("10/03/2025", "TARIFA", "5,00", tipo="D")])

    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())

    [duplicidade] = resultado.itens_de_titulo
    assert duplicidade.motivo is MotivoDeNaoIdentificacao.POSSIVEL_DUPLICIDADE
    assert sorted((t.sistema_de_origem, t.documento) for t in duplicidade.titulos) == [
        ("Omie", "FAT-9301"),
        ("Protheus", "NF 930001"),
    ]
    assert "R$ 2.400,00" in duplicidade.explicacao
    assert "10/03/2025" in duplicidade.explicacao
    assert resultado.fila_de_excecoes == []


def test_mesmo_valor_e_vencimento_no_mesmo_sistema_ou_em_franqueados_diferentes_nao_e_duplicidade(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 940001", valor="600,00", vencimento="10/03/2025"),
            TituloProtheus(documento="NF 940002", valor="600,00", vencimento="10/03/2025"),
        ],
    )
    escrever_omie(
        pasta, [TituloOmie(documento="FAT-9401", valor="600.00", vencimento="2025-03-10", cnpj="20400600000180")]
    )
    escrever_extrato_itau(pasta, [])

    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())

    assert resultado.itens_de_titulo == ()


def test_item_de_duplicidade_nao_entra_na_conferencia_de_soma_dos_creditos(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 950001", valor="1.500,00", vencimento="10/03/2025")])
    escrever_omie(pasta, [TituloOmie(documento="FAT-9501", valor="1500.00", vencimento="2025-03-10")])
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("10/03/2025", "PIX RECEBIDO FULANO", "250,00"),
            LinhaItau("10/03/2025", "TARIFA", "5,00", tipo="D"),
        ],
    )

    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    caminho = tmp_path / "saida.xlsx"
    gerar_workbook(resultado, caminho)

    assert len(resultado.itens_de_titulo) == 1
    assert [item.lancamento.historico for item in resultado.fila_de_excecoes] == ["PIX RECEBIDO FULANO"]
    workbook = load_workbook(caminho)
    resumo = {linha[0]: linha[1] for linha in workbook["Resumo"].iter_rows(values_only=True)}
    assert resumo["Fila de exceções"] == "1 lançamento(s), R$ 250,00"
    assert resumo["Itens de Título na Fila"] == 1
    linhas = list(workbook["Fila de exceções"].iter_rows(values_only=True))
    cabecalho = linhas[0]
    fila = [dict(zip(cabecalho, linha)) for linha in linhas[1:]]
    assert [linha["Motivo"] for linha in fila] == [
        MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR.value,
        MotivoDeNaoIdentificacao.POSSIVEL_DUPLICIDADE.value,
    ]
    duplicidade = fila[1]
    assert duplicidade["Valor"] is None
    assert duplicidade["Item"] == resultado.itens_de_titulo[0].identificador
    assert "NF 950001 (Protheus)" in str(duplicidade["Títulos sugeridos"])
    assert "FAT-9501 (Omie)" in str(duplicidade["Títulos sugeridos"])
    assert sum(linha["Valor"] for linha in fila if isinstance(linha["Valor"], (int, float))) == 250


def test_item_de_duplicidade_mantem_o_identificador_entre_execucoes(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 960001", valor="900,00", vencimento="10/03/2025")])
    escrever_omie(pasta, [TituloOmie(documento="FAT-9601", valor="900.00", vencimento="2025-03-10")])
    escrever_extrato_itau(pasta, [])

    primeira = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    segunda = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())

    assert [i.identificador for i in segunda.itens_de_titulo] == [i.identificador for i in primeira.itens_de_titulo]
    assert not primeira.itens_de_titulo[0].identificador.startswith("L")


def test_franqueado_guarda_os_nomes_de_cada_sistema_de_origem_entre_execucoes(
    tmp_path: Path, armazenamento: Armazenamento
) -> None:
    julho = tmp_path / "julho"
    julho.mkdir()
    escrever_protheus(julho, [TituloProtheus(documento="NF 1", valor="10,00", cliente="Padaria Teste Ltda")])
    escrever_omie(julho, [TituloOmie(documento="FAT-1", valor="20.00", cliente="PADARIA TESTE", cnpj=CNPJ_OMIE_FILIAL)])
    escrever_extrato_itau(julho, [])
    executar_conciliacao(julho, armazenamento, configuracao_itau(), PareceristaFalso())
    agosto = tmp_path / "agosto"
    agosto.mkdir()
    escrever_protheus(agosto, [TituloProtheus(documento="NF 2", valor="10,00", cliente="Padaria Teste Filial")])
    escrever_extrato_itau(agosto, [])
    executar_conciliacao(agosto, armazenamento, configuracao_itau(), PareceristaFalso())

    franqueado = armazenamento.franqueados()["10200300"]

    assert franqueado.nomes == {
        "Protheus": frozenset({"Padaria Teste Ltda", "Padaria Teste Filial"}),
        "Omie": frozenset({"PADARIA TESTE"}),
    }
    assert franqueado.todos_os_nomes == frozenset({"Padaria Teste Ltda", "Padaria Teste Filial", "PADARIA TESTE"})
