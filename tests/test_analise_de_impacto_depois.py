from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.configuracao import (
    Configuracao,
    ConfiguracaoDaAnaliseDeImpacto,
    ConfiguracaoInvalida,
    carregar_configuracao,
)
from conciliacao.decisoes import registrar_decisao_da_tela
from conciliacao.dominio import (
    Destino,
    Item,
    LinhaDeDecisao,
    MotivoDeNaoIdentificacao,
    Natureza,
    Resultado,
    TipoDeEvidencia,
    TotalDoDestino,
)
from conciliacao.execucao import executar_conciliacao
from conciliacao.gerador import Caso, ConfiguracaoDoGerador, LancamentoEsperado, gerar_dados_sinteticos
from conciliacao.impacto import (
    AnaliseDeImpacto,
    ChegadaDosCreditos,
    FonteDaVerdade,
    PorQueFicaSolto,
    analisar_impacto,
    gerar_workbook_do_impacto,
    main,
)
from conciliacao.saidas import ABA_DA_FILA_DE_EXCECOES
from tests.arquivos import LinhaItau, TituloProtheus, configuracao_itau, escrever_extrato_itau, escrever_protheus
from tests.parecerista import PareceristaFalso

CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.400.600/0001-80"
CNPJ_C = "30.600.900/0001-10"
CNPJ_D = "40.800.100/0001-20"
VENCIMENTO = "10/03/2025"
Motivo = MotivoDeNaoIdentificacao
Solto = PorQueFicaSolto


def _cenario(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus("NF A1", "1.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A, vencimento=VENCIMENTO),
            TituloProtheus("NF A2", "2.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A, vencimento=VENCIMENTO),
            TituloProtheus(
                "NF B1",
                "500,00",
                nosso_numero="000077700001",
                cliente="Brisa Serena Armazem Ltda",
                cnpj=CNPJ_B,
                vencimento=VENCIMENTO,
            ),
            TituloProtheus("NF C1", "800,00", cliente="Recanto Feliz Ltda", cnpj=CNPJ_C, vencimento=VENCIMENTO),
            TituloProtheus(
                "NF D1",
                "900,00",
                nosso_numero="000077700009",
                cliente="Doce Lar Mercearia Ltda",
                cnpj=CNPJ_D,
                vencimento=VENCIMENTO,
            ),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "TED RECEBIDA CNPJ 10200300000140", "3.000,00"),
            LinhaItau("05/03/2025", "PIX RECEBIDO BRISA SERENA ARM", "400,00"),
            LinhaItau("06/03/2025", "PIX RECEBIDO RECANTO FELIZ", "300,00"),
            LinhaItau("06/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "900,00", documento="000077700009"),
            LinhaItau("07/03/2025", "DEPOSITO EM DINHEIRO", "333,00"),
            LinhaItau("07/03/2025", "TARIFA BANCARIA", "2,50", tipo="D"),
        ],
    )


def _configuracao(*adesoes: int, semente: int = 1) -> Configuracao:
    return replace(
        configuracao_itau(),
        analise_de_impacto=ConfiguracaoDaAnaliseDeImpacto(
            semente=semente, faixa_de_adesao=tuple(Decimal(adesao) for adesao in adesoes)
        ),
    )


def _creditos(resultado: Resultado) -> list[Item]:
    return [*resultado.lista_de_baixa, *resultado.fila_de_excecoes]


def _dinheiro(resultado: Resultado) -> int:
    return sum(item.lancamento.valor_centavos for item in _creditos(resultado))


def _identificado(item: Item) -> bool:
    return bool({TipoDeEvidencia.NOSSO_NUMERO, TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA} & set(item.evidencia))


def _do_valor(resultado: Resultado, valor_centavos: int) -> Item:
    [item] = [item for item in _creditos(resultado) if item.lancamento.valor_centavos == valor_centavos]
    return item


def _documentos(item: Item) -> list[str]:
    return [titulo.documento for titulo in item.titulos]


def test_adesao_zero_da_depois_igual_ao_antes(pasta: Path) -> None:
    _cenario(pasta)

    analise = analisar_impacto(pasta, _configuracao(0), PareceristaFalso())

    def destinos(resultado: Resultado) -> list[tuple[object, ...]]:
        return sorted(
            (i.lancamento.data, i.lancamento.historico, i.lancamento.valor_centavos, i.destino, i.motivo, _documentos(i))
            for i in resultado.itens
        )

    [cenario] = analise.cenarios
    assert destinos(cenario.depois) == destinos(analise.antes)
    assert cenario.depois.resumo.totais == analise.antes.resumo.totais
    assert ChegadaDosCreditos.do_resultado(cenario.depois) == ChegadaDosCreditos.do_resultado(analise.antes)


def test_toda_a_adesao_identifica_o_credito_com_titulo_conhecido_e_valor_permitido(pasta: Path) -> None:
    _cenario(pasta)

    analise = analisar_impacto(pasta, _configuracao(100), PareceristaFalso())

    [cenario] = analise.cenarios
    depois = cenario.depois
    pela_soma = [i for i in depois.lista_de_baixa if i.lancamento.data.day == 5]
    assert sorted((i.lancamento.valor_centavos, tuple(_documentos(i))) for i in pela_soma) == [
        (100000, ("NF A1",)),
        (200000, ("NF A2",)),
    ]
    assert all(TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA in i.evidencia for i in pela_soma)
    parcial = _do_valor(depois, 30000)
    assert (parcial.destino, parcial.motivo, _documentos(parcial)) == (
        Destino.FILA_DE_EXCECOES,
        Motivo.VALOR_NAO_FECHA,
        ["NF C1"],
    )
    assert parcial.evidencia == (TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA,)
    assert "nosso número" in parcial.explicacao
    pelo_protheus = _do_valor(depois, 90000)
    assert (pelo_protheus.destino, pelo_protheus.evidencia) == (Destino.LISTA_DE_BAIXA, (TipoDeEvidencia.NOSSO_NUMERO,))
    assert not _identificado(_do_valor(depois, 40000))
    assert not _identificado(_do_valor(depois, 33300))
    assert cenario.soltos[Solto.NAO_IDENTIFICAVEL_PELO_VALOR] == TotalDoDestino(1, 40000)
    assert (cenario.soltos[Solto.SEM_TITULO_CONHECIDO].quantidade, cenario.soltos[Solto.FORA_DA_ADESAO].quantidade) == (
        1,
        0,
    )
    chegada = ChegadaDosCreditos.do_resultado(depois)
    assert (chegada.identificados.quantidade, chegada.soltos.quantidade) == (4, 2)
    assert _dinheiro(depois) == _dinheiro(analise.antes)


def test_pix_com_txid_so_pelo_valor_da_cobranca(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [TituloProtheus("NF A1", "1.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A, vencimento=VENCIMENTO)],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO AURORA PANIF", "1.000,00")])

    [cenario] = analisar_impacto(pasta, _configuracao(100), PareceristaFalso()).cenarios

    [item] = cenario.depois.lista_de_baixa
    assert item.lancamento.historico.startswith("PIX RECEBIDO")
    assert "txid" in item.explicacao


def test_valor_a_maior_fica_solto_como_nao_identificavel_pelo_valor(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [TituloProtheus("NF C1", "800,00", cliente="Recanto Feliz Ltda", cnpj=CNPJ_C, vencimento=VENCIMENTO)],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO RECANTO FELIZ", "900,00")])

    analise = analisar_impacto(pasta, _configuracao(100), PareceristaFalso())

    [cenario] = analise.cenarios
    assert cenario.soltos[Solto.NAO_IDENTIFICAVEL_PELO_VALOR].quantidade == 1
    assert [i.lancamento for i in cenario.depois.itens] == [i.lancamento for i in analise.antes.itens]


def test_titulo_ja_pago_por_credito_anterior_deixa_o_seguinte_solto(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [TituloProtheus("NF C1", "800,00", cliente="Recanto Feliz Ltda", cnpj=CNPJ_C, vencimento=VENCIMENTO)],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "PIX RECEBIDO RECANTO FELIZ", "800,00"),
            LinhaItau("06/03/2025", "DEPOSITO EM DINHEIRO", "800,00"),
        ],
    )

    analise = analisar_impacto(pasta, _configuracao(100), PareceristaFalso())

    [cenario] = analise.cenarios
    assert cenario.soltos[Solto.TITULO_SEM_COBRANCA] == TotalDoDestino(1, 80000)
    assert _dinheiro(cenario.depois) == _dinheiro(analise.antes)


def test_gabarito_e_a_primeira_fonte_da_verdade(pasta: Path) -> None:
    _cenario(pasta)
    sem_titulo = LancamentoEsperado(
        conta="1234-5",
        data=date(2025, 3, 6),
        historico="PIX RECEBIDO RECANTO FELIZ",
        documento="",
        valor_centavos=30000,
        natureza=Natureza.CREDITO,
        caso=Caso.ANTECIPACAO,
        destino=Destino.FILA_DE_EXCECOES,
        titulos=frozenset(),
        motivo=Motivo.CREDITO_SEM_TITULO_EM_ABERTO,
    )

    analise = analisar_impacto(pasta, _configuracao(100), PareceristaFalso(), gabarito=[sem_titulo])

    assert analise.fontes == {FonteDaVerdade.GABARITO: 1, FonteDaVerdade.PROPOSTA_OU_SUGESTAO: 4}
    [cenario] = analise.cenarios
    assert cenario.soltos[Solto.SEM_TITULO_CONHECIDO].quantidade == 2
    assert not _identificado(_do_valor(cenario.depois, 30000))


def test_decisoes_confirmadas_do_armazenamento_informado_vem_antes_das_sugestoes(pasta: Path, tmp_path: Path) -> None:
    _cenario(pasta)
    estado = tmp_path / "conciliacao.sqlite"
    armazenamento = Armazenamento(estado)
    primeira = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    parcial = _do_valor(primeira, 30000)
    registrar_decisao_da_tela(
        armazenamento,
        LinhaDeDecisao(
            ABA_DA_FILA_DE_EXCECOES, parcial.identificador, "conciliado fora", franqueado=CNPJ_C, referencia="acordo 7"
        ),
        "rodrigo",
        datetime(2025, 3, 8, 9, 0),
    )
    executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    armazenamento.fechar()
    conteudo = estado.read_bytes()

    analise = analisar_impacto(pasta, _configuracao(100), PareceristaFalso(), estado=estado)

    assert estado.read_bytes() == conteudo
    assert analise.fontes == {FonteDaVerdade.DECISOES_DA_ANALISTA: 1, FonteDaVerdade.PROPOSTA_OU_SUGESTAO: 4}
    [cenario] = analise.cenarios
    assert not _identificado(_do_valor(cenario.depois, 30000))
    assert cenario.soltos[Solto.SEM_TITULO_CONHECIDO].quantidade == 2


@pytest.fixture(scope="module")
def sintetico(tmp_path_factory: pytest.TempPathFactory) -> tuple[AnaliseDeImpacto, tuple[LancamentoEsperado, ...]]:
    dados = gerar_dados_sinteticos(ConfiguracaoDoGerador(semente=20260801, titulos=300, franqueados=40))
    pasta = tmp_path_factory.mktemp("sintetico")
    dados.escrever(pasta)
    nomes = {raiz: nomes for raiz, nomes in dados.nomes_dos_franqueados.items()}
    parecerista = PareceristaFalso(
        sim_para=(
            (esperado.nome_no_extrato, nome)
            for esperado in dados.gabarito
            if esperado.nome_no_extrato and esperado.pagador
            for nome in nomes[esperado.pagador]
        )
    )
    configuracao = Configuracao(
        contas=dados.contas,
        analise_de_impacto=ConfiguracaoDaAnaliseDeImpacto(
            semente=7, faixa_de_adesao=(Decimal(0), Decimal(30), Decimal(60), Decimal(100))
        ),
    )
    return analisar_impacto(pasta, configuracao, parecerista, gabarito=dados.gabarito), dados.gabarito


def test_sintetico_dinheiro_igual_e_fila_que_nunca_cresce_com_a_adesao(
    sintetico: tuple[AnaliseDeImpacto, tuple[LancamentoEsperado, ...]],
) -> None:
    analise, _ = sintetico

    filas = [analise.antes.resumo.totais[Destino.FILA_DE_EXCECOES]]
    identificados: list[set[tuple[object, ...]]] = []
    for cenario in analise.cenarios:
        assert _dinheiro(cenario.depois) == _dinheiro(analise.antes)
        filas.append(cenario.depois.resumo.totais[Destino.FILA_DE_EXCECOES])
        identificados.append(
            {
                (i.lancamento.conta, i.lancamento.data, i.lancamento.historico, i.lancamento.valor_centavos)
                for i in _creditos(cenario.depois)
                if _identificado(i)
            }
        )
    assert [f.quantidade for f in filas] == sorted((f.quantidade for f in filas), reverse=True)
    assert [f.valor_centavos for f in filas] == sorted((f.valor_centavos for f in filas), reverse=True)
    assert filas[-1].quantidade < filas[0].quantidade
    assert all(menor <= maior for menor, maior in zip(identificados, identificados[1:]))
    assert analise.fontes == {FonteDaVerdade.GABARITO: len(_creditos(analise.antes))}


def test_sintetico_com_toda_a_adesao_identifica_tudo_que_pode_e_nao_diverge_do_gabarito(
    sintetico: tuple[AnaliseDeImpacto, tuple[LancamentoEsperado, ...]],
) -> None:
    analise, gabarito = sintetico

    toda = analise.cenarios[-1]
    assert toda.soltos[Solto.FORA_DA_ADESAO].quantidade == 0
    soltos_planejados = sum(total.quantidade for total in toda.soltos.values())
    assert ChegadaDosCreditos.do_resultado(toda.depois).soltos.quantidade == soltos_planejados
    titulos_do_gabarito = {titulo for esperado in gabarito for titulo in esperado.titulos}
    por_chave = {(e.conta, e.data, e.historico, e.valor_centavos): e for e in gabarito}
    for item in toda.depois.lista_de_baixa:
        titulos = {(t.sistema_de_origem, t.documento) for t in item.titulos}
        esperado = por_chave.get(
            (item.lancamento.conta, item.lancamento.data, item.lancamento.historico, item.lancamento.valor_centavos)
        )
        if esperado is not None:
            assert titulos == esperado.titulos
        else:
            assert len(titulos) == 1 and titulos <= titulos_do_gabarito


def test_mesma_semente_mesmo_depois(pasta: Path) -> None:
    _cenario(pasta)

    primeira = analisar_impacto(pasta, _configuracao(50, semente=3), PareceristaFalso())
    segunda = analisar_impacto(pasta, _configuracao(50, semente=3), PareceristaFalso())

    def chaves(analise: AnaliseDeImpacto) -> list[tuple[object, ...]]:
        return [
            (i.lancamento.chave_natural, i.destino, i.motivo, tuple(_documentos(i))) for i in analise.cenarios[0].depois.itens
        ]

    assert chaves(primeira) == chaves(segunda)


def _linhas(caminho: Path, aba: str) -> list[tuple[object, ...]]:
    return [tuple(linha) for linha in load_workbook(caminho)[aba].iter_rows(values_only=True)]


def test_workbook_com_resumo_do_impacto_e_fila_por_motivo_para_cada_adesao(pasta: Path, tmp_path: Path) -> None:
    _cenario(pasta)
    saida = tmp_path / "impacto.xlsx"

    analise = analisar_impacto(pasta, _configuracao(50, 100), PareceristaFalso())
    gerar_workbook_do_impacto(analise, saida)

    assert load_workbook(saida).sheetnames == ["Resumo do impacto", "Ondas", "Fila por motivo", "Ranking de migração", "Premissas"]
    resumo = _linhas(saida, "Resumo do impacto")
    assert resumo[0][:3] == ("Cenário", "Adesão (%)", "Créditos identificados")
    antes, *depois = resumo[1:4]
    assert antes[:7] == ("Antes", None, 1, 900, 4, 4033, None)
    assert antes[8:10] == (3, 1033)
    toda = depois[1]
    assert toda[:10] == ("Depois, todos os Franqueados migrados", 100, 4, 4200, 2, 733, 1, 400, 3, 1033)
    assert toda[10] == pytest.approx(3 / 6)
    fila = _linhas(saida, "Fila por motivo")
    assert fila[0][:3] == ("Motivo de não identificação", "antes (Créditos)", "antes (valor)")
    assert fila[0][5:7] == ("Adesão 100% (Créditos)", "Adesão 100% (valor)")
    assert fila[-1][0] == "Total" and fila[-1][1:3] == (3, 1033) and fila[-1][5:7] == (3, 1033)
    premissas = {linha[0]: linha[1] for linha in _linhas(saida, "Premissas")}
    assert premissas["Faixa de Adesão (%)"] == "50%; 100%"
    assert premissas[f"Créditos pela fonte: {FonteDaVerdade.PROPOSTA_OU_SUGESTAO.value}"] == 5
    assert premissas[f"Créditos pela fonte: {FonteDaVerdade.GABARITO.value}"] == 0


def _escrever_configuracao(caminho: Path, extra: str = "") -> Path:
    caminho.write_text(
        '[[contas]]\npadrao = "itau_cobranca*.csv"\nbanco = "itau"\nconta = "1234-5"\n' + extra, encoding="utf-8"
    )
    return caminho


def test_faixa_de_adesao_na_configuracao(tmp_path: Path) -> None:
    configuracao = carregar_configuracao(
        _escrever_configuracao(tmp_path / "c.toml", "[analise_de_impacto]\nfaixa_de_adesao_percentual = [80, 0, 92.5]\n")
    )

    assert configuracao.analise_de_impacto.faixa_de_adesao == (Decimal("0"), Decimal("80"), Decimal("92.5"))
    assert ConfiguracaoDaAnaliseDeImpacto().faixa_de_adesao == (Decimal("50"), Decimal("80"), Decimal("95"))


@pytest.mark.parametrize("faixa", ["[]", "[101]", "[-1]", "[50, 50]", '["50"]', "[true]", "50"])
def test_faixa_de_adesao_invalida(tmp_path: Path, faixa: str) -> None:
    with pytest.raises(ConfiguracaoInvalida, match="faixa_de_adesao_percentual"):
        carregar_configuracao(
            _escrever_configuracao(tmp_path / "c.toml", f"[analise_de_impacto]\nfaixa_de_adesao_percentual = {faixa}\n")
        )


def test_comando_le_as_decisoes_do_armazenamento_informado(
    pasta: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    _cenario(pasta)
    configuracao = _escrever_configuracao(tmp_path / "c.toml", "[analise_de_impacto]\nfaixa_de_adesao_percentual = [50]\n")
    estado = tmp_path / "conciliacao.sqlite"
    armazenamento = Armazenamento(estado)
    executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    armazenamento.fechar()
    saida = tmp_path / "impacto.xlsx"

    assert main([str(pasta), "--configuracao", str(configuracao), "--saida", str(saida), "--estado", str(estado)]) == 0
    assert load_workbook(saida).sheetnames[0] == "Resumo do impacto"
    assert main([str(pasta), "--configuracao", str(configuracao), "--estado", str(tmp_path / "falta.sqlite")]) == 2
    assert "falta.sqlite" in capsys.readouterr().err
