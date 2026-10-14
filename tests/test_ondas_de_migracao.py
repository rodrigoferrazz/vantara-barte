from dataclasses import replace
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.configuracao import Configuracao, ConfiguracaoDaAnaliseDeImpacto
from conciliacao.dominio import Destino, Item, PorQueFicaSolto, Resultado, TipoDeEvidencia, TotalDoDestino
from conciliacao.gerador import ConfiguracaoDoGerador, gerar_dados_sinteticos
from conciliacao.impacto import (
    AnaliseDeImpacto,
    CenarioDeAdesao,
    analisar_impacto,
    gerar_workbook_do_impacto,
)
from tests.arquivos import LinhaItau, TituloProtheus, configuracao_itau, escrever_extrato_itau, escrever_protheus
from tests.parecerista import PareceristaFalso

CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.400.600/0001-80"
RAIZ_A = "10200300"
RAIZ_B = "20400600"
VENCIMENTO = "10/03/2025"


def _cenario(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus("NF A1", "1.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A, vencimento=VENCIMENTO),
            TituloProtheus("NF A2", "2.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A, vencimento=VENCIMENTO),
            TituloProtheus("NF B1", "500,00", cliente="Brisa Serena Armazem Ltda", cnpj=CNPJ_B, vencimento=VENCIMENTO),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "DEPOSITO EM DINHEIRO", "1.000,00"),
            LinhaItau("06/03/2025", "DEPOSITO EM DINHEIRO", "2.000,00"),
            LinhaItau("07/03/2025", "DEPOSITO EM DINHEIRO", "500,00"),
        ],
    )


def _configuracao(*adesoes: int, tamanho_da_onda: int = 1) -> Configuracao:
    return replace(
        configuracao_itau(),
        analise_de_impacto=ConfiguracaoDaAnaliseDeImpacto(
            tamanho_da_onda=tamanho_da_onda, faixa_de_adesao=tuple(Decimal(adesao) for adesao in adesoes)
        ),
    )


def _creditos(resultado: Resultado) -> list[Item]:
    return [*resultado.lista_de_baixa, *resultado.fila_de_excecoes]


def _identificado(item: Item) -> bool:
    return bool({TipoDeEvidencia.NOSSO_NUMERO, TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA} & set(item.evidencia))


def _identificados(cenario: CenarioDeAdesao) -> set[tuple[object, ...]]:
    return {
        (i.lancamento.conta, i.lancamento.data, i.lancamento.historico, i.lancamento.valor_centavos)
        for i in _creditos(cenario.depois)
        if _identificado(i)
    }


def _fila(resultado: Resultado) -> TotalDoDestino:
    return resultado.resumo.totais[Destino.FILA_DE_EXCECOES]


def test_cada_onda_acumulada_so_identifica_os_creditos_dos_franqueados_ja_migrados(pasta: Path) -> None:
    _cenario(pasta)

    analise = analisar_impacto(pasta, _configuracao(100), PareceristaFalso())

    assert [(onda.onda, onda.franqueados) for onda in analise.ondas] == [
        (1, frozenset({RAIZ_A})),
        (2, frozenset({RAIZ_A, RAIZ_B})),
    ]
    [primeira] = analise.ondas[0].cenarios
    assert sorted(i.lancamento.valor_centavos for i in _creditos(primeira.depois) if _identificado(i)) == [100000, 200000]
    assert primeira.soltos[PorQueFicaSolto.FRANQUEADO_NAO_MIGRADO] == TotalDoDestino(1, 50000)
    assert _fila(primeira.depois) == TotalDoDestino(1, 50000)
    [segunda] = analise.ondas[1].cenarios
    assert all(_identificado(i) for i in _creditos(segunda.depois))
    assert segunda.soltos[PorQueFicaSolto.FRANQUEADO_NAO_MIGRADO] == TotalDoDestino(0, 0)
    [todos] = analise.cenarios
    assert todos.soltos[PorQueFicaSolto.FRANQUEADO_NAO_MIGRADO] == TotalDoDestino(0, 0)


def test_cada_adesao_tem_a_sua_linha_em_cada_onda(pasta: Path) -> None:
    _cenario(pasta)

    analise = analisar_impacto(pasta, _configuracao(0, 100, tamanho_da_onda=2), PareceristaFalso())

    [onda] = analise.ondas
    assert [cenario.adesao for cenario in onda.cenarios] == [Decimal(0), Decimal(100)]
    assert _fila(onda.cenarios[0].depois) == _fila(analise.antes)
    assert _fila(onda.cenarios[1].depois) == TotalDoDestino(0, 0)


@pytest.fixture(scope="module")
def sintetico(tmp_path_factory: pytest.TempPathFactory) -> AnaliseDeImpacto:
    dados = gerar_dados_sinteticos(ConfiguracaoDoGerador(semente=20260801, titulos=300, franqueados=40))
    pasta = tmp_path_factory.mktemp("sintetico")
    dados.escrever(pasta)
    parecerista = PareceristaFalso(
        sim_para=(
            (esperado.nome_no_extrato, nome)
            for esperado in dados.gabarito
            if esperado.nome_no_extrato and esperado.pagador
            for nome in dados.nomes_dos_franqueados[esperado.pagador]
        )
    )
    configuracao = Configuracao(
        contas=dados.contas,
        analise_de_impacto=ConfiguracaoDaAnaliseDeImpacto(
            semente=7, tamanho_da_onda=6, faixa_de_adesao=(Decimal(0), Decimal(50), Decimal(100))
        ),
    )
    return analisar_impacto(pasta, configuracao, parecerista, gabarito=dados.gabarito)


def test_sintetico_so_identifica_creditos_de_franqueados_migrados_em_cada_onda(sintetico: AnaliseDeImpacto) -> None:
    ja_identificados = {i.lancamento.chave_natural for i in _creditos(sintetico.antes) if _identificado(i)}

    assert len(sintetico.ondas) > 1
    for onda in sintetico.ondas:
        for cenario in onda.cenarios:
            for item in _creditos(cenario.depois):
                if _identificado(item) and item.lancamento.chave_natural not in ja_identificados:
                    assert {titulo.franqueado for titulo in item.titulos} <= onda.franqueados


def test_sintetico_fila_nunca_aumenta_com_a_onda_ou_com_a_adesao(sintetico: AnaliseDeImpacto) -> None:
    linhas = [*(onda.cenarios for onda in sintetico.ondas), sintetico.cenarios]
    antes = _fila(sintetico.antes)

    for cenarios in linhas:
        filas = [antes, *(_fila(c.depois) for c in cenarios)]
        assert [f.quantidade for f in filas] == sorted((f.quantidade for f in filas), reverse=True)
        assert [f.valor_centavos for f in filas] == sorted((f.valor_centavos for f in filas), reverse=True)
    for coluna in range(len(sintetico.cenarios)):
        filas = [antes, *(_fila(cenarios[coluna].depois) for cenarios in linhas)]
        assert [f.quantidade for f in filas] == sorted((f.quantidade for f in filas), reverse=True)
        assert [f.valor_centavos for f in filas] == sorted((f.valor_centavos for f in filas), reverse=True)
        identificados = [_identificados(cenarios[coluna]) for cenarios in linhas]
        assert all(menor <= maior for menor, maior in zip(identificados, identificados[1:]))
    assert _fila(sintetico.ondas[-1].cenarios[-1].depois).quantidade < antes.quantidade


def _linhas(caminho: Path, aba: str) -> list[tuple[object, ...]]:
    return [tuple(linha) for linha in load_workbook(caminho)[aba].iter_rows(values_only=True)]


def test_aba_ondas_com_a_fila_e_o_resolvido_sozinho_por_onda_e_adesao(pasta: Path, tmp_path: Path) -> None:
    _cenario(pasta)
    saida = tmp_path / "impacto.xlsx"

    gerar_workbook_do_impacto(analisar_impacto(pasta, _configuracao(0, 100), PareceristaFalso()), saida)

    assert load_workbook(saida).sheetnames == [
        "Resumo do impacto",
        "Ondas",
        "Fila por motivo",
        "Ranking de migração",
        "Premissas",
    ]
    ondas = _linhas(saida, "Ondas")
    assert ondas[0] == (
        "Onda acumulada",
        "Franqueados migrados",
        "Adesão (%)",
        "Créditos identificados",
        "Créditos de Franqueados não migrados",
        "Créditos na Fila",
        "Valor na Fila",
        "% resolvido sozinho (quantidade)",
        "% resolvido sozinho (valor)",
    )
    assert ondas[1][:7] == ("Antes", 0, None, 0, None, 3, 3500)
    assert [linha[:7] for linha in ondas[2:8]] == [
        (1, 1, 0, 0, 1, 3, 3500),
        (1, 1, 100, 2, 1, 1, 500),
        (2, 2, 0, 0, 0, 3, 3500),
        (2, 2, 100, 3, 0, 0, 0),
        ("Todos os Franqueados", 2, 0, 0, 0, 3, 3500),
        ("Todos os Franqueados", 2, 100, 3, 0, 0, 0),
    ]
    assert ondas[3][7] == pytest.approx(2 / 3)
    assert ondas[1][7] == 0
