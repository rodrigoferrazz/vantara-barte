import threading
import urllib.error
import urllib.request
from collections.abc import Iterator
from dataclasses import replace
from decimal import Decimal
from fractions import Fraction
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.configuracao import ConfiguracaoDaAnaliseDeImpacto
from conciliacao.execucao import executar_conciliacao
from conciliacao.impacto import (
    ABA_DA_FILA_POR_MOTIVO,
    ABA_DAS_ONDAS,
    ABA_DE_PREMISSAS,
    ABA_DO_RANKING,
    ABA_DO_RESUMO,
    analisar_impacto,
    main,
)
from conciliacao.normalizacao import formatar_adesao, formatar_percentual, formatar_reais
from conciliacao.tela import criar_servidor
from tests.arquivos import LinhaItau, TituloProtheus, configuracao_itau, escrever_extrato_itau, escrever_protheus
from tests.pagina import Elemento, ler_pagina
from tests.parecerista import PareceristaFalso

CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.400.600/0001-80"
CNPJ_C = "30.600.900/0001-10"
CNPJ_D = "40.800.100/0001-20"
VENCIMENTO = "10/03/2025"


@pytest.fixture
def estado(tmp_path: Path) -> Path:
    return tmp_path / "estado.sqlite"


@pytest.fixture
def endereco(estado: Path) -> Iterator[str]:
    servidor = criar_servidor(estado, porta=0)
    linha = threading.Thread(target=servidor.serve_forever, daemon=True)
    linha.start()
    try:
        yield f"http://127.0.0.1:{servidor.server_address[1]}"
    finally:
        servidor.shutdown()
        servidor.server_close()


def _html(endereco: str, caminho: str = "/impacto") -> str:
    with urllib.request.urlopen(endereco + caminho) as resposta:
        assert resposta.status == 200
        assert resposta.headers.get_content_type() == "text/html"
        texto: str = resposta.read().decode("utf-8")
        return texto


def _abrir(endereco: str, caminho: str = "/impacto") -> Elemento:
    return ler_pagina(_html(endereco, caminho))


def _cenario(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus("NF A1", "1.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A, vencimento=VENCIMENTO),
            TituloProtheus("NF A2", "2.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A, vencimento=VENCIMENTO),
            TituloProtheus("NF B1", "500,00", cliente="Brisa Serena Armazem Ltda", cnpj=CNPJ_B, vencimento=VENCIMENTO),
            TituloProtheus("NF C1", "800,00", cliente="Recanto Feliz Ltda", cnpj=CNPJ_C, vencimento=VENCIMENTO),
            TituloProtheus("NF D1", "700,00", cliente="Doce Lar Mercearia Ltda", cnpj=CNPJ_D, vencimento=VENCIMENTO),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "TED RECEBIDA CNPJ 10200300000140", "3.000,00"),
            LinhaItau("05/03/2025", "PIX RECEBIDO BRISA SERENA ARM", "500,00"),
            LinhaItau("06/03/2025", "PIX RECEBIDO RECANTO FELIZ", "300,00"),
            LinhaItau("06/03/2025", "PIX RECEBIDO DOCE LAR MERC", "700,00"),
            LinhaItau("07/03/2025", "DEPOSITO EM DINHEIRO", "333,00"),
        ],
    )


def _configuracao(caminho: Path, semente: int = 1, faixa: str = "[50, 100]") -> Path:
    caminho.write_text(
        '[[contas]]\npadrao = "itau_cobranca*.csv"\nbanco = "itau"\nconta = "1234-5"\n'
        f"[analise_de_impacto]\nsemente = {semente}\ntamanho_da_onda = 1\nfaixa_de_adesao_percentual = {faixa}\n",
        encoding="utf-8",
    )
    return caminho


def _analisar(pasta: Path, tmp_path: Path, estado: Path, semente: int = 1, faixa: str = "[50, 100]") -> Path:
    saida = tmp_path / f"impacto_{semente}.xlsx"
    configuracao = _configuracao(tmp_path / "c.toml", semente, faixa)
    assert main([str(pasta), "--configuracao", str(configuracao), "--saida", str(saida), "--estado", str(estado)]) == 0
    return saida


def _linhas_do_workbook(caminho: Path, aba: str) -> list[tuple[object, ...]]:
    linhas: list[tuple[object, ...]] = []
    for linha in load_workbook(caminho)[aba].iter_rows(values_only=True):
        if all(celula is None for celula in linha):
            break
        linhas.append(tuple(linha))
    return linhas


def _como_na_pagina(coluna: str, valor: object) -> str:
    if valor is None:
        return ""
    numero = Decimal(str(valor)) if isinstance(valor, int | float) else None
    if numero is None:
        return str(valor)
    if coluna.startswith("%"):
        return formatar_percentual(Fraction(numero * 100))
    if coluna.startswith("Valor") or coluna.endswith("(valor)"):
        return formatar_reais(int(numero * 100))
    if coluna == "Adesão (%)":
        return formatar_adesao(numero)
    return str(valor)


def _tabela_do_workbook(caminho: Path, aba: str) -> list[list[str]]:
    cabecalho, *linhas = _linhas_do_workbook(caminho, aba)
    return [[str(c) for c in cabecalho]] + [
        [_como_na_pagina(str(coluna), valor) for coluna, valor in zip(cabecalho, linha)] for linha in linhas
    ]


def test_sem_analise_gravada_a_pagina_diz_que_nao_ha_analise_sem_inventar_dados(endereco: str, estado: Path) -> None:
    pagina = _abrir(endereco)

    assert "Não há análise de impacto" in pagina.por_id("sem-analise").texto
    assert pagina.todos("table") == [] and pagina.todos("svg") == []
    assert not estado.exists()


def test_armazenamento_so_com_execucao_diaria_tambem_diz_que_nao_ha_analise(
    endereco: str, pasta: Path, estado: Path
) -> None:
    _cenario(pasta)
    armazenamento = Armazenamento(estado)
    executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    armazenamento.fechar()

    pagina = _abrir(endereco)

    assert pagina.tem_id("sem-analise")
    assert pagina.todos("table") == []


def test_a_pagina_mostra_os_mesmos_numeros_do_workbook_da_mesma_analise(
    endereco: str, pasta: Path, estado: Path, tmp_path: Path
) -> None:
    _cenario(pasta)
    workbook = _analisar(pasta, tmp_path, estado)

    pagina = _abrir(endereco)

    cabecalho, *cenarios = _tabela_do_workbook(workbook, ABA_DO_RESUMO)
    antes_e_depois = pagina.por_id("antes-e-depois").linhas()
    assert antes_e_depois[0] == ["Medida", "Antes", "Adesão 50%", "Adesão 100%"]
    assert antes_e_depois[1:] == [[coluna, *(c[j] for c in cenarios)] for j, coluna in enumerate(cabecalho) if j >= 2]
    assert len(cenarios) == 3 and cenarios[0][0] == "Antes"
    assert pagina.por_id("ondas").linhas() == _tabela_do_workbook(workbook, ABA_DAS_ONDAS)
    assert pagina.por_id("fila-por-motivo").linhas() == _tabela_do_workbook(workbook, ABA_DA_FILA_POR_MOTIVO)
    assert pagina.por_id("ranking").linhas() == _tabela_do_workbook(workbook, ABA_DO_RANKING)
    premissas = pagina.por_id("premissas").linhas()
    do_workbook = _linhas_do_workbook(workbook, ABA_DE_PREMISSAS)
    assert [p[0] for p in premissas] == [str(p[0]) for p in do_workbook]
    for (_, na_pagina), (_, no_workbook) in zip(premissas[1:], do_workbook[1:]):
        if isinstance(no_workbook, str):
            assert na_pagina == no_workbook
        else:
            assert Decimal(na_pagina.replace(",", ".")) == Decimal(str(no_workbook))


def test_premissas_visiveis_com_fonte_da_verdade_semente_proporcoes_encargo_e_onda(
    endereco: str, pasta: Path, estado: Path, tmp_path: Path
) -> None:
    _cenario(pasta)
    _analisar(pasta, tmp_path, estado, semente=7)

    pagina = _abrir(endereco)

    premissas = dict((nome, valor) for nome, valor in pagina.por_id("premissas").linhas()[1:])
    assert premissas["Semente"] == "7"
    assert premissas["Tamanho da Onda de migração (Franqueados)"] == "1"
    assert premissas["Multa por atraso (%)"] == "2"
    assert "Tolerância do Encargo (R$)" in premissas and "Juros ao mês (%)" in premissas
    assert premissas["Proporções do simulador"].startswith("não se aplicam")
    assert premissas["Créditos pela fonte: Proposta da Lista ou sugestão da Fila do antes"] == "5"
    assert premissas["Limite da fonte da verdade"].startswith("5 de 5 Créditos")
    assert f"Pasta base {pasta.resolve()}, semente 7" in pagina.por_id("analise").texto


def test_a_pagina_le_a_ultima_analise_gravada_sem_recalcular(
    endereco: str, pasta: Path, estado: Path, tmp_path: Path
) -> None:
    _cenario(pasta)
    _analisar(pasta, tmp_path, estado, semente=3, faixa="[80]")
    _analisar(pasta, tmp_path, estado, semente=5, faixa="[0, 95]")
    for arquivo in pasta.iterdir():
        arquivo.unlink()
    (tmp_path / "c.toml").unlink()

    pagina = _abrir(endereco)

    assert pagina.por_id("antes-e-depois").linhas()[0] == ["Medida", "Antes", "Adesão 0%", "Adesão 95%"]
    assert "semente 5" in pagina.por_id("analise").texto


def test_graficos_da_fila_por_adesao_e_por_onda_desenhados_no_servidor_sem_recurso_externo(
    endereco: str, pasta: Path, estado: Path, tmp_path: Path
) -> None:
    _cenario(pasta)
    _analisar(pasta, tmp_path, estado)

    html = _html(endereco)

    pagina = ler_pagina(html)
    resumo = {linha[0]: linha[1:] for linha in pagina.por_id("antes-e-depois").linhas()}
    por_adesao = pagina.por_id("grafico-fila-por-adesao")
    assert por_adesao.tag == "svg"
    barras = por_adesao.todos("g", classe="barra")
    assert [b.texto for b in barras] == [
        f"{nome} {quantidade} Crédito(s), {valor}"
        for nome, quantidade, valor in zip(
            pagina.por_id("antes-e-depois").linhas()[0][1:], resumo["Créditos na Fila"], resumo["Valor na Fila"]
        )
    ]
    ondas = pagina.por_id("ondas").linhas()[1:]
    quantas_ondas = len({linha[0] for linha in ondas} - {"Antes", "Todos os Franqueados"})
    assert quantas_ondas >= 2
    for posicao, adesao in enumerate(("50%", "100%")):
        grafico = pagina.por_id(f"grafico-fila-por-onda-{posicao}")
        textos = [b.texto for b in grafico.todos("g", classe="barra")]
        assert len(textos) == 1 + quantas_ondas + 1
        assert textos[0].startswith("Antes") and textos[1].startswith("Onda 1 (1 Franqueado(s))")
        assert textos[-1].startswith("Todos os Franqueados")
        da_adesao = [linha for linha in ondas if linha[2] == adesao]
        fila = [f"{linha[5]} Crédito(s), {linha[6]}" for linha in da_adesao]
        assert len(fila) == len(textos[1:]) and all(t.endswith(f) for t, f in zip(textos[1:], fila))
    assert pagina.todos("script") == [] and pagina.todos("link") == [] and pagina.todos("img") == []
    assert "url(" not in html and "@import" not in html and "src=" not in html


def test_a_pagina_e_so_leitura(endereco: str, pasta: Path, estado: Path, tmp_path: Path) -> None:
    _cenario(pasta)
    _analisar(pasta, tmp_path, estado)

    pagina = _abrir(endereco)

    assert pagina.todos("form") == [] and pagina.todos("button") == []
    with pytest.raises(urllib.error.HTTPError) as erro:
        urllib.request.urlopen(urllib.request.Request(endereco + "/impacto", data=b"", method="POST"))
    assert erro.value.code == 404


def test_a_tela_da_execucao_da_acesso_a_pagina_de_impacto_e_volta(endereco: str) -> None:
    inicial = _abrir(endereco, "/")
    impacto = _abrir(endereco, "/impacto")

    assert any(a.atributos.get("href") == "/impacto" for a in inicial.por_id("navegacao").todos("a"))
    assert any(a.atributos.get("href") == "/" for a in impacto.por_id("navegacao").todos("a"))


def test_a_analise_gravada_volta_igual_do_armazenamento_sem_mudar_a_execucao_diaria(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    _cenario(pasta)
    diaria = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    configuracao = replace(
        configuracao_itau(),
        analise_de_impacto=ConfiguracaoDaAnaliseDeImpacto(tamanho_da_onda=1, faixa_de_adesao=(Decimal("92.5"),)),
    )
    relatorio = analisar_impacto(pasta, configuracao, PareceristaFalso()).relatorio

    assert armazenamento.ultima_analise_de_impacto() is None
    with armazenamento.transacao():
        armazenamento.registrar_analise_de_impacto(relatorio)

    assert armazenamento.ultima_analise_de_impacto() == relatorio
    assert armazenamento.ultima_execucao() == diaria.resumo.execucao
    assert armazenamento.carregar_resultado(diaria.resumo.execucao) == diaria


def test_comando_cria_o_armazenamento_que_nao_existe_e_grava_a_analise(
    pasta: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    _cenario(pasta)
    estado = tmp_path / "novo.sqlite"

    _analisar(pasta, tmp_path, estado)

    assert "novo.sqlite" in capsys.readouterr().out
    armazenamento = Armazenamento(estado)
    try:
        analise = armazenamento.ultima_analise_de_impacto()
        assert analise is not None and analise.semente == 1
        assert armazenamento.ultima_execucao() is None
    finally:
        armazenamento.fechar()
