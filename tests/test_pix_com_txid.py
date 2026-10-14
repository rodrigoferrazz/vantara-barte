import json
import re
from dataclasses import asdict, replace
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.configuracao import (
    Configuracao,
    ConfiguracaoDeMigracao,
    ConfiguracaoInvalida,
    ContaDeExtrato,
    carregar_configuracao,
)
from conciliacao.dominio import Destino, MotivoDeNaoIdentificacao, Resultado, TipoDeEvidencia
from conciliacao.emissao import emitir_cobrancas, main
from conciliacao.execucao import executar_conciliacao
from conciliacao.porta import CobrancaIdentificada, PortaSimulada
from conciliacao.saidas import gerar_workbook
from tests.arquivos import (
    LinhaItau,
    TituloOmie,
    TituloProtheus,
    configuracao_itau,
    escrever_extrato_itau,
    escrever_omie,
    escrever_protheus,
)
from tests.parecerista import PareceristaFalso

MIGRADO = "10200300"
CNPJ_DE_OUTRO_FRANQUEADO = "20.300.400/0001-50"
REFERENCIA = LinhaItau("05/03/2025", "TARIFA MANUTENCAO CONTA", "9,90", tipo="D")
_TXID = re.compile(r"[a-zA-Z0-9]{26,35}")


def _configuracao(migrados: tuple[str, ...] | None = (MIGRADO,), conta_do_pix: str = "1234-5") -> Configuracao:
    base = configuracao_itau()
    return replace(
        base,
        contas=(*base.contas, ContaDeExtrato(padrao="pix_recebidos*.json", banco="itau", conta=conta_do_pix)),
        migracao=None if migrados is None else ConfiguracaoDeMigracao(frozenset(migrados)),
    )


def _pix(txid: str | None, valor: str, horario: str = "2025-03-05T14:30:00.000Z", fim_do_e2e: str = "1") -> dict[str, object]:
    recebido: dict[str, object] = {
        "endToEndId": f"E00000000202503051430{fim_do_e2e:0>11}",
        "valor": valor,
        "horario": horario,
        "chave": "chave-pix-do-recebedor",
        "infoPagador": "pagamento",
    }
    if txid is not None:
        recebido["txid"] = txid
    return recebido


def _escrever_pix(pasta: Path, recebidos: list[dict[str, object]], nome: str = "pix_recebidos.json") -> Path:
    caminho = pasta / nome
    conteudo = {
        "parametros": {
            "inicio": "2025-03-01T00:00:00Z",
            "fim": "2025-03-31T23:59:59Z",
            "paginacao": {"paginaAtual": 0, "itensPorPagina": 100, "quantidadeDePaginas": 1, "quantidadeTotalDeItens": len(recebidos)},
        },
        "pix": recebidos,
    }
    caminho.write_text(json.dumps(conteudo), encoding="utf-8")
    return caminho


def _conciliar(pasta: Path, armazenamento: Armazenamento, configuracao: Configuracao) -> Resultado:
    return executar_conciliacao(pasta, armazenamento, configuracao, PareceristaFalso())


def _emitir(
    pasta: Path, armazenamento: Armazenamento, titulos: list[TituloProtheus], configuracao: Configuracao | None = None
) -> PortaSimulada:
    configuracao = configuracao or _configuracao()
    escrever_protheus(pasta, titulos)
    escrever_extrato_itau(pasta, [REFERENCIA])
    _conciliar(pasta, armazenamento, configuracao)
    porta = PortaSimulada()
    emitir_cobrancas(armazenamento, configuracao, porta)
    return porta


def _txids(porta: PortaSimulada) -> dict[str, str]:
    return {cobranca.documento: cobranca.txid for cobranca in porta.emitidas}


def test_sem_configuracao_de_migracao_nada_da_fase_2_executa(pasta: Path, armazenamento: Armazenamento) -> None:
    sem_migracao = _configuracao(migrados=None)
    porta = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00")], sem_migracao)
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("05/03/2025", "PIX RECEBIDO FULANO", "1.000,00")])
    _escrever_pix(pasta, [_pix("CI000000000000000000000000000001", "1000.00")])

    resultado = _conciliar(pasta, armazenamento, sem_migracao)

    assert porta.emitidas == []
    assert [recusado.nome for recusado in resultado.arquivos_recusados] == ["pix_recebidos.json"]
    [item] = resultado.fila_de_excecoes
    assert TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA not in item.evidencia
    assert resultado.lista_de_baixa == []


def test_emissao_so_considera_titulos_em_aberto_de_migrados_sem_cobranca_e_nao_vencidos(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    titulos = [
        TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="05/03/2025"),
        TituloProtheus(documento="NF 2", valor="200,00", vencimento="04/03/2025"),
        TituloProtheus(documento="NF 3", valor="300,00", vencimento="28/02/2025"),
        TituloProtheus(documento="NF 4", valor="400,00", vencimento="10/03/2025", cnpj=CNPJ_DE_OUTRO_FRANQUEADO),
        TituloProtheus(documento="NF 5", valor="500,00", vencimento="10/03/2025"),
    ]

    primeira = _emitir(pasta, armazenamento, titulos)
    segunda = PortaSimulada()
    emitir_cobrancas(armazenamento, _configuracao(), segunda)

    assert sorted(_txids(primeira)) == ["NF 1", "NF 2", "NF 5"]
    assert segunda.emitidas == []


def test_emissao_deixa_de_fora_titulos_em_possivel_duplicidade_ate_a_analista_decidir(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_omie(pasta, [TituloOmie(documento="FAT-1", valor="800.00", vencimento="2025-03-10")])
    primeira = _emitir(
        pasta,
        armazenamento,
        [
            TituloProtheus(documento="NF 1", valor="800,00", vencimento="10/03/2025"),
            TituloProtheus(documento="NF 2", valor="500,00", vencimento="10/03/2025"),
        ],
    )
    workbook = tmp_path / "conciliacao.xlsx"
    gerar_workbook(_conciliar(pasta, armazenamento, _configuracao()), workbook)
    planilha = load_workbook(workbook)
    fila = planilha["Fila de exceções"]
    cabecalho = [celula.value for celula in fila[1]]
    fila.cell(row=2, column=cabecalho.index("Decisão") + 1, value="dívidas distintas")
    planilha.properties.lastModifiedBy = "Ana Analista"
    planilha.save(workbook)
    executar_conciliacao(pasta, armazenamento, _configuracao(), PareceristaFalso(), workbook_anterior=workbook)
    depois_da_decisao = PortaSimulada()
    emitir_cobrancas(armazenamento, _configuracao(), depois_da_decisao)

    assert sorted(_txids(primeira)) == ["NF 2"]
    assert sorted(_txids(depois_da_decisao)) == ["FAT-1", "NF 1"]


def test_emissao_desconta_do_valor_as_baixas_confirmadas_e_ignora_titulo_quitado(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025", nosso_numero="0001")])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "1.000,00", documento="0001")])
    configuracao = _configuracao()
    workbook = pasta.parent / "conciliacao.xlsx"
    gerar_workbook(_conciliar(pasta, armazenamento, configuracao), workbook)
    planilha = load_workbook(workbook)
    lista = planilha["Lista de baixa"]
    cabecalho = [celula.value for celula in lista[1]]
    lista.cell(row=2, column=cabecalho.index("Decisão") + 1, value="confirmar")
    planilha.properties.lastModifiedBy = "Ana Analista"
    planilha.save(workbook)
    executar_conciliacao(pasta, armazenamento, configuracao, PareceristaFalso(), workbook_anterior=workbook)

    porta = PortaSimulada()
    emitir_cobrancas(armazenamento, configuracao, porta)

    assert porta.emitidas == []


def test_txid_alfanumerico_de_26_a_35_caracteres_unico_e_nunca_reaproveitado(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    primeira = _emitir(
        pasta,
        armazenamento,
        [TituloProtheus(documento=f"NF {numero}", valor="100,00", vencimento="10/03/2025") for numero in range(1, 4)],
    )
    escrever_protheus(pasta, [TituloProtheus(documento="NF 9", valor="100,00", vencimento="10/03/2025")])
    _conciliar(pasta, armazenamento, _configuracao())
    segunda = PortaSimulada()
    emitir_cobrancas(armazenamento, _configuracao(), segunda)

    txids = [cobranca.txid for cobranca in (*primeira.emitidas, *segunda.emitidas)]
    assert len(txids) == 4
    assert all(_TXID.fullmatch(txid) for txid in txids)
    assert len(set(txids)) == 4


class _PortaFora:
    def __init__(self) -> None:
        self.tentadas: list[CobrancaIdentificada] = []

    def emitir(self, cobranca: CobrancaIdentificada) -> None:
        self.tentadas.append(cobranca)
        raise ConnectionError("porta fora do ar")


def test_vinculo_e_gravado_antes_de_chamar_a_porta_e_emitir_de_novo_nao_duplica(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    escrever_extrato_itau(pasta, [REFERENCIA])
    _conciliar(pasta, armazenamento, _configuracao())
    fora = _PortaFora()

    with pytest.raises(ConnectionError):
        emitir_cobrancas(armazenamento, _configuracao(), fora)
    [tentada] = fora.tentadas
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "1.000,00")])
    _escrever_pix(pasta, [_pix(tentada.txid, "1000.00", horario="2025-03-06T10:00:00Z")])
    identificado = _conciliar(pasta, armazenamento, _configuracao())
    depois = PortaSimulada()
    emitir_cobrancas(armazenamento, _configuracao(), depois)
    de_novo = PortaSimulada()
    emitir_cobrancas(armazenamento, _configuracao(), de_novo)

    assert [item.titulos[0].documento for item in identificado.lista_de_baixa] == ["NF 1"]
    assert depois.emitidas == [tentada]
    assert de_novo.emitidas == []


def test_porta_recebe_so_titulo_valor_vencimento_cnpj_e_identificador(pasta: Path, armazenamento: Armazenamento) -> None:
    porta = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.234,56", vencimento="10/03/2025")])

    [cobranca] = porta.emitidas
    assert asdict(cobranca) == {
        "sistema_de_origem": "Protheus",
        "documento": "NF 1",
        "valor_centavos": 123456,
        "vencimento": date(2025, 3, 10),
        "cnpj": "10200300000140",
        "txid": cobranca.txid,
    }


def test_pix_com_txid_e_valor_exato_vai_para_a_lista_com_o_tipo_e_a_origem_do_identificador(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    porta = _emitir(
        pasta,
        armazenamento,
        [
            TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025"),
            TituloProtheus(documento="NF 2", valor="1.000,00", vencimento="12/03/2025"),
        ],
    )
    txid = _txids(porta)["NF 2"]
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("06/03/2025", "PIX RECEBIDO SOCIO DESCONHECIDO", "1.000,00")])
    _escrever_pix(pasta, [_pix(txid, "1000.00", horario="2025-03-06T10:00:00.000Z")])

    resultado = _conciliar(pasta, armazenamento, _configuracao())

    [item] = resultado.lista_de_baixa
    assert [titulo.documento for titulo in item.titulos] == ["NF 2"]
    assert item.evidencia == (TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA,)
    assert f"txid {txid}" in item.explicacao
    assert "plataforma" in item.explicacao
    assert resultado.fila_de_excecoes == []
    assert resultado.arquivos_recusados == ()


def test_pix_com_txid_e_valor_que_nao_fecha_vai_para_a_fila_apontando_o_titulo(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    porta = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "900,00")])
    _escrever_pix(pasta, [_pix(_txids(porta)["NF 1"], "900.00", horario="2025-03-06T10:00:00Z")])

    resultado = _conciliar(pasta, armazenamento, _configuracao())

    [item] = resultado.fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.VALOR_NAO_FECHA
    assert [titulo.documento for titulo in item.titulos] == ["NF 1"]
    assert item.evidencia == (TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA,)


def test_data_do_pix_e_a_de_brasilia_mesmo_quando_o_horario_vem_em_utc(pasta: Path, armazenamento: Armazenamento) -> None:
    porta = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "1.000,00")])
    _escrever_pix(pasta, [_pix(_txids(porta)["NF 1"], "1000.00", horario="2025-03-07T01:30:00Z")])

    resultado = _conciliar(pasta, armazenamento, _configuracao())

    assert [item.titulos[0].documento for item in resultado.lista_de_baixa] == ["NF 1"]


def test_retorno_sem_credito_no_extrato_nao_cria_credito(pasta: Path, armazenamento: Armazenamento) -> None:
    porta = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    _escrever_pix(pasta, [_pix(_txids(porta)["NF 1"], "1000.00")])

    resultado = _conciliar(pasta, armazenamento, _configuracao())

    assert resultado.lista_de_baixa == []
    assert resultado.fila_de_excecoes == []
    assert resultado.resumo.totais[Destino.LISTA_DE_BAIXA].quantidade == 0
    assert resultado.arquivos_recusados == ()


def test_retorno_so_enriquece_credito_com_par_unico_de_conta_data_e_valor(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    porta = _emitir(
        pasta,
        armazenamento,
        [
            TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025"),
            TituloProtheus(documento="NF 2", valor="700,00", vencimento="10/03/2025"),
        ],
        _configuracao(),
    )
    txids = _txids(porta)
    escrever_extrato_itau(
        pasta,
        [
            REFERENCIA,
            LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "1.000,00"),
            LinhaItau("06/03/2025", "PIX RECEBIDO BELTRANO", "1.000,00"),
            LinhaItau("06/03/2025", "PIX RECEBIDO CICLANO", "700,00"),
        ],
    )
    _escrever_pix(pasta, [_pix(txids["NF 1"], "1000.00", horario="2025-03-06T10:00:00Z")])
    _escrever_pix(
        pasta, [_pix(txids["NF 2"], "700.00", horario="2025-03-06T10:00:00Z", fim_do_e2e="2")], nome="pix_recebidos_outra_conta.json"
    )
    configuracao = replace(
        _configuracao(),
        contas=(
            *configuracao_itau().contas,
            ContaDeExtrato(padrao="pix_recebidos.json", banco="itau", conta="1234-5"),
            ContaDeExtrato(padrao="pix_recebidos_outra_conta.json", banco="itau", conta="9999-9"),
        ),
    )

    resultado = _conciliar(pasta, armazenamento, configuracao)

    assert resultado.lista_de_baixa == []
    assert all(TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA not in item.evidencia for item in resultado.fila_de_excecoes)


def test_credito_enriquecido_continua_identificado_sem_o_arquivo_de_pix_e_reimportar_nao_duplica(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    porta = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "1.000,00")])
    pix = _escrever_pix(pasta, [_pix(_txids(porta)["NF 1"], "1000.00", horario="2025-03-06T10:00:00Z")])
    _conciliar(pasta, armazenamento, _configuracao())

    reimportado = _conciliar(pasta, armazenamento, _configuracao())
    pix.unlink()
    sem_o_arquivo = _conciliar(pasta, armazenamento, _configuracao())

    for resultado in (reimportado, sem_o_arquivo):
        [item] = resultado.lista_de_baixa
        assert item.evidencia == (TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA,)
        assert resultado.fila_de_excecoes == []


def test_arquivo_de_pix_recebidos_e_identificado_pelo_conteudo_com_conta_pelo_nome(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    porta = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "1.000,00")])
    _escrever_pix(pasta, [_pix(_txids(porta)["NF 1"], "1000.00", horario="2025-03-06T10:00:00Z")], nome="pix_recebidos_marco.json")
    _escrever_pix(pasta, [_pix("CI000000000000000000000000000077", "5.00")], nome="sem_conta.json")
    sem_horario = _pix("CI000000000000000000000000000078", "5.00")
    del sem_horario["horario"]
    _escrever_pix(pasta, [sem_horario], nome="pix_recebidos_quebrado.json")
    campo_estranho = {**_pix("CI000000000000000000000000000079", "5.00"), "campoEstranho": 1}
    _escrever_pix(pasta, [campo_estranho], nome="pix_recebidos_estranho.json")

    resultado = _conciliar(pasta, armazenamento, _configuracao())

    assert "pix_recebidos_marco.json" in resultado.resumo.arquivos_lidos
    recusados = {recusado.nome: recusado.motivo for recusado in resultado.arquivos_recusados}
    assert sorted(recusados) == ["pix_recebidos_estranho.json", "pix_recebidos_quebrado.json", "sem_conta.json"]
    assert "conta" in recusados["sem_conta.json"]
    assert "horario" in recusados["pix_recebidos_quebrado.json"]
    assert "campoEstranho" in recusados["pix_recebidos_estranho.json"]
    [item] = resultado.lista_de_baixa
    assert item.evidencia == (TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA,)


def test_pix_recebido_sem_txid_nao_e_retorno_e_o_credito_segue_a_fase_1(pasta: Path, armazenamento: Armazenamento) -> None:
    _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "1.000,00")])
    _escrever_pix(pasta, [_pix(None, "1000.00", horario="2025-03-06T10:00:00Z")])

    resultado = _conciliar(pasta, armazenamento, _configuracao())

    assert resultado.arquivos_recusados == ()
    [item] = resultado.fila_de_excecoes
    assert TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA not in item.evidencia


def _escrever_configuracao(caminho: Path, migracao: bool = True) -> Path:
    texto = '[[contas]]\npadrao = "itau_cobranca*.csv"\nbanco = "itau"\nconta = "1234-5"\n'
    if migracao:
        texto += f'\n[migracao]\nfranqueados = ["{MIGRADO}"]\n'
    caminho.write_text(texto, encoding="utf-8")
    return caminho


def test_comando_de_emissao_emite_uma_vez_e_avisa_sem_configuracao_de_migracao(
    tmp_path: Path, pasta: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    estado = tmp_path / "estado.sqlite"
    armazenamento = Armazenamento(estado)
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    escrever_extrato_itau(pasta, [REFERENCIA])
    _conciliar(pasta, armazenamento, _configuracao())
    armazenamento.fechar()
    argumentos = ["--estado", str(estado), "--configuracao"]

    sem_migracao = main([*argumentos, str(_escrever_configuracao(tmp_path / "sem.toml", migracao=False))])
    aviso = capsys.readouterr().out
    primeira = main([*argumentos, str(_escrever_configuracao(tmp_path / "com.toml"))])
    emitidas = capsys.readouterr().out
    segunda = main([*argumentos, str(tmp_path / "com.toml")])
    nenhuma = capsys.readouterr().out

    assert (sem_migracao, primeira, segunda) == (0, 0, 0)
    assert "[migracao]" in aviso
    assert "1 Cobrança(s) identificada(s) emitida(s)" in emitidas
    assert "0 Cobrança(s) identificada(s) emitida(s)" in nenhuma


def test_configuracao_de_migracao_exige_raizes_de_cnpj(tmp_path: Path) -> None:
    caminho = tmp_path / "conciliacao.toml"
    caminho.write_text(
        '[[contas]]\npadrao = "x*.csv"\nbanco = "itau"\nconta = "1"\n\n[migracao]\nfranqueados = ["10.200.300"]\n',
        encoding="utf-8",
    )

    with pytest.raises(ConfiguracaoInvalida, match="raízes de CNPJ"):
        carregar_configuracao(caminho)
