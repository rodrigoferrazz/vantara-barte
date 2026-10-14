from dataclasses import replace
from decimal import Decimal
from fractions import Fraction
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.configuracao import Configuracao, ConfiguracaoInvalida, carregar_configuracao
from conciliacao.dominio import Resultado, TipoDeEvidencia
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import LinhaItau, TituloProtheus, configuracao_itau, escrever_extrato_itau, escrever_protheus
from tests.parecerista import MODELO_FALSO, VERSAO_DO_PROMPT_FALSA, PareceristaFalso

CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.300.400/0001-50"
FRANQUIA_A = "Franquia Teste Ltda"
FRANQUIA_B = "Padaria Outra Ltda"


def _executar(
    pasta: Path,
    armazenamento: Armazenamento,
    workbook: Path,
    configuracao: Configuracao | None = None,
    parecerista: PareceristaFalso | None = None,
) -> Resultado:
    resultado = executar_conciliacao(
        pasta,
        armazenamento,
        configuracao or configuracao_itau(),
        parecerista or PareceristaFalso(),
        workbook_anterior=workbook,
    )
    gerar_workbook(resultado, workbook)
    return resultado


def _resumo(workbook: Path) -> dict[object, object]:
    return {linha[0]: linha[1] for linha in load_workbook(workbook)["Resumo"].iter_rows(values_only=True)}


def _blocos(workbook: Path) -> dict[str, list[dict[str, object]]]:
    blocos: dict[str, list[dict[str, object]]] = {}
    linhas = list(load_workbook(workbook)["Métricas"].iter_rows(values_only=True))
    while linhas:
        titulo, cabecalho, *linhas = linhas
        corpo: list[dict[str, object]] = []
        while linhas and any(celula is not None for celula in linhas[0]):
            corpo.append(dict(zip([str(c) for c in cabecalho if c is not None], linhas.pop(0))))
        blocos[str(titulo[0])] = corpo
        while linhas and all(celula is None for celula in linhas[0]):
            linhas.pop(0)
    return blocos


def _decidir(workbook: Path, aba: str, decisoes: dict[str, dict[str, str | None]]) -> None:
    livro = load_workbook(workbook)
    planilha = livro[aba]
    cabecalho = [str(c.value) for c in planilha[1]]
    for linha in range(2, planilha.max_row + 1):
        item = str(planilha.cell(row=linha, column=cabecalho.index("Item") + 1).value)
        for coluna, valor in decisoes.get(item, {}).items():
            planilha.cell(row=linha, column=cabecalho.index(coluna) + 1).value = valor
    livro.properties.lastModifiedBy = "Ana Analista"
    livro.save(workbook)


def _por_nosso_numero(documento: str, nosso_numero: str, valor: str) -> TituloProtheus:
    return TituloProtheus(documento=documento, valor=valor, nosso_numero=nosso_numero)


def _liquidacao(data: str, valor: str, nosso_numero: str) -> LinhaItau:
    return LinhaItau(data, "LIQUIDACAO COBRANCA SIMPLES", valor, documento=nosso_numero)


def test_resumo_mostra_quanto_foi_resolvido_sozinho_em_quantidade_e_valor_e_os_conflitos(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        [
            _por_nosso_numero("NF 160001", "000016000011", "1.000,00"),
            _por_nosso_numero("NF 160002", "000016000022", "2.500,00"),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            _liquidacao("05/03/2025", "1.000,00", "000016000011"),
            _liquidacao("05/03/2025", "2.500,00", "000016000022"),
            LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "250,00"),
            LinhaItau("06/03/2025", "TARIFA BANCARIA", "2,50", tipo="D"),
        ],
    )
    workbook = tmp_path / "conciliacao.xlsx"

    resultado = _executar(pasta, armazenamento, workbook)

    automatica = resultado.resumo.parte_automatica
    assert automatica.percentual_da_quantidade == Fraction(200, 3)
    assert automatica.percentual_do_valor == Fraction(350000, 3750)
    resumo = _resumo(workbook)
    assert resumo["Resolvido sozinho"] == "2 de 3 Crédito(s) (66,7%); R$ 3.500,00 de R$ 3.750,00 (93,3%)"
    assert resumo["Alerta de queda da parte automática"] == "sem execução anterior para comparar"
    assert resumo["Conflitos entre Excel e tela"] == 0
    for linha in ("Arquivos lidos", "Arquivos recusados", "IA (parecer de nome)", "Conta que segura a Data de referência"):
        assert linha in resumo


def test_resumo_sem_credito_para_conciliar_nao_inventa_percentual(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_extrato_itau(pasta, [LinhaItau("06/03/2025", "TARIFA BANCARIA", "2,50", tipo="D")])
    workbook = tmp_path / "conciliacao.xlsx"

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.resumo.parte_automatica.percentual_da_quantidade is None
    assert _resumo(workbook)["Resolvido sozinho"] == "nenhum Crédito para conciliar nesta execução"


def _um_credito_resolvido_e_depois_um_sem_identificacao(pasta: Path) -> None:
    escrever_protheus(pasta, [_por_nosso_numero("NF 160101", "000016010011", "1.000,00")])
    escrever_extrato_itau(pasta, [_liquidacao("05/03/2025", "1.000,00", "000016010011")])


def _chega_credito_sem_identificacao(pasta: Path) -> None:
    escrever_extrato_itau(
        pasta, [LinhaItau("06/03/2025", "PIX RECEBIDO FULANO", "1.000,00")], nome="itau_cobranca_2.csv"
    )


def test_queda_da_parte_automatica_acima_do_limite_gera_alerta_no_resumo(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    configuracao = replace(configuracao_itau(), limite_de_queda=Decimal("49.9"))
    workbook = tmp_path / "conciliacao.xlsx"
    _um_credito_resolvido_e_depois_um_sem_identificacao(pasta)
    primeira = _executar(pasta, armazenamento, workbook, configuracao)
    _chega_credito_sem_identificacao(pasta)

    segunda = _executar(pasta, armazenamento, workbook, configuracao)

    assert primeira.resumo.alerta_de_queda is None
    assert segunda.resumo.alerta_de_queda == (
        f"ATENÇÃO: a parte resolvida sozinha caiu além do limite de 49,9 pontos percentuais em relação à "
        f"execução {primeira.resumo.execucao}: em quantidade de 100,0% para 50,0% (queda de 50,0); "
        "em valor de 100,0% para 50,0% (queda de 50,0). Verifique se alguma exportação veio quebrada."
    )
    assert _resumo(workbook)["Alerta de queda da parte automática"] == segunda.resumo.alerta_de_queda


def test_queda_igual_ao_limite_nao_gera_alerta(pasta: Path, armazenamento: Armazenamento, tmp_path: Path) -> None:
    configuracao = replace(configuracao_itau(), limite_de_queda=Decimal("50"))
    workbook = tmp_path / "conciliacao.xlsx"
    _um_credito_resolvido_e_depois_um_sem_identificacao(pasta)
    primeira = _executar(pasta, armazenamento, workbook, configuracao)
    _chega_credito_sem_identificacao(pasta)

    segunda = _executar(pasta, armazenamento, workbook, configuracao)

    assert segunda.resumo.alerta_de_queda is None
    assert _resumo(workbook)["Alerta de queda da parte automática"] == (
        f"sem alerta: na execução {primeira.resumo.execucao} eram 100,0% em quantidade e 100,0% em valor; "
        "limite de queda de 50 pontos percentuais"
    )


def test_queda_so_no_valor_tambem_gera_alerta(pasta: Path, armazenamento: Armazenamento, tmp_path: Path) -> None:
    configuracao = replace(configuracao_itau(), limite_de_queda=Decimal("10"))
    workbook = tmp_path / "conciliacao.xlsx"
    escrever_protheus(
        pasta,
        [
            _por_nosso_numero("NF 160201", "000016020011", "100,00"),
            _por_nosso_numero("NF 160202", "000016020022", "100,00"),
            _por_nosso_numero("NF 160203", "000016020033", "1,00"),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            _liquidacao("05/03/2025", "100,00", "000016020011"),
            _liquidacao("05/03/2025", "100,00", "000016020022"),
            LinhaItau("05/03/2025", "PIX RECEBIDO FULANO", "3,00"),
        ],
    )
    _executar(pasta, armazenamento, workbook, configuracao)
    escrever_extrato_itau(
        pasta,
        [_liquidacao("06/03/2025", "1,00", "000016020033"), LinhaItau("06/03/2025", "PIX RECEBIDO BELTRANO", "500,00")],
        nome="itau_cobranca_2.csv",
    )

    segunda = _executar(pasta, armazenamento, workbook, configuracao)

    assert segunda.resumo.parte_automatica.percentual_da_quantidade == 60
    assert segunda.resumo.alerta_de_queda is not None
    assert "em valor de 98,5% para 28,6% (queda de 70,0)" in segunda.resumo.alerta_de_queda
    assert "em quantidade" not in segunda.resumo.alerta_de_queda


def _configuracao_em_arquivo(tmp_path: Path, extra: str) -> Path:
    caminho = tmp_path / "conciliacao.toml"
    caminho.write_text(
        '[[contas]]\npadrao = "itau_cobranca*.csv"\nbanco = "itau"\nconta = "1234-5"\n' + extra, encoding="utf-8"
    )
    return caminho


def test_limite_de_queda_vem_da_configuracao(tmp_path: Path) -> None:
    caminho = _configuracao_em_arquivo(tmp_path, "[resumo]\nlimite_de_queda_pontos_percentuais = 12.5\n")

    assert carregar_configuracao(caminho).limite_de_queda == Decimal("12.5")


def test_sem_secao_de_resumo_vale_o_limite_padrao(tmp_path: Path) -> None:
    assert carregar_configuracao(_configuracao_em_arquivo(tmp_path, "")).limite_de_queda == Decimal("10")


@pytest.mark.parametrize(
    "secao",
    [
        '[resumo]\nlimite_de_queda_pontos_percentuais = "10%"\n',
        "[resumo]\nlimite_de_queda_pontos_percentuais = -1\n",
        "[resumo]\nlimite_de_queda_pontos_percentuais = 101\n",
        "[resumo]\nlimite = 10\n",
    ],
    ids=["texto", "negativo", "acima de 100", "campo desconhecido"],
)
def test_limite_de_queda_invalido_recusa_a_configuracao(tmp_path: Path, secao: str) -> None:
    with pytest.raises(ConfiguracaoInvalida, match="resumo"):
        carregar_configuracao(_configuracao_em_arquivo(tmp_path, secao))


def _tres_propostas_na_lista(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            _por_nosso_numero("NF 160301", "000016030011", "1.000,00"),
            _por_nosso_numero("NF 160302", "000016030022", "2.000,00"),
            TituloProtheus(documento="NF 160303", valor="700,00", cnpj=CNPJ_B, cliente=FRANQUIA_B),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            _liquidacao("05/03/2025", "1.000,00", "000016030011"),
            _liquidacao("05/03/2025", "2.000,00", "000016030022"),
            LinhaItau("05/03/2025", f"TED RECEBIDA CNPJ {CNPJ_B}", "700,00"),
        ],
    )


def test_relatorio_de_taxa_de_rejeicao_por_tipo_de_evidencia_e_por_regra(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _tres_propostas_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    por_documento = {item.titulos[0].documento: item.identificador for item in primeira.lista_de_baixa}
    _decidir(
        workbook,
        "Lista de baixa",
        {por_documento["NF 160302"]: {"Decisão": "rejeitar", "Motivo da rejeição": "valor de outro Título"}},
    )

    resultado = _executar(pasta, armazenamento, workbook)

    assert [(g.nome, g.confirmadas, g.rejeitadas, g.taxa_de_rejeicao) for g in resultado.metricas.por_evidencia] == [
        (TipoDeEvidencia.NOSSO_NUMERO.value, 1, 1, Fraction(1, 2)),
        (TipoDeEvidencia.CNPJ_CPF.value, 1, 0, Fraction(0)),
    ]
    blocos = _blocos(workbook)
    assert blocos["Taxa de rejeição por tipo de evidência"] == [
        {"Tipo de evidência": "nosso número", "Confirmadas": 1, "Rejeitadas": 1, "Taxa de rejeição": 0.5,
         "Revertidas depois de confirmadas": 0},
        {"Tipo de evidência": "CNPJ/CPF", "Confirmadas": 1, "Rejeitadas": 0, "Taxa de rejeição": 0,
         "Revertidas depois de confirmadas": 0},
    ]
    assert [linha["Regra"] for linha in blocos["Taxa de rejeição por regra"]] == [
        "nosso número + valor exato ou dentro do Encargo",
        "CNPJ/CPF + Valor único",
    ]
    assert blocos["Acerto da primeira sugestão da IA"] == []


def test_regra_com_muita_rejeicao_continua_propondo_para_a_lista(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _tres_propostas_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    _decidir(
        workbook,
        "Lista de baixa",
        {
            item.identificador: {"Decisão": "rejeitar", "Motivo da rejeição": "conferir"}
            for item in primeira.lista_de_baixa
        },
    )
    _executar(pasta, armazenamento, workbook)
    escrever_protheus(
        pasta,
        [
            _por_nosso_numero("NF 160301", "000016030011", "1.000,00"),
            _por_nosso_numero("NF 160302", "000016030022", "2.000,00"),
            TituloProtheus(documento="NF 160303", valor="700,00", cnpj=CNPJ_B, cliente=FRANQUIA_B),
            _por_nosso_numero("NF 160304", "000016030044", "3.000,00"),
            TituloProtheus(documento="NF 160305", valor="900,00", cnpj=CNPJ_A),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            _liquidacao("07/03/2025", "3.000,00", "000016030044"),
            LinhaItau("07/03/2025", f"TED RECEBIDA CNPJ {CNPJ_A}", "900,00"),
        ],
        nome="itau_cobranca_2.csv",
    )

    resultado = _executar(pasta, armazenamento, workbook)

    assert all(g.taxa_de_rejeicao == 1 for g in resultado.metricas.por_evidencia)
    assert sorted(t.documento for item in resultado.lista_de_baixa for t in item.titulos) == ["NF 160304", "NF 160305"]


def _fila_com_parecer(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 160401", valor="1.000,00", cliente=FRANQUIA_A, cnpj=CNPJ_A),
            TituloProtheus(documento="NF 160402", valor="2.000,00", cliente=FRANQUIA_B, cnpj=CNPJ_B),
            TituloProtheus(documento="NF 160403", valor="3.000,00", cliente=FRANQUIA_B, cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "PIX RECEBIDO FRQ TST", "1.000,00"),
            LinhaItau("05/03/2025", "PIX RECEBIDO XPTO COMERCIO", "2.000,00"),
            LinhaItau("05/03/2025", "PIX RECEBIDO PDR OUT", "3.000,00"),
        ],
    )


def test_relatorio_de_acerto_da_primeira_sugestao_da_ia(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _fila_com_parecer(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    parecerista = PareceristaFalso(sim_para=[("FRQ TST", FRANQUIA_A)])
    primeira = _executar(pasta, armazenamento, workbook, parecerista=parecerista)
    por_historico = {item.lancamento.historico: item.identificador for item in primeira.fila_de_excecoes}
    _decidir(
        workbook,
        "Fila de exceções",
        {
            por_historico["PIX RECEBIDO FRQ TST"]: {"Decisão": "confirmar"},
            por_historico["PIX RECEBIDO XPTO COMERCIO"]: {"Decisão": "confirmar"},
        },
    )

    resultado = _executar(pasta, armazenamento, workbook, parecerista=parecerista)

    [acerto] = resultado.metricas.acerto_da_ia
    assert (acerto.modelo, acerto.versao_do_prompt) == (MODELO_FALSO, VERSAO_DO_PROMPT_FALSA)
    assert (acerto.avaliados, acerto.acertos, acerto.aguardando_decisao) == (2, 1, 1)
    assert _blocos(workbook)["Acerto da primeira sugestão da IA"] == [
        {
            "Modelo": MODELO_FALSO,
            "Versão do prompt": VERSAO_DO_PROMPT_FALSA,
            "Avaliados": 2,
            "Acertos": 1,
            "Taxa de acerto": 0.5,
            "Aguardando decisão": 1,
        }
    ]


def test_acerto_da_ia_usa_o_primeiro_parecer_e_nao_o_ultimo(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _fila_com_parecer(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    _executar(pasta, armazenamento, workbook, parecerista=PareceristaFalso())
    segunda = _executar(pasta, armazenamento, workbook, parecerista=PareceristaFalso(sim_para=[("FRQ TST", FRANQUIA_A)]))
    [item] = [i for i in segunda.fila_de_excecoes if i.lancamento.historico == "PIX RECEBIDO FRQ TST"]
    _decidir(workbook, "Fila de exceções", {item.identificador: {"Decisão": "confirmar"}})

    resultado = _executar(pasta, armazenamento, workbook, parecerista=PareceristaFalso())

    [acerto] = resultado.metricas.acerto_da_ia
    assert (acerto.avaliados, acerto.acertos) == (1, 0)


def test_acerto_da_ia_ignora_execucao_em_que_ela_estava_indisponivel(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _fila_com_parecer(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    _executar(pasta, armazenamento, workbook, parecerista=PareceristaFalso(indisponivel=True))
    segunda = _executar(pasta, armazenamento, workbook, parecerista=PareceristaFalso(sim_para=[("FRQ TST", FRANQUIA_A)]))
    [item] = [i for i in segunda.fila_de_excecoes if i.lancamento.historico == "PIX RECEBIDO FRQ TST"]
    _decidir(workbook, "Fila de exceções", {item.identificador: {"Decisão": "confirmar"}})

    resultado = _executar(pasta, armazenamento, workbook)

    [acerto] = resultado.metricas.acerto_da_ia
    assert (acerto.avaliados, acerto.acertos, acerto.aguardando_decisao) == (1, 1, 2)
