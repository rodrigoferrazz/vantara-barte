from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.configuracao import Configuracao, ContaDeExtrato
from conciliacao.dominio import Resultado
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import (
    LinhaItau,
    TituloOmie,
    TituloProtheus,
    configuracao_itau,
    escrever_extrato_itau,
    escrever_omie,
    escrever_protheus,
)
from tests.parecerista import PareceristaFalso


def _duas_contas() -> Configuracao:
    return Configuracao(
        contas=(
            ContaDeExtrato(padrao="itau_cobranca*.csv", banco="itau", conta="1234-5"),
            ContaDeExtrato(padrao="itau_movimento*.csv", banco="itau", conta="9999-0"),
        )
    )


def _tarifa(data: str) -> LinhaItau:
    return LinhaItau(data, "TARIFA BANCARIA", "1,00", tipo="D")


def _executar(pasta: Path, armazenamento: Armazenamento, configuracao: Configuracao | None = None) -> Resultado:
    return executar_conciliacao(pasta, armazenamento, configuracao or configuracao_itau(), PareceristaFalso())


def _documentos_a_cobrar(resultado: Resultado) -> list[str]:
    return [t.titulo.documento for t in resultado.lista_de_cobranca]


def _documentos_em_analise(resultado: Resultado) -> list[str]:
    return [t.titulo.documento for t in resultado.em_analise]


def test_data_de_referencia_e_a_menor_das_ultimas_datas_entre_os_extratos_e_o_resumo_mostra_a_conta(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="700,00", vencimento="05/06/2025"),
            TituloProtheus(documento="NF 2", valor="800,00", vencimento="10/06/2025"),
        ],
    )
    escrever_extrato_itau(pasta, [_tarifa("12/06/2025")], nome="itau_cobranca.csv")
    escrever_extrato_itau(pasta, [_tarifa("02/06/2025"), _tarifa("09/06/2025")], nome="itau_movimento.csv")

    resultado = _executar(pasta, armazenamento, _duas_contas())

    referencia = resultado.resumo.data_de_referencia
    assert referencia.data == date(2025, 6, 9)
    assert [(c.banco, c.conta, c.ultima_data) for c in referencia.contas_que_seguram] == [
        ("itau", "9999-0", date(2025, 6, 9))
    ]
    assert _documentos_a_cobrar(resultado) == ["NF 1"]


def test_conta_configurada_sem_extrato_deixa_sem_data_de_referencia_e_sem_cobranca(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="700,00", vencimento="05/06/2025")])
    escrever_extrato_itau(pasta, [_tarifa("12/06/2025")], nome="itau_cobranca.csv")

    resultado = _executar(pasta, armazenamento, _duas_contas())

    referencia = resultado.resumo.data_de_referencia
    assert referencia.data is None
    assert [(c.conta, c.ultima_data) for c in referencia.contas_que_seguram] == [("9999-0", None)]
    assert resultado.lista_de_cobranca == ()
    assert resultado.em_analise == ()


def test_ultima_data_da_conta_considera_extratos_de_execucoes_anteriores(
    tmp_path: Path, armazenamento: Armazenamento
) -> None:
    primeira = tmp_path / "primeira"
    primeira.mkdir()
    escrever_extrato_itau(primeira, [_tarifa("09/06/2025")], nome="itau_movimento.csv")
    escrever_extrato_itau(primeira, [_tarifa("09/06/2025")], nome="itau_cobranca.csv")
    _executar(primeira, armazenamento, _duas_contas())
    segunda = tmp_path / "segunda"
    segunda.mkdir()
    escrever_extrato_itau(segunda, [_tarifa("11/06/2025")], nome="itau_cobranca.csv")

    resultado = _executar(segunda, armazenamento, _duas_contas())

    assert resultado.resumo.data_de_referencia.data == date(2025, 6, 9)


def test_titulo_proposto_na_lista_sai_da_cobranca_e_so_o_com_credito_candidato_na_fila_fica_em_analise(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF NA LISTA", valor="1.000,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF NA FILA", valor="500,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF SEM CREDITO", valor="730,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF A VENCER", valor="90,00", vencimento="20/06/2025"),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("02/06/2025", "PIX RECEBIDO 10200300000140", "1.000,00"),
            LinhaItau("03/06/2025", "DEPOSITO EM DINHEIRO", "500,00"),
            _tarifa("10/06/2025"),
        ],
    )

    resultado = _executar(pasta, armazenamento)

    assert [i.titulos[0].documento for i in resultado.lista_de_baixa] == ["NF NA LISTA"]
    assert [i.titulos[0].documento for i in resultado.fila_de_excecoes] == ["NF NA FILA"]
    assert _documentos_a_cobrar(resultado) == ["NF SEM CREDITO"]
    assert _documentos_em_analise(resultado) == ["NF NA FILA"]
    [na_fila] = [t for t in resultado.em_analise if t.titulo.documento == "NF NA FILA"]
    assert [c.historico for c in na_fila.creditos_candidatos] == ["DEPOSITO EM DINHEIRO"]


def test_titulos_citados_por_credito_que_fecha_com_mais_de_um_ficam_em_analise(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF A", valor="300,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF B", valor="300,00", vencimento="02/06/2025"),
        ],
    )
    escrever_extrato_itau(
        pasta, [LinhaItau("02/06/2025", "PIX RECEBIDO 10200300000140", "300,00"), _tarifa("10/06/2025")]
    )

    resultado = _executar(pasta, armazenamento)

    [ambiguo] = resultado.fila_de_excecoes
    assert ambiguo.titulos == ()
    assert resultado.lista_de_cobranca == ()
    assert sorted(_documentos_em_analise(resultado)) == ["NF A", "NF B"]


def test_titulos_do_franqueado_de_credito_sem_titulo_candidato_continuam_na_cobranca(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF X", valor="900,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF Y", valor="400,00", vencimento="02/06/2025"),
        ],
    )
    escrever_extrato_itau(
        pasta, [LinhaItau("03/06/2025", "PIX RECEBIDO 10200300000140", "650,00"), _tarifa("10/06/2025")]
    )

    resultado = _executar(pasta, armazenamento)

    assert [i.titulos for i in resultado.fila_de_excecoes] == [()]
    assert _documentos_a_cobrar(resultado) == ["NF X", "NF Y"]
    assert resultado.em_analise == ()


def test_titulo_unico_do_franqueado_com_valor_que_nao_fecha_fica_em_analise(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF X", valor="900,00", vencimento="02/06/2025")])
    escrever_extrato_itau(
        pasta, [LinhaItau("03/06/2025", "PIX RECEBIDO 10200300000140", "650,00"), _tarifa("10/06/2025")]
    )

    resultado = _executar(pasta, armazenamento)

    assert resultado.lista_de_cobranca == ()
    assert _documentos_em_analise(resultado) == ["NF X"]


def test_titulo_marcado_como_possivel_duplicidade_nao_entra_na_cobranca(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="800,00", vencimento="02/06/2025")])
    escrever_omie(pasta, [TituloOmie(documento="FAT-1", valor="800.00", vencimento="2025-06-02")])
    escrever_extrato_itau(pasta, [_tarifa("10/06/2025")])

    resultado = _executar(pasta, armazenamento)

    assert len(resultado.itens_de_titulo) == 1
    assert resultado.lista_de_cobranca == ()
    assert resultado.em_analise == ()


@pytest.mark.parametrize(
    ("nosso_numero", "vencimento", "ultima_data", "cobra"),
    [
        ("000077700001", "05/06/2025", "09/06/2025", False),
        ("000077700001", "05/06/2025", "10/06/2025", True),
        ("", "05/06/2025", "06/06/2025", True),
        ("", "07/06/2025", "09/06/2025", False),
        ("", "07/06/2025", "10/06/2025", True),
        ("000077700001", "18/06/2025", "23/06/2025", False),
        ("000077700001", "18/06/2025", "24/06/2025", True),
    ],
    ids=[
        "boleto no segundo dia util",
        "boleto no terceiro dia util",
        "sem boleto no dia seguinte",
        "sabado vence efetivamente na segunda",
        "sabado vencido na terca",
        "boleto com feriado na janela",
        "boleto depois da janela com feriado",
    ],
)
def test_vencido_pelo_vencimento_efetivo_e_boleto_respeita_dois_dias_uteis_de_compensacao(
    pasta: Path, armazenamento: Armazenamento, nosso_numero: str, vencimento: str, ultima_data: str, cobra: bool
) -> None:
    escrever_protheus(
        pasta, [TituloProtheus(documento="NF 1", valor="700,00", vencimento=vencimento, nosso_numero=nosso_numero)]
    )
    escrever_extrato_itau(pasta, [_tarifa(ultima_data)])

    resultado = _executar(pasta, armazenamento)

    assert _documentos_a_cobrar(resultado) == (["NF 1"] if cobra else [])


def test_cobranca_mostra_o_saldo_que_o_erp_passou_a_informar(tmp_path: Path, armazenamento: Armazenamento) -> None:
    primeira = tmp_path / "primeira"
    primeira.mkdir()
    escrever_protheus(primeira, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="02/06/2025")])
    escrever_extrato_itau(primeira, [_tarifa("05/06/2025")])
    _executar(primeira, armazenamento)
    segunda = tmp_path / "segunda"
    segunda.mkdir()
    escrever_protheus(segunda, [TituloProtheus(documento="NF 1", valor="600,00", vencimento="02/06/2025")])
    escrever_extrato_itau(segunda, [_tarifa("06/06/2025")])

    resultado = _executar(segunda, armazenamento)

    [titulo] = resultado.lista_de_cobranca
    assert titulo.saldo_centavos == 60000


def test_workbook_traz_a_lista_de_cobranca_com_em_analise_a_parte_e_o_resumo_com_a_data_de_referencia(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF NA FILA", valor="500,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF SEM CREDITO", valor="730,00", vencimento="02/06/2025"),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("03/06/2025", "DEPOSITO EM DINHEIRO", "500,00"), _tarifa("10/06/2025")])
    resultado = _executar(pasta, armazenamento)
    caminho = tmp_path / "saida.xlsx"

    gerar_workbook(resultado, caminho)

    workbook = load_workbook(caminho)
    assert workbook.sheetnames.index("Lista de cobrança") == workbook.sheetnames.index("Fila de exceções") + 1
    linhas = [linha for linha in workbook["Lista de cobrança"].iter_rows(values_only=True)]
    cabecalho = [str(c) for c in linhas[0]]
    a_cobrar = dict(zip(cabecalho, linhas[1]))
    assert a_cobrar["Título"] == "NF SEM CREDITO"
    assert a_cobrar["Saldo"] == 730
    assert all(c is None for c in linhas[2])
    assert linhas[3][0] == "Em análise"
    em_analise = dict(zip([str(c) for c in linhas[4]], linhas[5]))
    assert em_analise["Título"] == "NF NA FILA"
    [item] = resultado.fila_de_excecoes
    assert em_analise["Créditos candidatos"] == item.identificador
    resumo = {linha[0]: linha[1] for linha in workbook["Resumo"].iter_rows(values_only=True)}
    assert resumo["Data de referência"] == "10/06/2025"
    assert resumo["Conta que segura a Data de referência"] == "itau 1234-5 (último lançamento em 10/06/2025)"
    assert resumo["Lista de cobrança"] == "1 Título(s), R$ 730,00"
    assert resumo["Títulos em análise"] == "1 Título(s), R$ 500,00"
