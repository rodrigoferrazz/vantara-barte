from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from conciliacao.armazenamento import Armazenamento
from conciliacao.dominio import Destino, MotivoDeNaoIdentificacao, Resultado, TipoDeDecisao, TipoDeEvidencia
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import (
    LinhaItau,
    TituloOmie,
    TituloProtheus,
    configuracao_itau,
    escrever_extrato_itau,
    escrever_omie,
    escrever_protheus,
)
from tests.parecerista import MODELO_FALSO, VERSAO_DO_PROMPT_FALSA, PareceristaFalso

CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.300.400/0001-50"
FRANQUIA_A = "Franquia Teste Ltda"
FRANQUIA_B = "Padaria Outra Ltda"

DECISAO = "Decisão"
TITULOS_ESCOLHIDOS = "Títulos escolhidos"
FRANQUEADO = "Franqueado do Conciliado fora (CNPJ)"
REFERENCIA = "Referência do Conciliado fora"


def _executar(
    pasta: Path, armazenamento: Armazenamento, workbook: Path, parecerista: PareceristaFalso | None = None
) -> Resultado:
    resultado = executar_conciliacao(
        pasta, armazenamento, configuracao_itau(), parecerista or PareceristaFalso(), workbook_anterior=workbook
    )
    gerar_workbook(resultado, workbook)
    return resultado


def _linhas(aba: Worksheet) -> list[dict[str, object]]:
    cabecalho, *linhas = aba.iter_rows(values_only=True)
    return [dict(zip([str(c) for c in cabecalho], linha)) for linha in linhas]


def _resolver(workbook: Path, decisoes: dict[str, dict[str, str]]) -> None:
    livro = load_workbook(workbook)
    aba = livro["Fila de exceções"]
    cabecalho = [str(c.value) for c in aba[1]]
    for linha in range(2, aba.max_row + 1):
        item = str(aba.cell(row=linha, column=cabecalho.index("Item") + 1).value)
        for coluna, valor in decisoes.get(item, {}).items():
            aba.cell(row=linha, column=cabecalho.index(coluna) + 1).value = valor
    livro.properties.lastModifiedBy = "Ana Analista"
    livro.save(workbook)


def _titulos(*titulos: TituloProtheus) -> list[TituloProtheus]:
    return list(titulos)


def _extrato(pasta: Path, nome: str, *linhas: LinhaItau) -> None:
    escrever_extrato_itau(pasta, list(linhas), nome=nome)


def test_item_da_fila_resolvido_com_titulo_escolhido_segue_o_caminho_da_confirmacao(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        _titulos(
            TituloProtheus(documento="NF 130001", valor="1.000,00", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 130002", valor="1.000,00", cnpj=CNPJ_A),
        ),
    )
    _extrato(
        pasta,
        "itau_cobranca.csv",
        LinhaItau("05/03/2025", f"TED RECEBIDA CNPJ {CNPJ_A}", "1.000,00"),
        LinhaItau("20/03/2025", "TARIFA BANCARIA", "10,00", tipo="D"),
    )
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.MAIS_DE_UMA_COMBINACAO
    _resolver(workbook, {item.identificador: {DECISAO: "confirmar", TITULOS_ESCOLHIDOS: "NF 130002 (Protheus)"}})

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.fila_de_excecoes == []
    assert resultado.lista_de_baixa == []
    assert [(b.item, b.titulo.documento) for b in resultado.baixa_manual] == [(item.identificador, "NF 130002")]
    assert resultado.resumo.decisoes.confirmadas == 1
    [decisao] = armazenamento.decisoes()
    assert decisao.tipo is TipoDeDecisao.CONFIRMAR
    assert [t.documento for t in decisao.titulos] == ["NF 130002"]
    assert decisao.explicacao == item.explicacao
    assert decisao.evidencia == (TipoDeEvidencia.CNPJ_CPF,)
    assert [t.titulo.documento for t in resultado.cobranca] == ["NF 130001"]


def test_confirmar_na_fila_sem_titulo_escolhido_usa_os_sugeridos(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(pasta, _titulos(TituloProtheus(documento="NF 130101", valor="1.000,00", cliente=FRANQUIA_A)))
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "1.000,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {DECISAO: "Confirmar"}})

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.fila_de_excecoes == []
    assert [b.titulo.documento for b in resultado.baixa_manual] == ["NF 130101"]


def test_confirmacao_por_nome_cria_pagador_conhecido_e_o_proximo_credito_vai_para_a_lista(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        _titulos(
            TituloProtheus(documento="NF 130201", valor="1.000,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130202", valor="2.000,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130203", valor="500,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130204", valor="700,00", cliente=FRANQUIA_A),
        ),
    )
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "1.000,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    _resolver(workbook, {item.identificador: {DECISAO: "confirmar"}})
    _extrato(
        pasta,
        "itau_cobranca_2.csv",
        LinhaItau("06/03/2025", "TED RECEBIDA FRANQUIA TESTE", "2.000,00"),
        LinhaItau("06/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "1.200,00"),
    )

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.fila_de_excecoes == []
    por_valor = {i.lancamento.valor_centavos: i for i in resultado.lista_de_baixa}
    assert [t.documento for t in por_valor[200000].titulos] == ["NF 130202"]
    assert [t.documento for t in por_valor[120000].titulos] == ["NF 130203", "NF 130204"]
    for proposta in por_valor.values():
        assert proposta.evidencia == (TipoDeEvidencia.PAGADOR_CONHECIDO,)
        assert "Pagador conhecido" in proposta.explicacao
    assert por_valor[200000].regra is not None and por_valor[120000].regra is not None
    assert por_valor[200000].regra != por_valor[120000].regra


def test_pagador_conhecido_com_valor_que_nao_e_unico_vai_para_a_fila(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        _titulos(
            TituloProtheus(documento="NF 130301", valor="1.000,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130302", valor="2.000,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130303", valor="2.000,00", cliente=FRANQUIA_A),
        ),
    )
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "1.000,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {DECISAO: "confirmar"}})
    _extrato(pasta, "itau_cobranca_2.csv", LinhaItau("06/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "2.000,00"))

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.lista_de_baixa == []
    [novo] = resultado.fila_de_excecoes
    assert novo.motivo is MotivoDeNaoIdentificacao.MAIS_DE_UMA_COMBINACAO
    assert novo.evidencia == (TipoDeEvidencia.PAGADOR_CONHECIDO,)


def test_nome_confirmado_para_outro_franqueado_deixa_de_ser_pagador_conhecido(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        _titulos(
            TituloProtheus(documento="NF 130401", valor="1.000,00", cliente=FRANQUIA_A, cnpj=CNPJ_A),
            TituloProtheus(documento="NF 130402", valor="2.000,00", cliente=FRANQUIA_A, cnpj=CNPJ_A),
            TituloProtheus(documento="NF 130403", valor="3.000,00", cliente=FRANQUIA_B, cnpj=CNPJ_B),
        ),
    )
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "1.000,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [primeiro] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {primeiro.identificador: {DECISAO: "confirmar"}})
    _extrato(pasta, "itau_cobranca_2.csv", LinhaItau("06/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "3.000,00"))
    [segundo] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    assert segundo.evidencia == (TipoDeEvidencia.PAGADOR_CONHECIDO,)
    _resolver(workbook, {segundo.identificador: {DECISAO: "confirmar", TITULOS_ESCOLHIDOS: "NF 130403"}})
    _extrato(pasta, "itau_cobranca_3.csv", LinhaItau("07/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "2.000,00"))

    terceira = _executar(pasta, armazenamento, workbook)
    _extrato(pasta, "itau_cobranca_4.csv", LinhaItau("08/03/2025", "TED RECEBIDA FRANQUIA TESTE", "2.000,00"))
    quarta = _executar(pasta, armazenamento, workbook)

    assert [b.titulo.documento for b in terceira.baixa_manual] == ["NF 130403"]
    for resultado in (terceira, quarta):
        assert resultado.lista_de_baixa == []
        for item in resultado.fila_de_excecoes:
            assert item.motivo is MotivoDeNaoIdentificacao.NOME_AMBIGUO
            assert TipoDeEvidencia.PAGADOR_CONHECIDO not in item.evidencia
            assert "mais de um Franqueado" in item.explicacao
    assert len(quarta.fila_de_excecoes) == 2


def test_confirmacao_por_valor_ou_soma_nao_cria_pagador_conhecido(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        _titulos(
            TituloProtheus(documento="NF 130501", valor="1.000,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130502", valor="600,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130503", valor="900,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130504", valor="4.000,00", cliente=FRANQUIA_A),
        ),
    )
    _extrato(
        pasta,
        "itau_cobranca.csv",
        LinhaItau("05/03/2025", "PIX RECEBIDO JOAO DA SILVA", "1.000,00"),
        LinhaItau("05/03/2025", "PIX RECEBIDO MARIA SOUZA", "1.500,00"),
    )
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    por_valor = {i.lancamento.valor_centavos: i for i in primeira.fila_de_excecoes}
    assert por_valor[100000].evidencia == (TipoDeEvidencia.SO_VALOR,)
    assert [t.documento for t in por_valor[150000].titulos] == ["NF 130502", "NF 130503"]
    _resolver(workbook, {item.identificador: {DECISAO: "confirmar"} for item in primeira.fila_de_excecoes})
    _extrato(
        pasta,
        "itau_cobranca_2.csv",
        LinhaItau("06/03/2025", "PIX RECEBIDO JOAO DA SILVA", "4.000,00"),
        LinhaItau("06/03/2025", "PIX RECEBIDO MARIA SOUZA", "4.000,00"),
    )

    resultado = _executar(pasta, armazenamento, workbook)

    assert len(resultado.baixa_manual) == 3
    assert resultado.lista_de_baixa == []
    assert len(resultado.fila_de_excecoes) == 2
    assert all(TipoDeEvidencia.PAGADOR_CONHECIDO not in item.evidencia for item in resultado.fila_de_excecoes)


def test_decisao_registra_modelo_e_versao_do_prompt_e_parecer_confirmado_vira_pagador_conhecido(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        _titulos(
            TituloProtheus(documento="NF 130601", valor="1.000,00", cliente=FRANQUIA_A),
            TituloProtheus(documento="NF 130602", valor="2.000,00", cliente=FRANQUIA_A),
        ),
    )
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "PIX RECEBIDO FRQ TST", "1.000,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    parecerista = PareceristaFalso(sim_para=[("FRQ TST", FRANQUIA_A)])
    [item] = _executar(pasta, armazenamento, workbook, parecerista).fila_de_excecoes
    assert item.evidencia == (TipoDeEvidencia.PARECER_DE_LLM,)
    _resolver(workbook, {item.identificador: {DECISAO: "confirmar"}})
    _extrato(pasta, "itau_cobranca_2.csv", LinhaItau("06/03/2025", "PIX RECEBIDO FRQ TST", "2.000,00"))

    resultado = _executar(pasta, armazenamento, workbook, parecerista)

    [decisao] = armazenamento.decisoes()
    assert [(p.parecer.modelo, p.parecer.versao_do_prompt) for p in decisao.pareceres] == [
        (MODELO_FALSO, VERSAO_DO_PROMPT_FALSA)
    ]
    [proposta] = resultado.lista_de_baixa
    assert [t.documento for t in proposta.titulos] == ["NF 130602"]
    assert len(parecerista.chamadas) == 1


def _duplicidade(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        _titulos(
            TituloProtheus(documento="NF 130701", valor="800,00", cnpj=CNPJ_A, vencimento="05/03/2025"),
            TituloProtheus(documento="NF 130702", valor="900,00", cnpj=CNPJ_A, vencimento="05/03/2025"),
        ),
    )
    escrever_omie(pasta, [TituloOmie(documento="FAT-1307", valor="800.00", cnpj="10200300000140")])
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("20/03/2025", "TARIFA BANCARIA", "10,00", tipo="D"))


def test_mesma_divida_tira_o_outro_titulo_da_cobranca_para_sempre_e_lista_para_cancelamento(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _duplicidade(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    [duplicidade] = primeira.itens_de_titulo
    assert [t.titulo.documento for t in primeira.cobranca] == ["NF 130702"]
    _resolver(workbook, {duplicidade.identificador: {DECISAO: "mesma dívida", TITULOS_ESCOLHIDOS: "NF 130701 (Protheus)"}})

    segunda = _executar(pasta, armazenamento, workbook)
    terceira = _executar(pasta, armazenamento, workbook)

    for resultado in (segunda, terceira):
        assert resultado.itens_de_titulo == ()
        assert [t.titulo.documento for t in resultado.cobranca] == ["NF 130701", "NF 130702"]
        [cancelar] = resultado.cancelar_na_origem
        assert (cancelar.titulo.documento, cancelar.titulo.sistema_de_origem) == ("FAT-1307", "Omie")
        assert cancelar.oficial.documento == "NF 130701"
    aba = _linhas(load_workbook(workbook)["Cancelar na origem"])
    assert [(linha["Título"], linha["Sistema de origem"], linha["Título oficial"]) for linha in aba] == [
        ("FAT-1307", "Omie", "NF 130701 (Protheus)")
    ]
    assert segunda.resumo.decisoes.duplicidades == 1


def test_credito_que_fechava_com_os_dois_titulos_da_mesma_divida_casa_com_o_oficial(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _duplicidade(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    [duplicidade] = _executar(pasta, armazenamento, workbook).itens_de_titulo
    _resolver(workbook, {duplicidade.identificador: {DECISAO: "Mesma divida", TITULOS_ESCOLHIDOS: "FAT-1307"}})
    _extrato(pasta, "itau_cobranca_2.csv", LinhaItau("05/03/2025", f"TED RECEBIDA CNPJ {CNPJ_A}", "800,00"))

    resultado = _executar(pasta, armazenamento, workbook)

    [proposta] = resultado.lista_de_baixa
    assert [t.documento for t in proposta.titulos] == ["FAT-1307"]
    assert [c.titulo.documento for c in resultado.cancelar_na_origem] == ["NF 130701"]


def test_dividas_distintas_tira_o_item_da_fila_e_volta_os_titulos_para_a_cobranca(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _duplicidade(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    [duplicidade] = _executar(pasta, armazenamento, workbook).itens_de_titulo
    _resolver(workbook, {duplicidade.identificador: {DECISAO: "dívidas distintas"}})

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.itens_de_titulo == ()
    assert resultado.cancelar_na_origem == ()
    assert sorted(t.titulo.documento for t in resultado.cobranca) == ["FAT-1307", "NF 130701", "NF 130702"]


def test_mesma_divida_sem_titulo_oficial_do_grupo_nao_e_aplicada(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _duplicidade(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    [duplicidade] = _executar(pasta, armazenamento, workbook).itens_de_titulo
    _resolver(workbook, {duplicidade.identificador: {DECISAO: "mesma dívida", TITULOS_ESCOLHIDOS: "NF 130702"}})

    resultado = _executar(pasta, armazenamento, workbook)

    assert [i.identificador for i in resultado.itens_de_titulo] == [duplicidade.identificador]
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert duplicidade.identificador in aviso


def test_conciliado_fora_tira_o_credito_da_fila_e_mantem_a_contagem_na_causa_raiz(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(pasta, _titulos(TituloProtheus(documento="NF 130801", valor="1.000,00", vencimento="20/03/2025")))
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "777,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(
        workbook,
        {item.identificador: {DECISAO: "conciliado fora", FRANQUEADO: "30.400.500/0001-60", REFERENCIA: "planilha U3 linha 12"}},
    )

    segunda = _executar(pasta, armazenamento, workbook)
    terceira = _executar(pasta, armazenamento, workbook)

    for resultado in (segunda, terceira):
        assert resultado.fila_de_excecoes == []
        assert resultado.causa_raiz.quantidade == 1
        assert resultado.causa_raiz.valor_centavos == 77700
        por_motivo = {g.nome: g.quantidade for g in resultado.causa_raiz.por_motivo}
        assert por_motivo[MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR.value] == 1
        assert [g.nome for g in resultado.causa_raiz.por_franqueado] == ["raiz 30400500"]
        assert [c.identificador for c in resultado.conciliados_fora] == [item.identificador]
    assert segunda.resumo.decisoes.conciliados_fora == 1
    [decisao] = armazenamento.decisoes()
    assert decisao.tipo is TipoDeDecisao.CONCILIADO_FORA
    assert (decisao.franqueado, decisao.referencia) == ("30400500", "planilha U3 linha 12")


def test_conciliado_fora_sem_franqueado_ou_referencia_nao_e_aplicado(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _extrato(
        pasta,
        "itau_cobranca.csv",
        LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "777,00"),
        LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "888,00"),
    )
    workbook = tmp_path / "conciliacao.xlsx"
    [um, outro] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(
        workbook,
        {
            um.identificador: {DECISAO: "conciliado fora", REFERENCIA: "planilha U3"},
            outro.identificador: {DECISAO: "conciliado fora", FRANQUEADO: "30.400.500/0001-60"},
        },
    )

    resultado = _executar(pasta, armazenamento, workbook)

    assert len(resultado.fila_de_excecoes) == 2
    assert len(resultado.resumo.decisoes.nao_aplicadas) == 2
    assert armazenamento.decisoes() == ()


def test_credito_sem_titulo_classificado_como_devolucao_ou_erro_sai_da_fila(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        _titulos(
            TituloProtheus(documento="NF 130901", valor="1.000,00", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 130902", valor="1.000,00", cnpj=CNPJ_B, cliente=FRANQUIA_B),
        ),
    )
    workbook = tmp_path / "conciliacao.xlsx"
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "TARIFA", "1,00", tipo="D"))
    _executar(pasta, armazenamento, workbook)
    escrever_protheus(pasta, [])
    _extrato(
        pasta,
        "itau_cobranca_2.csv",
        LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ_A}", "300,00"),
        LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ_B}", "400,00"),
    )
    segunda = _executar(pasta, armazenamento, workbook)
    assert {i.motivo for i in segunda.fila_de_excecoes} == {MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO}
    [devolucao, erro] = segunda.fila_de_excecoes
    _resolver(workbook, {devolucao.identificador: {DECISAO: "devolução"}, erro.identificador: {DECISAO: "erro"}})

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.fila_de_excecoes == []
    assert resultado.resumo.decisoes.classificadas == 2
    assert [d.tipo for d in armazenamento.decisoes()] == [TipoDeDecisao.DEVOLUCAO, TipoDeDecisao.ERRO]


def test_antecipacao_continua_viva_e_casa_quando_o_titulo_aparece(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(pasta, _titulos(TituloProtheus(documento="NF 131001", valor="1.000,00", cnpj=CNPJ_A)))
    workbook = tmp_path / "conciliacao.xlsx"
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "TARIFA", "1,00", tipo="D"))
    _executar(pasta, armazenamento, workbook)
    escrever_protheus(pasta, [])
    _extrato(pasta, "itau_cobranca_2.csv", LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ_A}", "300,00"))
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {DECISAO: "antecipação"}})

    segunda = _executar(pasta, armazenamento, workbook)
    escrever_protheus(pasta, _titulos(TituloProtheus(documento="NF 131002", valor="300,00", cnpj=CNPJ_A)))
    terceira = _executar(pasta, armazenamento, workbook)

    [na_fila] = segunda.fila_de_excecoes
    assert na_fila.motivo is MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO
    assert "antecipação" in na_fila.explicacao
    assert segunda.resumo.decisoes.classificadas == 1
    [proposta] = terceira.lista_de_baixa
    assert [t.documento for t in proposta.titulos] == ["NF 131002"]


def test_classificacao_so_vale_para_credito_sem_titulo_em_aberto(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "777,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {DECISAO: "devolução"}})

    resultado = _executar(pasta, armazenamento, workbook)

    assert [i.identificador for i in resultado.fila_de_excecoes] == [item.identificador]
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert item.identificador in aviso and "Crédito sem título em aberto" in aviso


def test_titulo_escolhido_por_duas_decisoes_nao_e_aplicado_em_nenhuma(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(pasta, _titulos(TituloProtheus(documento="NF 131101", valor="1.000,00", vencimento="20/03/2025")))
    _extrato(
        pasta,
        "itau_cobranca.csv",
        LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "1.000,00"),
        LinhaItau("06/03/2025", "DEPOSITO NAO IDENTIFICADO", "1.000,00"),
    )
    workbook = tmp_path / "conciliacao.xlsx"
    fila = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {DECISAO: "confirmar", TITULOS_ESCOLHIDOS: "NF 131101"} for item in fila})

    resultado = _executar(pasta, armazenamento, workbook)

    assert len(resultado.fila_de_excecoes) == 2
    assert armazenamento.decisoes() == ()
    assert len(resultado.resumo.decisoes.nao_aplicadas) == 2


def test_titulo_escolhido_que_nao_esta_em_aberto_nao_e_aplicado(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "1.000,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {DECISAO: "confirmar", TITULOS_ESCOLHIDOS: "NF 999999"}})

    resultado = _executar(pasta, armazenamento, workbook)

    assert [i.identificador for i in resultado.fila_de_excecoes] == [item.identificador]
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert "NF 999999" in aviso


def test_fila_sai_no_workbook_com_colunas_de_decisao_vazias(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "1.000,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    _executar(pasta, armazenamento, workbook)

    [linha] = _linhas(load_workbook(workbook)["Fila de exceções"])

    assert [linha[coluna] for coluna in (DECISAO, TITULOS_ESCOLHIDOS, FRANQUEADO, REFERENCIA)] == [None] * 4
    assert Destino.FILA_DE_EXCECOES.value in load_workbook(workbook).sheetnames


def test_credito_classificado_como_possivel_fonte_nao_integrada_continua_vivo_na_fila_com_esse_motivo(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "777,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR
    _resolver(workbook, {item.identificador: {DECISAO: "Possivel fonte nao integrada"}})

    segunda = _executar(pasta, armazenamento, workbook)
    [fila_resolvida] = [
        str(linha[1])
        for linha in load_workbook(workbook)["Resumo"].iter_rows(values_only=True)
        if linha[0] == "Fila resolvida pelo Excel"
    ]
    terceira = _executar(pasta, armazenamento, workbook)

    assert "1 classificado(s) como Possível fonte não integrada" in fila_resolvida
    for resultado in (segunda, terceira):
        [na_fila] = resultado.fila_de_excecoes
        assert na_fila.lancamento == item.lancamento
        assert na_fila.motivo is MotivoDeNaoIdentificacao.POSSIVEL_FONTE_NAO_INTEGRADA
        assert "Classificado pela Analista como Possível fonte não integrada" in na_fila.explicacao
        por_motivo = {g.nome: (g.quantidade, g.valor_centavos) for g in resultado.causa_raiz.por_motivo}
        assert por_motivo[MotivoDeNaoIdentificacao.POSSIVEL_FONTE_NAO_INTEGRADA.value] == (1, 77700)
        assert por_motivo[MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR.value] == (0, 0)
    assert segunda.resumo.decisoes.fontes_nao_integradas == 1
    assert terceira.resumo.decisoes.fontes_nao_integradas == 0
    [decisao] = armazenamento.decisoes()
    assert (decisao.tipo, decisao.origem) == (TipoDeDecisao.POSSIVEL_FONTE_NAO_INTEGRADA, "Excel")


def test_possivel_fonte_nao_integrada_marcada_conciliado_fora_conta_na_causa_raiz_com_esse_motivo(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "777,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {DECISAO: "possível fonte não integrada"}})
    [classificado] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(
        workbook,
        {classificado.identificador: {DECISAO: "conciliado fora", FRANQUEADO: "30400500", REFERENCIA: "planilha U3"}},
    )

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.fila_de_excecoes == []
    por_motivo = {g.nome: g.quantidade for g in resultado.causa_raiz.por_motivo}
    assert por_motivo[MotivoDeNaoIdentificacao.POSSIVEL_FONTE_NAO_INTEGRADA.value] == 1
    assert [g.nome for g in resultado.causa_raiz.por_franqueado] == ["raiz 30400500"]


def test_credito_classificado_como_possivel_fonte_nao_integrada_casa_quando_o_titulo_aparece(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(pasta, _titulos(TituloProtheus(documento="NF 131201", valor="1.000,00", cnpj=CNPJ_A)))
    workbook = tmp_path / "conciliacao.xlsx"
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("05/03/2025", "TARIFA", "1,00", tipo="D"))
    _executar(pasta, armazenamento, workbook)
    escrever_protheus(pasta, [])
    _extrato(pasta, "itau_cobranca_2.csv", LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ_A}", "300,00"))
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {DECISAO: "possível fonte não integrada"}})
    _executar(pasta, armazenamento, workbook)
    escrever_protheus(pasta, _titulos(TituloProtheus(documento="NF 131202", valor="300,00", cnpj=CNPJ_A)))

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.fila_de_excecoes == []
    [proposta] = resultado.lista_de_baixa
    assert [t.documento for t in proposta.titulos] == ["NF 131202"]


def test_classificacao_de_possivel_fonte_nao_integrada_nao_vale_para_a_lista_de_baixa(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(pasta, _titulos(TituloProtheus(documento="NF 131301", valor="300,00", cnpj=CNPJ_A)))
    _extrato(pasta, "itau_cobranca.csv", LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ_A}", "300,00"))
    workbook = tmp_path / "conciliacao.xlsx"
    [proposta] = _executar(pasta, armazenamento, workbook).lista_de_baixa
    livro = load_workbook(workbook)
    aba = livro["Lista de baixa"]
    cabecalho = [str(c.value) for c in aba[1]]
    aba.cell(row=2, column=cabecalho.index(DECISAO) + 1).value = "possível fonte não integrada"
    livro.properties.lastModifiedBy = "Ana Analista"
    livro.save(workbook)

    resultado = _executar(pasta, armazenamento, workbook)

    assert [i.identificador for i in resultado.lista_de_baixa] == [proposta.identificador]
    assert armazenamento.decisoes() == ()
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert proposta.identificador in aviso
