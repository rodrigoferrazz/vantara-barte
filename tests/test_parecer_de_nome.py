from pathlib import Path

from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.dominio import Item, MotivoDeNaoIdentificacao, Resposta, Resultado, TipoDeEvidencia
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import (
    LinhaItau,
    TituloOmie,
    TituloProtheus,
    configuracao_itau,
    escrever_extrato_itau,
    escrever_omie,
    escrever_protheus,
)
from tests.parecerista import MODELO_FALSO, VERSAO_DO_PROMPT_FALSA, PareceristaFalso

CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.400.600/0001-80"
NOME_A = "Confeitaria Alvorada Comercio de Alimentos Ltda"
NOME_B = "Mercearia Horizonte Ltda"
PIX_ABREVIADO = "PIX RECEBIDO CONF ALVORADA COM"


def _executar(pasta: Path, armazenamento: Armazenamento, parecerista: PareceristaFalso) -> Resultado:
    return executar_conciliacao(pasta, armazenamento, configuracao_itau(), parecerista)


def _unico_da_fila(resultado: Resultado) -> Item:
    assert resultado.lista_de_baixa == []
    [item] = resultado.fila_de_excecoes
    return item


def _dois_franqueados(pasta: Path, valor_a: str = "3.200,00", valor_b: str = "1.000,00") -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor=valor_a, cliente=NOME_A, cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor=valor_b, cliente=NOME_B, cnpj=CNPJ_B),
        ],
    )


def test_parecer_positivo_vai_para_a_fila_como_pagador_provavel_com_o_parecer_e_nunca_para_a_lista(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    _dois_franqueados(pasta)
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", PIX_ABREVIADO, "3.200,00")])
    parecerista = PareceristaFalso(sim_para=[("CONF ALVORADA COM", NOME_A)])

    item = _unico_da_fila(_executar(pasta, armazenamento, parecerista))

    assert item.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert [t.documento for t in item.titulos] == ["NF 1"]
    assert item.evidencia == (TipoDeEvidencia.PARECER_DE_LLM,)
    [parecer] = item.pareceres
    assert parecer.franqueado == "10200300"
    assert parecer.parecer.resposta is Resposta.SIM
    assert "abrevia" in parecer.parecer.explicacao
    assert parecer.parecer.modelo == MODELO_FALSO
    assert parecer.parecer.versao_do_prompt == VERSAO_DO_PROMPT_FALSA
    assert "abrevia" in item.explicacao
    assert not item.ia_indisponivel


def test_a_ia_so_ve_os_nomes_do_franqueado_com_que_o_valor_fecha(pasta: Path, armazenamento: Armazenamento) -> None:
    _dois_franqueados(pasta)
    escrever_omie(pasta, [TituloOmie(documento="OM 1", valor="50.00", cliente="CONF ALVORADA NORTE", cnpj="10200300000140")])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", PIX_ABREVIADO, "3.200,00")])
    parecerista = PareceristaFalso()

    _executar(pasta, armazenamento, parecerista)

    assert parecerista.chamadas == [("CONF ALVORADA COM", ("CONF ALVORADA NORTE", NOME_A))]


def test_nome_por_prefixo_nao_chama_a_ia(pasta: Path, armazenamento: Armazenamento) -> None:
    _dois_franqueados(pasta)
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO CONFEITARIA ALVORADA CO", "3.200,00")])
    parecerista = PareceristaFalso(sim_para=[("CONFEITARIA ALVORADA CO", NOME_A)])

    item = _unico_da_fila(_executar(pasta, armazenamento, parecerista))

    assert parecerista.chamadas == []
    assert item.pareceres == ()
    assert item.evidencia == (TipoDeEvidencia.NOME_POR_PREFIXO,)


def test_nome_sem_franqueado_com_que_o_valor_fecha_nao_chama_a_ia(pasta: Path, armazenamento: Armazenamento) -> None:
    _dois_franqueados(pasta)
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", PIX_ABREVIADO, "777,00")])
    parecerista = PareceristaFalso(sim_para=[("CONF ALVORADA COM", NOME_A)])

    item = _unico_da_fila(_executar(pasta, armazenamento, parecerista))

    assert parecerista.chamadas == []
    assert item.motivo is MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR
    assert item.titulos == ()


def test_credito_sem_nome_nao_chama_a_ia(pasta: Path, armazenamento: Armazenamento) -> None:
    _dois_franqueados(pasta)
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "3.200,00")])
    parecerista = PareceristaFalso()

    _executar(pasta, armazenamento, parecerista)

    assert parecerista.chamadas == []


def test_parecer_negativo_vira_sem_identificacao_do_pagador_mantendo_a_sugestao_por_valor(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    _dois_franqueados(pasta)
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", PIX_ABREVIADO, "3.200,00")])

    item = _unico_da_fila(_executar(pasta, armazenamento, PareceristaFalso()))

    assert item.motivo is MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR
    assert [t.documento for t in item.titulos] == ["NF 1"]
    assert item.evidencia == (TipoDeEvidencia.SO_VALOR,)
    [parecer] = item.pareceres
    assert parecer.parecer.resposta is Resposta.NAO
    assert "não lembra" in item.explicacao


def test_parecer_positivo_com_soma_de_titulos_do_franqueado_sugere_a_combinacao(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="1.200,00", cliente=NOME_A, cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="800,00", cliente=NOME_A, cnpj=CNPJ_A),
            TituloProtheus(documento="NF 3", valor="450,00", cliente=NOME_A, cnpj=CNPJ_A),
            TituloProtheus(documento="NF 4", valor="9.999,00", cliente=NOME_B, cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", PIX_ABREVIADO, "2.000,00")])
    parecerista = PareceristaFalso(sim_para=[("CONF ALVORADA COM", NOME_A)])

    item = _unico_da_fila(_executar(pasta, armazenamento, parecerista))

    assert item.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert sorted(t.documento for t in item.titulos) == ["NF 1", "NF 2"]
    assert len(parecerista.chamadas) == 1


def test_parecer_positivo_com_mais_de_uma_combinacao_do_franqueado_nao_sugere(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="500,00", cliente=NOME_A, cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="500,00", cliente=NOME_A, cnpj=CNPJ_A),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", PIX_ABREVIADO, "500,00")])
    parecerista = PareceristaFalso(sim_para=[("CONF ALVORADA COM", NOME_A)])

    item = _unico_da_fila(_executar(pasta, armazenamento, parecerista))

    assert item.motivo is MotivoDeNaoIdentificacao.MAIS_DE_UMA_COMBINACAO
    assert item.titulos == ()


def test_parecer_positivo_para_mais_de_um_franqueado_vira_nome_ambiguo_sem_sugestao(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    _dois_franqueados(pasta, valor_a="3.200,00", valor_b="3.200,00")
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", PIX_ABREVIADO, "3.200,00")])
    parecerista = PareceristaFalso(sim_para=[("CONF ALVORADA COM", NOME_A), ("CONF ALVORADA COM", NOME_B)])

    item = _unico_da_fila(_executar(pasta, armazenamento, parecerista))

    assert item.motivo is MotivoDeNaoIdentificacao.NOME_AMBIGUO
    assert item.titulos == ()
    assert len(item.pareceres) == 2


def test_ia_indisponivel_nao_para_a_execucao_e_marca_o_item_e_o_resumo(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_franqueados(pasta)
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", PIX_ABREVIADO, "3.200,00"),
            LinhaItau("05/03/2025", "PIX RECEBIDO MERC HORIZ", "1.000,00"),
            LinhaItau("05/03/2025", "PIX RECEBIDO FULANO", "12,34"),
        ],
    )

    resultado = _executar(pasta, armazenamento, PareceristaFalso(indisponivel=True))

    assert len(resultado.fila_de_excecoes) == 3
    marcados = [item for item in resultado.fila_de_excecoes if item.ia_indisponivel]
    assert [item.lancamento.historico for item in marcados] == [PIX_ABREVIADO, "PIX RECEBIDO MERC HORIZ"]
    assert all(item.motivo is MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR for item in marcados)
    assert [t.documento for t in marcados[0].titulos] == ["NF 1"]
    assert all("IA indisponível" in item.explicacao for item in marcados)
    assert resultado.resumo.itens_ia_indisponivel == 2
    caminho = tmp_path / "saida.xlsx"
    gerar_workbook(resultado, caminho)
    resumo = {linha[0]: linha[1] for linha in load_workbook(caminho)["Resumo"].iter_rows(values_only=True)}
    assert "indisponível" in str(resumo["IA (parecer de nome)"])
    assert "2" in str(resumo["IA (parecer de nome)"])


def test_parecer_e_exibido_no_item_da_fila_do_workbook_com_modelo_e_versao_do_prompt(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_franqueados(pasta)
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", PIX_ABREVIADO, "3.200,00")])
    resultado = _executar(pasta, armazenamento, PareceristaFalso(sim_para=[("CONF ALVORADA COM", NOME_A)]))
    caminho = tmp_path / "saida.xlsx"

    gerar_workbook(resultado, caminho)

    workbook = load_workbook(caminho)
    linhas = list(workbook["Fila de exceções"].iter_rows(values_only=True))
    [linha] = [dict(zip(linhas[0], valores)) for valores in linhas[1:]]
    parecer = str(linha["Parecer da IA"])
    assert "sim" in parecer
    assert "abrevia" in parecer
    assert MODELO_FALSO in parecer
    assert VERSAO_DO_PROMPT_FALSA in parecer
    assert linha["Evidência"] == TipoDeEvidencia.PARECER_DE_LLM.value
    resumo = {valores[0]: valores[1] for valores in workbook["Resumo"].iter_rows(values_only=True)}
    assert "funcionando" in str(resumo["IA (parecer de nome)"])
