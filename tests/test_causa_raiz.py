from collections.abc import Sequence
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.configuracao import Configuracao, ContaDeExtrato
from conciliacao.dominio import Canal, CausaRaiz, GrupoDaCausaRaiz, MotivoDeNaoIdentificacao, Resultado
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import LinhaItau, TituloProtheus, configuracao_itau, escrever_extrato_itau, escrever_protheus
from tests.parecerista import PareceristaFalso

CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.400.600/0001-80"
CNPJ_C = "30.600.900/0001-10"


def _executar(pasta: Path, armazenamento: Armazenamento, configuracao: Configuracao | None = None) -> Resultado:
    return executar_conciliacao(pasta, armazenamento, configuracao or configuracao_itau(), PareceristaFalso())


def _por_nome(grupos: tuple[GrupoDaCausaRaiz, ...]) -> dict[str, tuple[int, int]]:
    return {grupo.nome: (grupo.quantidade, grupo.valor_centavos) for grupo in grupos}


def _cenario_com_varios_motivos(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF A1", valor="1.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF A2", valor="2.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF B1", valor="500,00", cliente="Brisa Serena Armazem Ltda", cnpj=CNPJ_B),
            TituloProtheus(
                documento="NF C1", valor="800,00", cliente="Recanto Feliz Ltda", cnpj=CNPJ_C, nosso_numero="000077700001"
            ),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "PIX RECEBIDO AURORA PANIF", "1.000,00"),
            LinhaItau("05/03/2025", "TED RECEBIDA CNPJ 10200300000140", "1.234,00"),
            LinhaItau("06/03/2025", "PIX RECEBIDO BRISA SERENA ARM", "500,00"),
            LinhaItau("06/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "750,00", documento="000077700001"),
            LinhaItau("07/03/2025", "DEPOSITO EM DINHEIRO", "333,00"),
            LinhaItau("07/03/2025", "PIX RECEBIDO 30600900000110", "800,00"),
            LinhaItau("07/03/2025", "TARIFA BANCARIA", "2,50", tipo="D"),
        ],
    )


def test_cada_credito_da_fila_conta_em_um_so_motivo_e_todos_os_motivos_aparecem(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    _cenario_com_varios_motivos(pasta)

    resultado = _executar(pasta, armazenamento)

    causa_raiz = resultado.causa_raiz
    assert [grupo.nome for grupo in causa_raiz.por_motivo] == [motivo.value for motivo in MotivoDeNaoIdentificacao]
    assert _por_nome(causa_raiz.por_motivo) == {
        "sem identificação do pagador": (1, 33300),
        "nome ambíguo": (0, 0),
        "valor não fecha": (1, 75000),
        "Crédito sem título em aberto": (1, 123400),
        "mais de uma combinação possível": (0, 0),
        "Possível duplicidade entre sistemas": (0, 0),
        "Possível fonte não integrada": (0, 0),
        "pagador provável, primeira vez": (2, 150000),
    }
    fila = resultado.fila_de_excecoes
    assert (causa_raiz.quantidade, causa_raiz.valor_centavos) == (
        len(fila),
        sum(item.lancamento.valor_centavos for item in fila),
    )


def test_todas_as_agregacoes_somam_100_por_cento_da_fila(pasta: Path, armazenamento: Armazenamento) -> None:
    _cenario_com_varios_motivos(pasta)

    causa_raiz = _executar(pasta, armazenamento).causa_raiz

    for agregacao in (causa_raiz.por_motivo, causa_raiz.por_franqueado, causa_raiz.por_banco, causa_raiz.por_canal):
        assert sum(grupo.quantidade for grupo in agregacao) == causa_raiz.quantidade == 5
        assert sum(grupo.valor_centavos for grupo in agregacao) == causa_raiz.valor_centavos == 381700
        assert sum(causa_raiz.fracao_da_quantidade(grupo) or 0 for grupo in agregacao) == 1


def test_agregacao_por_franqueado_ordena_por_concentracao_e_poe_sem_franqueado_a_parte(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    _cenario_com_varios_motivos(pasta)

    causa_raiz = _executar(pasta, armazenamento).causa_raiz

    assert [(g.nome, g.quantidade, g.valor_centavos) for g in causa_raiz.por_franqueado] == [
        ("Aurora Panificadora Ltda (raiz 10200300)", 2, 223400),
        ("Recanto Feliz Ltda (raiz 30600900)", 1, 75000),
        ("Brisa Serena Armazem Ltda (raiz 20400600)", 1, 50000),
        (CausaRaiz.SEM_FRANQUEADO, 1, 33300),
    ]


def test_credito_cujo_nome_aponta_mais_de_um_franqueado_fica_sem_franqueado(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF A1", valor="700,00", cliente="Quitanda Estrela Leste", cnpj=CNPJ_A),
            TituloProtheus(documento="NF B1", valor="700,00", cliente="Quitanda Estrela Oeste", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO QUITANDA ESTRELA", "700,00")])

    resultado = _executar(pasta, armazenamento)

    [item] = resultado.fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.NOME_AMBIGUO
    assert _por_nome(resultado.causa_raiz.por_franqueado) == {CausaRaiz.SEM_FRANQUEADO: (1, 70000)}


def test_valor_desempata_o_nome_entre_franqueados_e_atribui_o_credito_ao_franqueado(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF A1", valor="700,00", cliente="Quitanda Estrela Leste", cnpj=CNPJ_A),
            TituloProtheus(documento="NF B1", valor="900,00", cliente="Quitanda Estrela Oeste", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO QUITANDA ESTRELA", "900,00")])

    resultado = _executar(pasta, armazenamento)

    [item] = resultado.fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert _por_nome(resultado.causa_raiz.por_franqueado) == {"Quitanda Estrela Oeste (raiz 20400600)": (1, 90000)}


def test_franqueado_pelo_parecer_positivo_da_ia(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF A1", valor="700,00", cliente="Aroma e Companhia Ltda")])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO AROMA CIA", "700,00")])
    parecerista = PareceristaFalso(sim_para=[("AROMA CIA", "Aroma e Companhia Ltda")])

    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), parecerista)

    assert _por_nome(resultado.causa_raiz.por_franqueado) == {"Aroma e Companhia Ltda (raiz 10200300)": (1, 70000)}


def test_canal_inferido_do_historico_e_todos_os_canais_aparecem(pasta: Path, armazenamento: Armazenamento) -> None:
    _cenario_com_varios_motivos(pasta)
    escrever_extrato_itau(
        pasta, [LinhaItau("08/03/2025", "CREDITO DIVERSO", "10,00")], nome="itau_cobranca_2.csv"
    )

    causa_raiz = _executar(pasta, armazenamento).causa_raiz

    assert [grupo.nome for grupo in causa_raiz.por_canal] == [canal.value for canal in Canal]
    assert _por_nome(causa_raiz.por_canal) == {
        "Pix": (2, 150000),
        "TED": (1, 123400),
        "boleto": (1, 75000),
        "depósito": (1, 33300),
        "canal não identificado": (1, 1000),
    }


def test_agregacao_por_banco(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "DEPOSITO EM DINHEIRO", "100,00")])
    escrever_extrato_itau(
        pasta,
        [LinhaItau("05/03/2025", "DEPOSITO EM DINHEIRO", "40,00"), LinhaItau("05/03/2025", "DEPOSITO", "60,00")],
        nome="itau_outro.csv",
    )
    configuracao = Configuracao(
        contas=(
            ContaDeExtrato(padrao="itau_cobranca*.csv", banco="itau", conta="1234-5"),
            ContaDeExtrato(padrao="itau_outro*.csv", banco="itau", conta="9999-0"),
        )
    )

    causa_raiz = _executar(pasta, armazenamento, configuracao).causa_raiz

    assert _por_nome(causa_raiz.por_banco) == {"itau": (3, 20000)}


def test_fila_vazia_tem_causa_raiz_zerada_sem_dividir_por_zero(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "TARIFA BANCARIA", "2,50", tipo="D")])

    causa_raiz = _executar(pasta, armazenamento).causa_raiz

    assert causa_raiz.quantidade == 0
    assert causa_raiz.por_franqueado == ()
    assert all(causa_raiz.fracao_do_valor(grupo) is None for grupo in causa_raiz.por_motivo)


def test_workbook_traz_a_aba_causa_raiz_com_as_quatro_agregacoes(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _cenario_com_varios_motivos(pasta)
    resultado = _executar(pasta, armazenamento)
    caminho = tmp_path / "saida.xlsx"

    gerar_workbook(resultado, caminho)

    workbook = load_workbook(caminho)
    assert workbook.sheetnames.index("Causa raiz") == workbook.sheetnames.index("Lista de cobrança") + 1
    linhas = list(workbook["Causa raiz"].iter_rows(values_only=True))
    blocos = _blocos(linhas)
    assert list(blocos) == ["Por Motivo de não identificação", "Por Franqueado", "Por banco", "Por canal"]
    for bloco in blocos.values():
        assert bloco["Total"] == (5, 1, 3817, 1)
    motivos = blocos["Por Motivo de não identificação"]
    assert motivos["valor não fecha"] == pytest.approx((1, 1 / 5, 750, 750 / 3817))
    assert motivos["Possível fonte não integrada"] == (0, 0, 0, 0)
    assert list(blocos["Por Franqueado"])[0] == "Aurora Panificadora Ltda (raiz 10200300)"
    assert blocos["Por Franqueado"][CausaRaiz.SEM_FRANQUEADO] == (1, 1 / 5, 333, 333 / 3817)


def _blocos(linhas: Sequence[tuple[object, ...]]) -> dict[str, dict[str, tuple[object, ...]]]:
    blocos: dict[str, dict[str, tuple[object, ...]]] = {}
    atual: dict[str, tuple[object, ...]] = {}
    for linha in linhas:
        if not any(celula is not None for celula in linha):
            continue
        primeira = str(linha[0])
        if primeira.startswith("Por "):
            atual = blocos.setdefault(primeira, {})
            continue
        if primeira == "Grupo":
            assert list(linha[:5]) == ["Grupo", "Quantidade", "% quantidade", "Valor", "% valor"]
            continue
        atual[primeira] = tuple(float(str(c)) if c is not None else None for c in linha[1:5])
    return blocos
