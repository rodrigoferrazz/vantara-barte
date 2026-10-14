from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.configuracao import (
    Configuracao,
    ConfiguracaoDaAnaliseDeImpacto,
    ConfiguracaoInvalida,
    carregar_configuracao,
)
from conciliacao.dominio import (
    CausaRaiz,
    Destino,
    Item,
    Lancamento,
    MotivoDeNaoIdentificacao,
    Natureza,
    RankingDeMigracao,
)
from conciliacao.encargo import ParametrosDeEncargo
from conciliacao.execucao import executar_conciliacao
from conciliacao.gerador import ConfiguracaoDoGerador, gerar_dados_sinteticos
from conciliacao.impacto import analisar_impacto, gerar_workbook_do_impacto, main
from tests.arquivos import LinhaItau, TituloProtheus, configuracao_itau, escrever_extrato_itau, escrever_protheus
from tests.parecerista import MODELO_FALSO, PareceristaFalso

RAIZ = Path(__file__).parent.parent
CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.400.600/0001-80"
CNPJ_C = "30.600.900/0001-10"
Motivo = MotivoDeNaoIdentificacao


def _cenario(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF A1", valor="1.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF A2", valor="2.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF B1", valor="500,00", cliente="Brisa Serena Armazem Ltda", cnpj=CNPJ_B),
            TituloProtheus(documento="NF B2", valor="700,00", cliente="Brisa Serena Armazem Ltda", cnpj=CNPJ_B),
            TituloProtheus(
                documento="NF C1", valor="800,00", cliente="Recanto Feliz Ltda", cnpj=CNPJ_C, nosso_numero="000077700001"
            ),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "PIX RECEBIDO AURORA PANIF", "1.000,00"),
            LinhaItau("05/03/2025", "TED RECEBIDA CNPJ 10200300000140", "1.234,00"),
            LinhaItau("06/03/2025", "PIX RECEBIDO BRISA SERENA ARM", "500,00"),
            LinhaItau("06/03/2025", "DEPOSITO EM DINHEIRO", "700,00"),
            LinhaItau("06/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "750,00", documento="000077700001"),
            LinhaItau("07/03/2025", "DEPOSITO EM DINHEIRO", "333,00"),
            LinhaItau("07/03/2025", "TARIFA BANCARIA", "2,50", tipo="D"),
        ],
    )


def _linhas(caminho: Path, aba: str) -> list[tuple[object, ...]]:
    return [tuple(linha) for linha in load_workbook(caminho)[aba].iter_rows(values_only=True)]


def test_ranking_conta_creditos_na_fila_que_o_identificador_resolve_por_franqueado(pasta: Path) -> None:
    _cenario(pasta)

    ranking = analisar_impacto(pasta, configuracao_itau(), PareceristaFalso()).ranking

    assert [(f.nome, f.quantidade, f.valor_centavos) for f in ranking.franqueados] == [
        ("Brisa Serena Armazem Ltda (raiz 20400600)", 2, 120000),
        ("Aurora Panificadora Ltda (raiz 10200300)", 1, 100000),
    ]
    assert (ranking.sem_franqueado.quantidade, ranking.sem_franqueado.valor_centavos) == (1, 33300)


def test_credito_sem_pagador_conta_para_o_franqueado_do_titulo_sugerido(pasta: Path) -> None:
    _cenario(pasta)

    analise = analisar_impacto(pasta, configuracao_itau(), PareceristaFalso())

    [deposito] = [item for item in analise.antes.fila_de_excecoes if item.lancamento.valor_centavos == 70000]
    assert deposito.motivo is Motivo.SEM_IDENTIFICACAO_DO_PAGADOR
    assert deposito.franqueado is None
    assert [t.documento for t in deposito.titulos] == ["NF B2"]
    assert analise.ranking.franqueados[0].raiz == "20400600"


def _lancamento(valor_centavos: int, ordinal: int) -> Lancamento:
    return Lancamento("itau", "1234-5", date(2025, 3, 5), "PIX RECEBIDO", "", valor_centavos, Natureza.CREDITO, ordinal)


def _item(motivo: Motivo | None, franqueado: str | None, valor_centavos: int, ordinal: int) -> Item:
    return Item(
        identificador=f"C{ordinal}",
        destino=Destino.FILA_DE_EXCECOES if motivo else Destino.LISTA_DE_BAIXA,
        lancamento=_lancamento(valor_centavos, ordinal),
        titulos=(),
        explicacao="",
        motivo=motivo,
        evidencia=(),
        pareceres=(),
        franqueado=franqueado,
    )


def test_ranking_deixa_de_fora_motivos_que_o_identificador_nao_resolve() -> None:
    fila = [_item(motivo, f"{numero:08d}", 1000, numero) for numero, motivo in enumerate(Motivo, start=1)]

    ranking = RankingDeMigracao.da_fila(fila, {}, tamanho_da_onda=10)

    assert {f.raiz for f in ranking.franqueados} == {
        f"{numero:08d}"
        for numero, motivo in enumerate(Motivo, start=1)
        if motivo in {
            Motivo.SEM_IDENTIFICACAO_DO_PAGADOR,
            Motivo.NOME_AMBIGUO,
            Motivo.MAIS_DE_UMA_COMBINACAO,
            Motivo.PAGADOR_PROVAVEL_PRIMEIRA_VEZ,
        }
    }


def test_ranking_so_olha_a_fila_de_excecoes() -> None:
    ranking = RankingDeMigracao.da_fila([_item(None, "11111111", 1000, 1)], {}, tamanho_da_onda=10)

    assert ranking.franqueados == ()


def test_empate_em_quantidade_e_desfeito_pelo_valor() -> None:
    fila = [
        _item(Motivo.NOME_AMBIGUO, "11111111", 1000, 1),
        _item(Motivo.MAIS_DE_UMA_COMBINACAO, "22222222", 900, 2),
        _item(Motivo.PAGADOR_PROVAVEL_PRIMEIRA_VEZ, "22222222", 900, 3),
        _item(Motivo.SEM_IDENTIFICACAO_DO_PAGADOR, "33333333", 5000, 4),
        _item(Motivo.PAGADOR_PROVAVEL_PRIMEIRA_VEZ, "11111111", 1000, 5),
    ]

    ranking = RankingDeMigracao.da_fila(fila, {"11111111": ["Alfa Ltda", "ALFA LTDA"]}, tamanho_da_onda=10)

    assert [(f.raiz, f.quantidade, f.valor_centavos) for f in ranking.franqueados] == [
        ("11111111", 2, 2000),
        ("22222222", 2, 1800),
        ("33333333", 1, 5000),
    ]
    assert ranking.franqueados[0].nome == "ALFA LTDA / Alfa Ltda (raiz 11111111)"
    assert ranking.franqueados[1].nome == "raiz 22222222"


def test_ranking_dividido_em_ondas_do_tamanho_configurado() -> None:
    fila = [
        _item(Motivo.SEM_IDENTIFICACAO_DO_PAGADOR, f"{franqueado:08d}", 1000, franqueado * 100 + credito)
        for franqueado in range(1, 26)
        for credito in range(franqueado)
    ]

    ranking = RankingDeMigracao.da_fila(fila, {}, tamanho_da_onda=10)

    assert [len(onda) for onda in ranking.ondas] == [10, 10, 5]
    assert [f.onda for f in ranking.franqueados] == [1] * 10 + [2] * 10 + [3] * 5
    assert [f.raiz for f in ranking.ondas[0]] == [f"{n:08d}" for n in range(25, 15, -1)]


def test_tamanho_da_onda_padrao_e_10_franqueados() -> None:
    assert Configuracao(contas=()).analise_de_impacto.tamanho_da_onda == 10


def test_antes_e_a_mesma_conciliacao_da_execucao_diaria_com_armazenamento_temporario(
    pasta: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _cenario(pasta)
    monkeypatch.chdir(tmp_path)
    arquivos_antes = sorted(p.relative_to(tmp_path) for p in tmp_path.rglob("*"))

    primeira = analisar_impacto(pasta, configuracao_itau(), PareceristaFalso())
    segunda = analisar_impacto(pasta, configuracao_itau(), PareceristaFalso())

    assert sorted(p.relative_to(tmp_path) for p in tmp_path.rglob("*")) == arquivos_antes
    diaria = executar_conciliacao(pasta, Armazenamento(tmp_path / "diaria.sqlite"), configuracao_itau(), PareceristaFalso())

    def destinos(itens: tuple[Item, ...]) -> list[tuple[object, ...]]:
        return [(i.lancamento.chave_natural, i.destino, i.motivo, i.titulos, i.explicacao) for i in itens]

    assert destinos(primeira.antes.itens) == destinos(segunda.antes.itens) == destinos(diaria.itens)
    assert primeira.antes.resumo.execucao == segunda.antes.resumo.execucao == 1


def test_antes_usa_o_parecerista_de_nome_informado(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF A1", valor="1.000,00", cliente="Aurora Panificadora Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF B1", valor="1.000,00", cliente="Brisa Serena Armazem Ltda", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO AUR PANIFICADORA", "1.000,00")])
    parecerista = PareceristaFalso(sim_para=[("AUR PANIFICADORA", "Aurora Panificadora Ltda")])

    analise = analisar_impacto(pasta, configuracao_itau(), parecerista)

    assert parecerista.chamadas
    [item] = analise.antes.fila_de_excecoes
    assert item.motivo is Motivo.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert {p.parecer.modelo for p in item.pareceres} == {MODELO_FALSO}
    assert [f.raiz for f in analise.ranking.franqueados] == ["10200300"]


def test_workbook_com_ranking_em_ondas_e_premissas(pasta: Path, tmp_path: Path) -> None:
    _cenario(pasta)
    configuracao = Configuracao(
        contas=configuracao_itau().contas,
        encargo=ParametrosDeEncargo(Decimal("2"), Decimal("1"), 5),
        analise_de_impacto=ConfiguracaoDaAnaliseDeImpacto(semente=42, tamanho_da_onda=1),
    )
    saida = tmp_path / "impacto.xlsx"

    gerar_workbook_do_impacto(analisar_impacto(pasta, configuracao, PareceristaFalso()), saida)

    assert load_workbook(saida).sheetnames == ["Ranking de migração", "Premissas"]
    ranking = _linhas(saida, "Ranking de migração")
    assert ranking[0] == ("Onda", "Posição", "Franqueado", "Raiz do CNPJ", "Créditos na Fila", "Valor")
    assert ranking[1:3] == [
        (1, 1, "Brisa Serena Armazem Ltda (raiz 20400600)", "20400600", 2, 1200),
        (2, 2, "Aurora Panificadora Ltda (raiz 10200300)", "10200300", 1, 1000),
    ]
    assert (None, None, CausaRaiz.SEM_FRANQUEADO, None, 1, 333) in ranking
    premissas = {linha[0]: linha[1] for linha in _linhas(saida, "Premissas")}
    assert premissas["Semente"] == 42
    assert premissas["Multa por atraso (%)"] == 2
    assert premissas["Juros ao mês (%)"] == 1
    assert premissas["Tolerância do Encargo (R$)"] == 0.05
    assert premissas["Tamanho da Onda de migração (Franqueados)"] == 1
    assert "nome ambíguo" in str(premissas["Motivos do Ranking de migração"])


def test_ranking_do_sintetico_soma_os_creditos_da_fila_com_os_motivos_do_ranking(tmp_path: Path) -> None:
    dados = gerar_dados_sinteticos(ConfiguracaoDoGerador(semente=20260801, titulos=300, franqueados=40))
    pasta = tmp_path / "sintetico"
    dados.escrever(pasta)

    analise = analisar_impacto(pasta, Configuracao(contas=dados.contas), PareceristaFalso())

    do_ranking = [item for item in analise.antes.fila_de_excecoes if item.motivo in RankingDeMigracao.MOTIVOS]
    ranking = analise.ranking
    assert len(do_ranking) > 20
    assert sum(f.quantidade for f in ranking.franqueados) + ranking.sem_franqueado.quantidade == len(do_ranking)
    assert sum(f.valor_centavos for f in ranking.franqueados) + ranking.sem_franqueado.valor_centavos == sum(
        item.lancamento.valor_centavos for item in do_ranking
    )
    chaves = [(-f.quantidade, -f.valor_centavos) for f in ranking.franqueados]
    assert chaves == sorted(chaves)
    assert all(len(onda) <= 10 for onda in ranking.ondas)


def test_amostra_de_julho_gera_o_ranking_sobre_a_mesma_conciliacao(tmp_path: Path) -> None:
    configuracao = carregar_configuracao(RAIZ / "conciliacao.toml")
    parecerista = PareceristaFalso(
        sim_para=[("DOCE PONTO COM ALIME", "Doce Ponto Comercio de Alimentos"), ("SABOR CIA", "Rede Sabor & Cia Ltda")]
    )

    analise = analisar_impacto(RAIZ / "data", configuracao, parecerista)
    saida = tmp_path / "impacto.xlsx"
    gerar_workbook_do_impacto(analise, saida)

    assert analise.antes.resumo.data_de_referencia.data == date(2026, 7, 31)
    assert [(f.raiz, f.quantidade, f.valor_centavos, f.onda) for f in analise.ranking.franqueados] == [
        ("98765432", 1, 4763050, 1),
        ("44555666", 1, 1298000, 1),
        ("33444555", 1, 710000, 1),
        ("88999000", 1, 567000, 1),
        ("22333444", 1, 340000, 1),
        ("12345678", 1, 320000, 1),
    ]
    assert analise.ranking.sem_franqueado.quantidade == 0
    assert load_workbook(saida).sheetnames == ["Ranking de migração", "Premissas"]


def _escrever_configuracao(caminho: Path, extra: str = "") -> Path:
    caminho.write_text(
        '[[contas]]\npadrao = "itau_cobranca*.csv"\nbanco = "itau"\nconta = "1234-5"\n' + extra, encoding="utf-8"
    )
    return caminho


def test_configuracao_da_analise_de_impacto(tmp_path: Path) -> None:
    configuracao = carregar_configuracao(
        _escrever_configuracao(tmp_path / "c.toml", "[analise_de_impacto]\nsemente = 7\ntamanho_da_onda = 5\n")
    )

    assert configuracao.analise_de_impacto == ConfiguracaoDaAnaliseDeImpacto(semente=7, tamanho_da_onda=5)
    assert carregar_configuracao(_escrever_configuracao(tmp_path / "padrao.toml")).analise_de_impacto == (
        ConfiguracaoDaAnaliseDeImpacto()
    )


@pytest.mark.parametrize(
    "secao",
    [
        "[analise_de_impacto]\ntamanho_da_onda = 0\n",
        "[analise_de_impacto]\ntamanho_da_onda = true\n",
        "[analise_de_impacto]\ntamanho_da_onda = 2.5\n",
        "[analise_de_impacto]\nsemente = -1\n",
        "[analise_de_impacto]\nsemente = \"x\"\n",
        "[analise_de_impacto]\nadesao = 0.5\n",
        "[[analise_de_impacto]]\nsemente = 3\n",
    ],
)
def test_configuracao_da_analise_de_impacto_invalida(tmp_path: Path, secao: str) -> None:
    with pytest.raises(ConfiguracaoInvalida, match="analise_de_impacto"):
        carregar_configuracao(_escrever_configuracao(tmp_path / "c.toml", secao))


def test_comando_gera_o_workbook_da_analise(pasta: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _cenario(pasta)
    saida = tmp_path / "impacto.xlsx"

    codigo = main([str(pasta), "--configuracao", str(_escrever_configuracao(tmp_path / "c.toml")), "--saida", str(saida)])

    assert codigo == 0
    assert str(saida) in capsys.readouterr().out
    assert load_workbook(saida).sheetnames == ["Ranking de migração", "Premissas"]


def test_comando_falha_sem_pasta_ou_sem_configuracao(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    configuracao = _escrever_configuracao(tmp_path / "c.toml")

    assert main([str(tmp_path / "nao_existe"), "--configuracao", str(configuracao)]) == 2
    assert "nao_existe" in capsys.readouterr().err
    assert main([str(tmp_path), "--configuracao", str(tmp_path / "falta.toml")]) == 2
    assert "falta.toml" in capsys.readouterr().err
