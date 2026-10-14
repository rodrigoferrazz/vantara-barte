from dataclasses import replace
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.configuracao import Configuracao
from conciliacao.dominio import Destino, Item, MotivoDeNaoIdentificacao, RegraDaProposta, Resultado, TipoDeEvidencia
from conciliacao.emissao import emitir_cobrancas
from conciliacao.encargo import ParametrosDeEncargo
from conciliacao.execucao import executar_conciliacao
from conciliacao.porta import PortaSimulada
from conciliacao.saidas import gerar_workbook
from tests.arquivos import (
    FRANQUEADO_DOS_TITULOS,
    LinhaItau,
    TituloProtheus,
    configuracao_com_migracao,
    escrever_extrato_itau,
    escrever_pix_recebidos,
    escrever_protheus,
    pix_recebido,
)
from tests.parecerista import PareceristaFalso

CNPJ_DO_MIGRADO = "10.200.300/0001-40"
CNPJ_DE_OUTRO_FRANQUEADO = "20.300.400/0001-50"
REFERENCIA = LinhaItau("05/03/2025", "TARIFA MANUTENCAO CONTA", "9,90", tipo="D")
PAGO_EM = "2025-03-12T10:00:00Z"


def _conciliar(
    pasta: Path, armazenamento: Armazenamento, configuracao: Configuracao | None = None, workbook: Path | None = None
) -> Resultado:
    resultado = executar_conciliacao(
        pasta, armazenamento, configuracao or configuracao_com_migracao(), PareceristaFalso(), workbook_anterior=workbook
    )
    if workbook is not None:
        gerar_workbook(resultado, workbook)
    return resultado


def _emitir(
    pasta: Path, armazenamento: Armazenamento, titulos: list[TituloProtheus], configuracao: Configuracao | None = None
) -> dict[str, str]:
    configuracao = configuracao or configuracao_com_migracao()
    escrever_protheus(pasta, titulos)
    escrever_extrato_itau(pasta, [REFERENCIA])
    _conciliar(pasta, armazenamento, configuracao)
    porta = PortaSimulada()
    emitir_cobrancas(armazenamento, configuracao, porta)
    return {cobranca.documento: cobranca.txid for cobranca in porta.emitidas}


def _pago(pasta: Path, txid: str, valor: str, historico: str = "PIX RECEBIDO FULANO", **componentes: str) -> None:
    escrever_extrato_itau(pasta, [REFERENCIA, LinhaItau("12/03/2025", historico, valor)])
    escrever_pix_recebidos(
        pasta,
        [
            pix_recebido(
                txid,
                valor.replace(".", "").replace(",", "."),
                horario=PAGO_EM,
                componentes={nome: {"valor": v} for nome, v in componentes.items()} if componentes else None,
            )
        ],
    )


def _documentos(item: Item) -> list[str]:
    return [titulo.documento for titulo in item.titulos]


def _decidir_na_lista(workbook: Path, decisao: str, motivo: str | None = None) -> None:
    livro = load_workbook(workbook)
    aba = livro["Lista de baixa"]
    cabecalho = [celula.value for celula in aba[1]]
    aba.cell(row=2, column=cabecalho.index("Decisão") + 1, value=decisao)
    aba.cell(row=2, column=cabecalho.index("Motivo da rejeição") + 1, value=motivo)
    livro.properties.lastModifiedBy = "Ana Analista"
    livro.save(workbook)


def test_atraso_dentro_do_encargo_vai_para_a_lista_com_principal_multa_e_juros_do_retorno(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    configuracao = replace(
        configuracao_com_migracao(), encargo=ParametrosDeEncargo(juros_ao_mes_percentual=Decimal("1"))
    )
    txids = _emitir(
        pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")], configuracao
    )
    _pago(pasta, txids["NF 1"], "1.020,67", original="1000.00", multa="20.00", juros="0.67")

    resultado = _conciliar(pasta, armazenamento, configuracao)

    [item] = resultado.lista_de_baixa
    assert _documentos(item) == ["NF 1"]
    assert item.regra is RegraDaProposta.IDENTIFICADOR_E_VALOR
    assert "Encargo R$ 20,67" in item.explicacao
    assert "O Retorno de cobrança separa principal R$ 1.000,00, multa R$ 20,00 e juros R$ 0,67" in item.explicacao


def test_atraso_dentro_do_encargo_sem_componentes_no_retorno_vai_para_a_lista_com_o_encargo_calculado(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    txids = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    _pago(pasta, txids["NF 1"], "1.020,00")

    resultado = _conciliar(pasta, armazenamento)

    [item] = resultado.lista_de_baixa
    assert "principal R$ 1.000,00 e Encargo R$ 20,00" in item.explicacao
    assert "Retorno de cobrança separa" not in item.explicacao


def test_atraso_pago_pelo_valor_original_vai_para_a_lista_como_encargo_nao_pago(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    txids = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    _pago(pasta, txids["NF 1"], "1.000,00")

    resultado = _conciliar(pasta, armazenamento)

    [item] = resultado.lista_de_baixa
    assert _documentos(item) == ["NF 1"]
    assert "Encargo não pago" in item.explicacao


def test_valor_acima_do_encargo_vai_para_a_fila_com_valor_nao_fecha_apontando_o_titulo(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    txids = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="10/03/2025")])
    _pago(pasta, txids["NF 1"], "1.050,00")

    resultado = _conciliar(pasta, armazenamento)

    [item] = resultado.fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.VALOR_NAO_FECHA
    assert _documentos(item) == ["NF 1"]
    assert "Valor a maior" in item.explicacao


def test_motor_nao_procura_soma_para_credito_identificado(pasta: Path, armazenamento: Armazenamento) -> None:
    txids = _emitir(
        pasta,
        armazenamento,
        [
            TituloProtheus(documento="NF 1", valor="600,00", vencimento="12/03/2025"),
            TituloProtheus(documento="NF 2", valor="400,00", vencimento="12/03/2025"),
        ],
    )
    _pago(pasta, txids["NF 1"], "1.000,00", historico=f"TED RECEBIDA CNPJ {CNPJ_DO_MIGRADO}")

    resultado = _conciliar(pasta, armazenamento)

    [item] = resultado.fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.VALOR_NAO_FECHA
    assert _documentos(item) == ["NF 1"]
    assert item.evidencia == (TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA,)
    assert "NF 2" not in item.explicacao


def test_identificador_vale_mais_que_o_pagador(pasta: Path, armazenamento: Armazenamento) -> None:
    txids = _emitir(
        pasta,
        armazenamento,
        [
            TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="12/03/2025"),
            TituloProtheus(documento="NF 9", valor="1.000,00", vencimento="12/03/2025", cnpj=CNPJ_DE_OUTRO_FRANQUEADO),
        ],
    )
    _pago(pasta, txids["NF 1"], "1.000,00", historico=f"TED RECEBIDA CNPJ {CNPJ_DE_OUTRO_FRANQUEADO}")

    resultado = _conciliar(pasta, armazenamento)

    [item] = resultado.lista_de_baixa
    assert _documentos(item) == ["NF 1"]
    assert item.franqueado == FRANQUEADO_DOS_TITULOS
    assert item.evidencia == (TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA,)


def test_identificador_de_titulo_ja_baixado_vai_para_a_fila_como_credito_sem_titulo_em_aberto(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    txids = _emitir(
        pasta,
        armazenamento,
        [
            TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="12/03/2025"),
            TituloProtheus(documento="NF 2", valor="1.000,00", vencimento="12/03/2025"),
        ],
    )
    escrever_protheus(pasta, [TituloProtheus(documento="NF 2", valor="1.000,00", vencimento="12/03/2025")])
    _pago(pasta, txids["NF 1"], "1.000,00")

    resultado = _conciliar(pasta, armazenamento)

    [item] = resultado.fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO
    assert item.titulos == ()
    assert item.franqueado == FRANQUEADO_DOS_TITULOS
    assert item.evidencia == (TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA,)
    assert f"txid {txids['NF 1']}" in item.explicacao
    assert "NF 1 (Protheus)" in item.explicacao
    assert resultado.lista_de_baixa == []


def test_segundo_pagamento_do_mesmo_identificador_depois_da_baixa_confirmada_e_pego_como_duplicado(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    txids = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="12/03/2025")])
    _pago(pasta, txids["NF 1"], "1.000,00")
    workbook = tmp_path / "conciliacao.xlsx"
    _conciliar(pasta, armazenamento, workbook=workbook)
    _decidir_na_lista(workbook, "confirmar")
    escrever_extrato_itau(pasta, [LinhaItau("13/03/2025", "PIX RECEBIDO FULANO", "1.000,00")], nome="itau_cobranca_2.csv")
    escrever_pix_recebidos(
        pasta,
        [pix_recebido(txids["NF 1"], "1000.00", horario="2025-03-13T10:00:00Z", fim_do_e2e="2")],
        nome="pix_recebidos_2.json",
    )

    resultado = _conciliar(pasta, armazenamento, workbook=workbook)

    [item] = resultado.fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO
    assert item.lancamento.data.day == 13


def test_identificador_desconhecido_segue_o_caminho_da_fase_1_sem_possivel_fonte_nao_integrada(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="12/03/2025")])
    desconhecido = "XY000000000000000000000000000001"
    escrever_extrato_itau(
        pasta,
        [
            REFERENCIA,
            LinhaItau("12/03/2025", f"TED RECEBIDA CNPJ {CNPJ_DO_MIGRADO}", "1.000,00"),
            LinhaItau("12/03/2025", "DEPOSITO NAO IDENTIFICADO", "333,00"),
        ],
    )
    escrever_pix_recebidos(
        pasta,
        [
            pix_recebido(desconhecido, "1000.00", horario=PAGO_EM),
            pix_recebido("XY000000000000000000000000000002", "333.00", horario=PAGO_EM, fim_do_e2e="2"),
        ],
    )

    resultado = _conciliar(pasta, armazenamento)

    [proposta] = resultado.lista_de_baixa
    assert proposta.evidencia == (TipoDeEvidencia.CNPJ_CPF,)
    assert proposta.regra is RegraDaProposta.DOCUMENTO_E_VALOR_UNICO
    assert f"txid {desconhecido}" in proposta.explicacao
    [na_fila] = resultado.fila_de_excecoes
    assert na_fila.motivo is MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR
    assert TipoDeEvidencia.IDENTIFICADOR_DE_COBRANCA not in na_fila.evidencia


def test_proposta_do_identificador_rejeitada_segue_o_caminho_da_fase_1_sem_o_titulo_rejeitado(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    txids = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="12/03/2025")])
    _pago(pasta, txids["NF 1"], "1.000,00", historico=f"TED RECEBIDA CNPJ {CNPJ_DO_MIGRADO}")
    workbook = tmp_path / "conciliacao.xlsx"
    _conciliar(pasta, armazenamento, workbook=workbook)
    _decidir_na_lista(workbook, "rejeitar", "não é deste Título")

    resultado = _conciliar(pasta, armazenamento, workbook=workbook)

    [item] = resultado.fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO
    assert item.evidencia == (TipoDeEvidencia.CNPJ_CPF,)
    assert "rejeitada pela Analista" in item.explicacao


def test_pagar_cobranca_identificada_nunca_cria_pagador_conhecido(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    txids = _emitir(pasta, armazenamento, [TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="12/03/2025")])
    _pago(pasta, txids["NF 1"], "1.000,00", historico="PIX RECEBIDO SOCIO DA REDE")
    workbook = tmp_path / "conciliacao.xlsx"
    _conciliar(pasta, armazenamento, workbook=workbook)
    _decidir_na_lista(workbook, "confirmar")
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="12/03/2025"),
            TituloProtheus(documento="NF 2", valor="1.500,00", vencimento="20/03/2025"),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("13/03/2025", "PIX RECEBIDO SOCIO DA REDE", "1.500,00")], nome="itau_cobranca_2.csv")

    resultado = _conciliar(pasta, armazenamento, workbook=workbook)

    [item] = resultado.fila_de_excecoes
    assert _documentos(item) == ["NF 2"]
    assert TipoDeEvidencia.PAGADOR_CONHECIDO not in item.evidencia
    assert resultado.lista_de_baixa == []


def _o_que_a_analista_ve(resultado: Resultado) -> list[tuple[object, ...]]:
    return [
        (item.destino, item.motivo, tuple(_documentos(item)), item.explicacao, item.evidencia, item.regra, item.franqueado)
        for item in (*resultado.lista_de_baixa, *resultado.fila_de_excecoes)
    ]


def test_credito_solto_de_franqueado_migrado_tem_exatamente_o_resultado_da_fase_1(tmp_path: Path) -> None:
    titulos = [
        TituloProtheus(documento="NF 1", valor="1.000,00", vencimento="12/03/2025"),
        TituloProtheus(documento="NF 2", valor="700,00", vencimento="12/03/2025"),
        TituloProtheus(documento="NF 3", valor="450,00", vencimento="10/03/2025"),
    ]
    soltos = [
        REFERENCIA,
        LinhaItau("12/03/2025", f"TED RECEBIDA CNPJ {CNPJ_DO_MIGRADO}", "1.000,00"),
        LinhaItau("12/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "700,00"),
        LinhaItau("12/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "459,00"),
        LinhaItau("12/03/2025", "DEPOSITO NAO IDENTIFICADO", "123,00"),
    ]
    resultados = []
    for nome, migrados in (("fase1", None), ("migrado", (FRANQUEADO_DOS_TITULOS,))):
        pasta = tmp_path / nome
        pasta.mkdir()
        armazenamento = Armazenamento(tmp_path / f"{nome}.sqlite")
        configuracao = configuracao_com_migracao(migrados=migrados)
        _emitir(pasta, armazenamento, titulos, configuracao)
        escrever_extrato_itau(pasta, soltos)
        escrever_pix_recebidos(pasta, [pix_recebido(None, "700.00", horario=PAGO_EM)])
        resultados.append(_conciliar(pasta, armazenamento, configuracao))
        armazenamento.fechar()

    fase_1, migrado = resultados
    assert len(_o_que_a_analista_ve(fase_1)) == 4
    assert _o_que_a_analista_ve(migrado) == _o_que_a_analista_ve(fase_1)


def test_componentes_do_valor_fora_da_api_pix_recusam_o_arquivo(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_extrato_itau(pasta, [REFERENCIA])
    sem_original = pix_recebido("CI000000000000000000000000000001", "20.00", componentes={"multa": {"valor": "20.00"}})
    estranho = pix_recebido(
        "CI000000000000000000000000000002", "20.00", componentes={"original": {"valor": "20.00"}, "taxa": {"valor": "1.00"}}
    )
    escrever_pix_recebidos(pasta, [sem_original], nome="pix_recebidos_sem_original.json")
    escrever_pix_recebidos(pasta, [estranho], nome="pix_recebidos_estranho.json")

    resultado = _conciliar(pasta, armazenamento)

    recusados = {recusado.nome: recusado.motivo for recusado in resultado.arquivos_recusados}
    assert "original" in recusados["pix_recebidos_sem_original.json"]
    assert "taxa" in recusados["pix_recebidos_estranho.json"]
