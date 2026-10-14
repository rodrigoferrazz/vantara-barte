import time
from collections import defaultdict
from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.configuracao import Configuracao, ConfiguracaoDaAnaliseDeImpacto, carregar_configuracao
from conciliacao.dominio import (
    Destino,
    Item,
    MotivoDeNaoIdentificacao,
    Natureza,
    PorQueFicaSolto,
    Resultado,
    TotalDoDestino,
)
from conciliacao.gerador import ConfiguracaoDoGerador, DadosSinteticos, LancamentoEsperado, gerar_dados_sinteticos
from conciliacao.impacto import (
    ABA_DE_PREMISSAS,
    AnaliseDeImpacto,
    CenarioDeAdesao,
    ChegadaDosCreditos,
    FonteDaVerdade,
    analisar_impacto,
    gerar_workbook_do_impacto,
)
from conciliacao.parecerista import PareceristaDeNome, PareceristaOllama
from tests.parecerista import PareceristaFalso, ollama_do_projeto

RAIZ = Path(__file__).parent.parent
AMOSTRA = RAIZ / "data"
SEMENTE_DO_GERADOR = 20260801
TAMANHO_DA_ONDA_NO_SINTETICO = 50
SEGUNDOS_ACEITAVEIS = 120

FILA = Destino.FILA_DE_EXCECOES
Motivo = MotivoDeNaoIdentificacao
Solto = PorQueFicaSolto

DOCE_PONTO = ("DOCE PONTO COM ALIME", "Doce Ponto Comercio de Alimentos")
SABOR_CIA = ("SABOR CIA", "Rede Sabor & Cia Ltda")


def _centavos(valor: str) -> int:
    return int(valor.replace(".", "").replace(",", ""))


def _total(quantidade: int, valor: str) -> TotalDoDestino:
    return TotalDoDestino(quantidade, _centavos(valor))


@dataclass(frozen=True)
class EsperadoDaAdesao:
    adesao: int
    identificados: TotalDoDestino
    fila: TotalDoDestino
    sorteio_acima_da_adesao: TotalDoDestino


FAIXA_DE_JULHO = (0, 50, 80, 95, 100)
ANTES_DE_JULHO = EsperadoDaAdesao(-1, _total(5, "188.370,00"), _total(9, "159.150,50"), _total(0, "0,00"))
DEPOIS_DE_JULHO = (
    EsperadoDaAdesao(0, _total(5, "188.370,00"), _total(9, "159.150,50"), _total(5, "48.608,00")),
    EsperadoDaAdesao(50, _total(9, "231.308,00"), _total(6, "139.570,50"), _total(1, "5.670,00")),
    EsperadoDaAdesao(80, _total(10, "236.978,00"), _total(5, "133.900,50"), _total(0, "0,00")),
    EsperadoDaAdesao(95, _total(10, "236.978,00"), _total(5, "133.900,50"), _total(0, "0,00")),
    EsperadoDaAdesao(100, _total(10, "236.978,00"), _total(5, "133.900,50"), _total(0, "0,00")),
)
SOLTOS_EM_TODA_ADESAO_DE_JULHO = {
    Solto.SEM_TITULO_CONHECIDO: _total(1, "1.850,00"),
    Solto.TITULO_SEM_COBRANCA: _total(0, "0,00"),
    Solto.NAO_IDENTIFICAVEL_PELO_VALOR: _total(4, "76.980,50"),
    Solto.FRANQUEADO_NAO_MIGRADO: _total(0, "0,00"),
}
FILA_DE_JULHO_COM_TODA_A_ADESAO = (
    ("14/07", "47.630,50", Motivo.PAGADOR_PROVAVEL_PRIMEIRA_VEZ),
    ("16/07", "62.500,00", Motivo.VALOR_NAO_FECHA),
    ("17/07", "14.820,00", Motivo.VALOR_NAO_FECHA),
    ("24/07", "7.100,00", Motivo.PAGADOR_PROVAVEL_PRIMEIRA_VEZ),
    ("30/07", "1.850,00", Motivo.CREDITO_SEM_TITULO_EM_ABERTO),
)


def _creditos(resultado: Resultado) -> list[Item]:
    return [*resultado.lista_de_baixa, *resultado.fila_de_excecoes]


def _dinheiro(resultado: Resultado) -> int:
    return sum(item.lancamento.valor_centavos for item in _creditos(resultado))


def _fila(resultado: Resultado) -> TotalDoDestino:
    return resultado.resumo.totais[FILA]


def _linhas_de_cenarios(analise: AnaliseDeImpacto) -> list[tuple[CenarioDeAdesao, ...]]:
    return [*(onda.cenarios for onda in analise.ondas), analise.cenarios]


def _todos_os_depois(analise: AnaliseDeImpacto) -> list[CenarioDeAdesao]:
    return [cenario for cenarios in _linhas_de_cenarios(analise) for cenario in cenarios]


def _nunca_aumenta(filas: list[TotalDoDestino]) -> bool:
    return all(
        depois.quantidade <= antes.quantidade and depois.valor_centavos <= antes.valor_centavos
        for antes, depois in zip(filas, filas[1:])
    )


def _conferir_dinheiro_e_fila(analise: AnaliseDeImpacto) -> None:
    assert {_dinheiro(c.depois) for c in _todos_os_depois(analise)} == {_dinheiro(analise.antes)}
    antes = _fila(analise.antes)
    linhas = _linhas_de_cenarios(analise)
    for cenarios in linhas:
        assert _nunca_aumenta([antes, *(_fila(c.depois) for c in cenarios)])
    for coluna in range(len(analise.cenarios)):
        assert _nunca_aumenta([antes, *(_fila(cenarios[coluna].depois) for cenarios in linhas)])


@pytest.fixture(
    scope="module",
    params=["ia falsa", pytest.param("ollama", marks=pytest.mark.integracao)],
)
def julho(request: pytest.FixtureRequest) -> AnaliseDeImpacto:
    parecerista: PareceristaDeNome
    if request.param == "ollama":
        parecerista = PareceristaOllama(ollama_do_projeto())
    else:
        parecerista = PareceristaFalso(sim_para=[DOCE_PONTO, SABOR_CIA])
    configuracao = carregar_configuracao(RAIZ / "conciliacao.toml")
    faixa = tuple(Decimal(adesao) for adesao in FAIXA_DE_JULHO)
    configuracao = replace(configuracao, analise_de_impacto=replace(configuracao.analise_de_impacto, faixa_de_adesao=faixa))
    return analisar_impacto(AMOSTRA, configuracao, parecerista)


def test_julho_antes_com_os_creditos_identificados_e_a_fila_da_fase_1(julho: AnaliseDeImpacto) -> None:
    chegada = ChegadaDosCreditos.do_resultado(julho.antes)

    assert julho.antes.resumo.data_de_referencia.data == date(2026, 7, 31)
    assert chegada.identificados == ANTES_DE_JULHO.identificados
    assert _fila(julho.antes) == ANTES_DE_JULHO.fila


@pytest.mark.parametrize("esperado", DEPOIS_DE_JULHO, ids=lambda e: f"Adesão {e.adesao}%")
def test_julho_depois_de_cada_adesao(julho: AnaliseDeImpacto, esperado: EsperadoDaAdesao) -> None:
    [cenario] = [c for c in julho.cenarios if c.adesao == esperado.adesao]

    assert ChegadaDosCreditos.do_resultado(cenario.depois).identificados == esperado.identificados
    assert _fila(cenario.depois) == esperado.fila
    assert cenario.soltos == {**SOLTOS_EM_TODA_ADESAO_DE_JULHO, Solto.FORA_DA_ADESAO: esperado.sorteio_acima_da_adesao}


def test_julho_adesao_zero_da_o_antes(julho: AnaliseDeImpacto) -> None:
    [zero] = [c for c in julho.cenarios if c.adesao == 0]

    assert sorted((i.lancamento.chave_natural, i.destino, i.motivo) for i in zero.depois.itens) == sorted(
        (i.lancamento.chave_natural, i.destino, i.motivo) for i in julho.antes.itens
    )


def test_julho_com_toda_a_adesao_so_fica_na_fila_o_que_o_identificador_nao_resolve(julho: AnaliseDeImpacto) -> None:
    toda = julho.cenarios[-1]

    assert sorted(
        (f"{i.lancamento.data:%d/%m}", i.lancamento.valor_centavos, i.motivo) for i in toda.depois.fila_de_excecoes
    ) == sorted((data, _centavos(valor), motivo) for data, valor, motivo in FILA_DE_JULHO_COM_TODA_A_ADESAO)


def test_julho_dinheiro_igual_e_fila_que_nunca_aumenta_com_a_adesao_ou_a_onda(julho: AnaliseDeImpacto) -> None:
    [onda] = julho.ondas

    assert onda.franqueados == {f.raiz for f in julho.ranking.franqueados}
    _conferir_dinheiro_e_fila(julho)


def test_julho_premissas_dizem_que_a_verdade_vem_das_propostas_e_sugestoes(
    julho: AnaliseDeImpacto, tmp_path: Path
) -> None:
    saida = tmp_path / "impacto.xlsx"
    gerar_workbook_do_impacto(julho, saida)

    premissas = {linha[0]: linha[1] for linha in load_workbook(saida)[ABA_DE_PREMISSAS].iter_rows(values_only=True)}
    assert julho.fontes == {FonteDaVerdade.PROPOSTA_OU_SUGESTAO: 15}
    assert premissas[f"Créditos pela fonte: {FonteDaVerdade.PROPOSTA_OU_SUGESTAO.value}"] == 15
    assert premissas[f"Créditos pela fonte: {FonteDaVerdade.GABARITO.value}"] == 0
    assert premissas[f"Créditos pela fonte: {FonteDaVerdade.DECISOES_DA_ANALISTA.value}"] == 0
    assert str(premissas["Limite da fonte da verdade"]).startswith("15 de 15 Créditos não têm gabarito nem decisão")


@dataclass(frozen=True)
class Sintetico:
    dados: DadosSinteticos
    analise: AnaliseDeImpacto
    segundos: float


@pytest.fixture(scope="module")
def sintetico(tmp_path_factory: pytest.TempPathFactory) -> Sintetico:
    pasta = tmp_path_factory.mktemp("sintetico")
    inicio = time.perf_counter()
    dados = gerar_dados_sinteticos(ConfiguracaoDoGerador(semente=SEMENTE_DO_GERADOR))
    dados.escrever(pasta)
    parecerista = PareceristaFalso(
        sim_para=(
            (esperado.nome_no_extrato, nome)
            for esperado in dados.gabarito
            if esperado.nome_no_extrato and esperado.pagador
            for nome in dados.nomes_dos_franqueados[esperado.pagador]
        )
    )
    configuracao = Configuracao(
        contas=dados.contas,
        analise_de_impacto=ConfiguracaoDaAnaliseDeImpacto(tamanho_da_onda=TAMANHO_DA_ONDA_NO_SINTETICO),
    )
    analise = analisar_impacto(pasta, configuracao, parecerista, gabarito=dados.gabarito)
    return Sintetico(dados, analise, time.perf_counter() - inicio)


type _Chave = tuple[str, date, str, int]


def _chave(item: Item) -> _Chave:
    lancamento = item.lancamento
    return (lancamento.conta, lancamento.data, lancamento.historico, lancamento.valor_centavos)


def _titulos(item: Item) -> frozenset[tuple[str, str]]:
    return frozenset((titulo.sistema_de_origem, titulo.documento) for titulo in item.titulos)


def _divergencias_da_lista(depois: Resultado, gabarito: tuple[LancamentoEsperado, ...]) -> list[str]:
    creditos = [e for e in gabarito if e.natureza is Natureza.CREDITO]
    por_chave = {(e.conta, e.data, e.historico, e.valor_centavos): e for e in creditos}
    no_depois = {_chave(item) for item in depois.itens}
    pagos_por_identificador: dict[date, set[tuple[str, str]]] = defaultdict(set)
    for chave, saiu in por_chave.items():
        if chave not in no_depois:
            pagos_por_identificador[saiu.data] |= saiu.titulos
    divergencias = []
    for item in depois.lista_de_baixa:
        esperado = por_chave.get(_chave(item))
        if esperado is None:
            if len(item.titulos) != 1 or not _titulos(item) <= pagos_por_identificador[item.lancamento.data]:
                divergencias.append(f"pagamento identificado {item.lancamento.historico}: {item.explicacao}")
        elif _titulos(item) != esperado.titulos:
            divergencias.append(f"{item.lancamento.historico}: {item.explicacao}")
    return divergencias


def test_sintetico_tem_cerca_de_1900_titulos_de_4_contas_e_mais_de_uma_onda(sintetico: Sintetico) -> None:
    titulos = {titulo for esperado in sintetico.dados.gabarito for titulo in esperado.titulos}

    assert len(titulos) > 1_500
    assert len({(conta.banco, conta.conta) for conta in sintetico.dados.contas}) == 4
    assert len(sintetico.analise.ondas) > 1
    assert sintetico.analise.fontes == {FonteDaVerdade.GABARITO: len(_creditos(sintetico.analise.antes))}


def test_sintetico_nenhum_item_da_lista_do_depois_diverge_do_gabarito(sintetico: Sintetico) -> None:
    for cenario in _todos_os_depois(sintetico.analise):
        assert _divergencias_da_lista(cenario.depois, sintetico.dados.gabarito) == []
    toda = sintetico.analise.cenarios[-1].depois
    assert len(toda.lista_de_baixa) > len(sintetico.analise.antes.lista_de_baixa)


def test_sintetico_dinheiro_igual_e_fila_que_nunca_aumenta_com_a_adesao_ou_a_onda(sintetico: Sintetico) -> None:
    _conferir_dinheiro_e_fila(sintetico.analise)
    assert _fila(sintetico.analise.cenarios[-1].depois).quantidade < _fila(sintetico.analise.antes).quantidade


def test_sintetico_termina_em_tempo_aceitavel_com_a_ia_falsa(sintetico: Sintetico) -> None:
    assert sintetico.segundos < SEGUNDOS_ACEITAVEIS

