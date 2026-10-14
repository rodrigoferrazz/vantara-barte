from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from conciliacao.armazenamento import Armazenamento
from conciliacao.dominio import (
    Destino,
    MotivoDeNaoIdentificacao,
    Resultado,
    SituacaoNoErp,
    TipoDeDecisao,
)
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import (
    LinhaItau,
    TituloOmie,
    TituloProtheus,
    configuracao_itau,
    escrever_extrato_itau,
    escrever_omie,
    escrever_protheus,
)
from tests.parecerista import PareceristaFalso

CNPJ = "10.200.300/0001-40"
OUTRO_CNPJ = "20.300.400/0001-50"
ANALISTA = "Ana Analista"


def _executar(pasta: Path, armazenamento: Armazenamento, workbook: Path) -> Resultado:
    resultado = executar_conciliacao(
        pasta, armazenamento, configuracao_itau(), PareceristaFalso(), workbook_anterior=workbook
    )
    gerar_workbook(resultado, workbook)
    return resultado


def _linhas(aba: Worksheet) -> list[dict[str, object]]:
    cabecalho, *linhas = aba.iter_rows(values_only=True)
    return [dict(zip([str(c) for c in cabecalho], linha)) for linha in linhas]


def _resumo(workbook: Path) -> dict[object, object]:
    return {linha[0]: linha[1] for linha in load_workbook(workbook)["Resumo"].iter_rows(values_only=True)}


def _salvar(workbook: Path) -> None:
    livro = load_workbook(workbook)
    livro.properties.lastModifiedBy = ANALISTA
    livro.save(workbook)


def _resolver(workbook: Path, decisoes: dict[str, dict[str, str]]) -> None:
    livro = load_workbook(workbook)
    aba = livro["Fila de exceções"]
    cabecalho = [str(c.value) for c in aba[1]]
    for linha in range(2, aba.max_row + 1):
        item = str(aba.cell(row=linha, column=cabecalho.index("Item") + 1).value)
        for coluna, valor in decisoes.get(item, {}).items():
            aba.cell(row=linha, column=cabecalho.index(coluna) + 1).value = valor
    livro.properties.lastModifiedBy = ANALISTA
    livro.save(workbook)


def _reverter(workbook: Path, reversoes: dict[str, str | None]) -> None:
    livro = load_workbook(workbook)
    aba = livro["Reverter confirmação"]
    assert [str(c.value) for c in aba[1]] == ["Item", "Motivo da reversão"]
    for item, motivo in reversoes.items():
        aba.append((item, motivo))
    livro.properties.lastModifiedBy = ANALISTA
    livro.save(workbook)


def _protheus(pasta: Path, *titulos: TituloProtheus) -> None:
    escrever_protheus(pasta, list(titulos))


def _extrato(pasta: Path, *linhas: LinhaItau, nome: str = "itau_cobranca.csv") -> None:
    escrever_extrato_itau(pasta, list(linhas), nome=nome)


def _textos_do_workbook(workbook: Path) -> str:
    livro = load_workbook(workbook)
    return " ".join(
        str(valor) for aba in livro.worksheets for linha in aba.iter_rows(values_only=True) for valor in linha
    )


def _situacoes(resultado: Resultado) -> list[tuple[SituacaoNoErp, str]]:
    return [(t.situacao, t.titulo.documento) for t in resultado.acompanhamento_do_erp]


A = TituloProtheus(documento="NF 140001", valor="1.000,00", nosso_numero="000014000011", cnpj=CNPJ)
B = TituloProtheus(documento="NF 140002", valor="500,00", cnpj=CNPJ)
LIQUIDACAO_A = LinhaItau("05/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "1.000,00", documento="000014000011")
TARIFA = LinhaItau("20/03/2025", "TARIFA BANCARIA", "10,00", tipo="D")


def _a_confirmado(pasta: Path, armazenamento: Armazenamento, workbook: Path) -> str:
    _protheus(pasta, A, B)
    _extrato(pasta, LIQUIDACAO_A, TARIFA)
    [item] = _executar(pasta, armazenamento, workbook).lista_de_baixa
    _salvar(workbook)
    return item.identificador


def test_titulo_confirmado_que_some_da_exportacao_sai_de_todas_as_saidas(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    _a_confirmado(pasta, armazenamento, workbook)
    _protheus(pasta, B)

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.resumo.decisoes.confirmadas == 1
    assert resultado.acompanhamento_do_erp == ()
    assert resultado.baixa_manual == ()
    assert [t.titulo.documento for t in resultado.cobranca] == ["NF 140002"]
    assert "NF 140001" not in _textos_do_workbook(workbook)


def test_confirmado_e_ainda_aberto_fica_aguardando_baixa_no_erp_fora_da_cobranca_ate_sumir(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    item = _a_confirmado(pasta, armazenamento, workbook)

    segunda = _executar(pasta, armazenamento, workbook)

    [aguardando] = segunda.acompanhamento_do_erp
    assert (aguardando.situacao, aguardando.titulo.documento, aguardando.item) == (
        SituacaoNoErp.AGUARDANDO_BAIXA_NO_ERP,
        "NF 140001",
        item,
    )
    assert aguardando.saldo_centavos == 0
    assert [t.titulo.documento for t in segunda.cobranca] == ["NF 140002"]
    acompanhamento = _linhas(load_workbook(workbook)["Acompanhamento do ERP"])
    assert [(linha["Situação"], linha["Título"], linha["Item"]) for linha in acompanhamento] == [
        ("Aguardando baixa no ERP", "NF 140001", item)
    ]
    assert _resumo(workbook)["Aguardando baixa no ERP"] == 1

    _protheus(pasta, B)
    terceira = _executar(pasta, armazenamento, workbook)

    assert terceira.acompanhamento_do_erp == ()
    assert "NF 140001" not in _textos_do_workbook(workbook)


def test_titulo_nao_proposto_que_some_e_baixado_por_fora(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    outro = TituloProtheus(documento="NF 140101", valor="300,00", cnpj=OUTRO_CNPJ)
    _protheus(pasta, A, outro)
    _extrato(pasta, LIQUIDACAO_A, TARIFA)
    _executar(pasta, armazenamento, workbook)
    _protheus(pasta, A)

    segunda = _executar(pasta, armazenamento, workbook)

    assert _situacoes(segunda) == [(SituacaoNoErp.BAIXADO_POR_FORA, "NF 140101")]
    acompanhamento = _linhas(load_workbook(workbook)["Acompanhamento do ERP"])
    assert [(linha["Situação"], linha["Título"]) for linha in acompanhamento] == [("Baixado por fora", "NF 140101")]
    assert _resumo(workbook)["Baixados por fora"] == 1

    terceira = _executar(pasta, armazenamento, workbook)

    assert terceira.acompanhamento_do_erp == ()


def test_titulo_proposto_e_ainda_nao_confirmado_que_some_nao_e_baixado_por_fora(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    _protheus(pasta, A, B)
    _extrato(pasta, LIQUIDACAO_A, TARIFA)
    _executar(pasta, armazenamento, workbook)
    _protheus(pasta, B)

    segunda = _executar(pasta, armazenamento, workbook)

    assert segunda.acompanhamento_do_erp == ()
    [credito] = [item for item in segunda.itens if item.destino is not Destino.FORA_DO_ESCOPO]
    assert credito.destino is Destino.FILA_DE_EXCECOES


def test_baixa_de_titulo_diferente_do_proposto_gera_divergencia_e_volta_para_a_fila(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    item = _a_confirmado(pasta, armazenamento, workbook)
    _executar(pasta, armazenamento, workbook)
    _protheus(pasta, A)

    resultado = _executar(pasta, armazenamento, workbook)

    assert _situacoes(resultado) == [(SituacaoNoErp.DIVERGENCIA_DE_BAIXA, "NF 140002")]
    assert resultado.acompanhamento_do_erp[0].item == item
    [na_fila] = resultado.fila_de_excecoes
    assert na_fila.identificador == item
    assert na_fila.motivo is MotivoDeNaoIdentificacao.MAIS_DE_UMA_COMBINACAO
    assert [t.documento for t in na_fila.titulos] == ["NF 140001"]
    assert "Divergência de baixa" in na_fila.explicacao
    assert "NF 140002" in na_fila.explicacao
    assert resultado.lista_de_baixa == []
    assert [t.titulo.documento for t in resultado.em_analise] == ["NF 140001"]
    assert _resumo(workbook)["Divergências de baixa"] == 1

    seguinte = _executar(pasta, armazenamento, workbook)

    assert [i.identificador for i in seguinte.fila_de_excecoes] == [item]
    assert seguinte.acompanhamento_do_erp == ()


def _pagamento_parcial_confirmado(pasta: Path, armazenamento: Armazenamento, workbook: Path) -> None:
    _protheus(pasta, TituloProtheus(documento="NF 140201", valor="1.000,00", cnpj=CNPJ))
    _extrato(pasta, LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ}", "400,00"), TARIFA)
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    assert item.motivo is MotivoDeNaoIdentificacao.VALOR_NAO_FECHA
    _resolver(workbook, {item.identificador: {"Decisão": "confirmar"}})


def test_pagamento_parcial_confirmado_deixa_o_titulo_em_aberto_e_na_cobranca_so_pelo_saldo(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    _pagamento_parcial_confirmado(pasta, armazenamento, workbook)

    resultado = _executar(pasta, armazenamento, workbook)

    [baixa] = resultado.baixa_manual
    assert (baixa.principal_centavos, baixa.encargo_centavos) == (40000, 0)
    [cobrado] = resultado.lista_de_cobranca
    assert (cobrado.titulo.documento, cobrado.saldo_centavos) == ("NF 140201", 60000)
    [aguardando] = resultado.acompanhamento_do_erp
    assert (aguardando.situacao, aguardando.saldo_centavos) == (SituacaoNoErp.AGUARDANDO_BAIXA_NO_ERP, 60000)
    cobranca = _linhas(load_workbook(workbook)["Lista de cobrança"])
    assert (cobranca[0]["Valor"], cobranca[0]["Saldo"]) == (1000, 600)


def test_titulo_em_possivel_duplicidade_com_pagamento_parcial_confirmado_continua_fora_da_cobranca(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    _protheus(pasta, TituloProtheus(documento="NF 140501", valor="1.000,00", cnpj=CNPJ))
    escrever_omie(pasta, [TituloOmie(documento="FAT-14051", valor="1000.00")])
    _extrato(pasta, LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ}", "400,00"), TARIFA)
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {"Decisão": "confirmar", "Títulos escolhidos": "NF 140501"}})

    resultado = _executar(pasta, armazenamento, workbook)

    [baixa] = resultado.baixa_manual
    assert baixa.principal_centavos == 40000
    [duplicidade] = resultado.itens_de_titulo
    assert duplicidade.motivo is MotivoDeNaoIdentificacao.POSSIVEL_DUPLICIDADE
    assert resultado.cobranca == ()


def test_valor_igual_entre_exportacoes_tem_saldo_pelas_decisoes_e_casa_o_resto(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    _pagamento_parcial_confirmado(pasta, armazenamento, workbook)
    _executar(pasta, armazenamento, workbook)
    _extrato(pasta, LinhaItau("07/03/2025", f"TED RECEBIDA CNPJ {CNPJ}", "600,00"), nome="itau_cobranca_2.csv")

    resultado = _executar(pasta, armazenamento, workbook)

    [na_lista] = resultado.lista_de_baixa
    assert na_lista.lancamento.valor_centavos == 60000
    assert [t.documento for t in na_lista.titulos] == ["NF 140201"]
    assert resultado.em_analise == ()
    assert "NF 140201" not in [t.titulo.documento for t in resultado.lista_de_cobranca]


def test_valor_que_muda_entre_exportacoes_vira_o_saldo_sem_descontar_de_novo_as_decisoes(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    _pagamento_parcial_confirmado(pasta, armazenamento, workbook)
    _executar(pasta, armazenamento, workbook)
    _protheus(pasta, TituloProtheus(documento="NF 140201", valor="600,00", cnpj=CNPJ))

    resultado = _executar(pasta, armazenamento, workbook)

    [cobrado] = resultado.lista_de_cobranca
    assert (cobrado.titulo.valor_centavos, cobrado.saldo_centavos) == (60000, 60000)
    assert resultado.acompanhamento_do_erp == ()


def test_confirmacao_de_varios_titulos_com_credito_menor_que_a_soma_nao_e_aplicada(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    _protheus(
        pasta,
        TituloProtheus(documento="NF 140301", valor="1.000,00", cnpj=CNPJ),
        TituloProtheus(documento="NF 140302", valor="1.000,00", cnpj=CNPJ),
    )
    _extrato(pasta, LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ}", "1.500,00"), TARIFA)
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {"Decisão": "confirmar", "Títulos escolhidos": "NF 140301; NF 140302"}})

    resultado = _executar(pasta, armazenamento, workbook)

    assert [i.identificador for i in resultado.fila_de_excecoes] == [item.identificador]
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert item.identificador in aviso and "menor que a soma" in aviso


def test_reversao_com_titulo_ainda_aberto_volta_o_caso_para_a_fila(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    item = _a_confirmado(pasta, armazenamento, workbook)
    _executar(pasta, armazenamento, workbook)
    _reverter(workbook, {item: "Título errado"})

    resultado = _executar(pasta, armazenamento, workbook)

    [na_fila] = resultado.fila_de_excecoes
    assert na_fila.identificador == item
    assert "Título errado" in na_fila.explicacao
    assert resultado.acompanhamento_do_erp == ()
    assert resultado.resumo.decisoes.revertidas == 1
    assert "NF 140001" in [t.titulo.documento for t in resultado.cobranca]
    reversao = armazenamento.decisoes()[-1]
    assert (reversao.tipo, reversao.item, reversao.motivo, reversao.autor) == (
        TipoDeDecisao.REVERTER,
        item,
        "Título errado",
        ANALISTA,
    )
    assert [t.documento for t in reversao.titulos] == ["NF 140001"]
    assert "1 revertida(s)" in str(_resumo(workbook)["Decisões lidas do Excel"])


def test_reversao_com_titulo_ja_sumido_marca_estorno_necessario_ate_ele_reaparecer(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    item = _a_confirmado(pasta, armazenamento, workbook)
    _protheus(pasta, B)
    _executar(pasta, armazenamento, workbook)
    _reverter(workbook, {item: "Crédito era de outro Franqueado"})

    revertida = _executar(pasta, armazenamento, workbook)

    assert _situacoes(revertida) == [(SituacaoNoErp.ESTORNO_NECESSARIO_NO_ERP, "NF 140001")]
    assert revertida.acompanhamento_do_erp[0].item == item
    assert [i.identificador for i in revertida.fila_de_excecoes] == [item]
    assert _resumo(workbook)["Estorno necessário no ERP"] == 1

    ainda_sumido = _executar(pasta, armazenamento, workbook)

    assert _situacoes(ainda_sumido) == [(SituacaoNoErp.ESTORNO_NECESSARIO_NO_ERP, "NF 140001")]

    _protheus(pasta, A, B)
    reapareceu = _executar(pasta, armazenamento, workbook)

    assert reapareceu.acompanhamento_do_erp == ()
    assert [i.identificador for i in reapareceu.fila_de_excecoes] == [item]

    _protheus(pasta, B)
    sumiu_de_novo = _executar(pasta, armazenamento, workbook)

    assert SituacaoNoErp.ESTORNO_NECESSARIO_NO_ERP not in [t.situacao for t in sumiu_de_novo.acompanhamento_do_erp]


def test_reversao_desfaz_o_pagador_conhecido_aprendido_na_decisao(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    _protheus(
        pasta,
        TituloProtheus(documento="NF 140401", valor="1.000,00", cliente="Franquia Teste Ltda"),
        TituloProtheus(documento="NF 140402", valor="2.000,00", cliente="Franquia Teste Ltda"),
    )
    _extrato(pasta, LinhaItau("05/03/2025", "PIX RECEBIDO FRANQUIA TESTE", "1.000,00"))
    [item] = _executar(pasta, armazenamento, workbook).fila_de_excecoes
    _resolver(workbook, {item.identificador: {"Decisão": "confirmar"}})
    _executar(pasta, armazenamento, workbook)
    _reverter(workbook, {item.identificador: "Pagador não é o Franqueado"})
    _extrato(pasta, LinhaItau("06/03/2025", "TED RECEBIDA FRANQUIA TESTE", "2.000,00"), nome="itau_cobranca_2.csv")

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.lista_de_baixa == []
    novo = next(i for i in resultado.fila_de_excecoes if i.lancamento.valor_centavos == 200000)
    assert novo.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ


def test_reversao_sem_motivo_ou_sem_confirmacao_em_vigor_nao_e_aplicada(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    workbook = tmp_path / "conciliacao.xlsx"
    item = _a_confirmado(pasta, armazenamento, workbook)
    _executar(pasta, armazenamento, workbook)
    _reverter(workbook, {item: None, "L999": "não existe"})

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.resumo.decisoes.revertidas == 0
    assert len(resultado.resumo.decisoes.nao_aplicadas) == 2
    assert [t.situacao for t in resultado.acompanhamento_do_erp] == [SituacaoNoErp.AGUARDANDO_BAIXA_NO_ERP]
