import getpass
import socket
import subprocess
import sys
import threading
import urllib.error
import urllib.parse
import urllib.request
from collections.abc import Iterator
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.dominio import DecisaoRegistrada, MotivoDeNaoIdentificacao, Resultado, TipoDeDecisao
from conciliacao.encargo import ParametrosDeEncargo
from conciliacao.execucao import executar_conciliacao
from conciliacao.normalizacao import formatar_reais
from conciliacao.pagina_da_analista import MOTIVO_EM_LINGUAGEM_SIMPLES, POR_QUE_DA_REGRA
from conciliacao.saidas import gerar_workbook
from conciliacao.tela import criar_servidor
from tests.arquivos import (
    LinhaItau,
    TituloOmie,
    TituloProtheus,
    configuracao_itau,
    escrever_extrato_itau,
    escrever_omie,
    escrever_protheus,
)
from tests.pagina import Elemento, ler_pagina
from tests.parecerista import PareceristaFalso

CNPJ_B = "20.400.600/0001-80"
NOME_B = "Confeitaria Alvorada Comercio de Alimentos Ltda"
CONFIGURACAO = '[[contas]]\npadrao = "itau_cobranca*.csv"\nbanco = "itau"\nconta = "1234-5"\n'


@pytest.fixture
def estado(tmp_path: Path) -> Path:
    return tmp_path / "estado.sqlite"


@pytest.fixture
def saida(tmp_path: Path) -> Path:
    return tmp_path / "conciliacao.xlsx"


@pytest.fixture
def endereco(estado: Path, pasta: Path, saida: Path, tmp_path: Path) -> Iterator[str]:
    configuracao = tmp_path / "conciliacao.toml"
    configuracao.write_text(CONFIGURACAO, encoding="utf-8")
    servidor = criar_servidor(
        estado,
        porta=0,
        entrada=pasta,
        configuracao=configuracao,
        saida=saida,
        parecerista_de=lambda _: PareceristaFalso(),
    )
    linha = threading.Thread(target=servidor.serve_forever, daemon=True)
    linha.start()
    try:
        yield f"http://127.0.0.1:{servidor.server_address[1]}"
    finally:
        servidor.shutdown()
        servidor.server_close()


def _html(endereco: str, caminho: str = "/") -> str:
    with urllib.request.urlopen(endereco + caminho) as resposta:
        assert resposta.status == 200
        assert resposta.headers.get_content_type() == "text/html"
        texto: str = resposta.read().decode("utf-8")
        return texto


def _abrir(endereco: str, caminho: str = "/") -> Elemento:
    return ler_pagina(_html(endereco, caminho))


def _executar(pasta: Path, armazenamento: Armazenamento, parecerista: PareceristaFalso | None = None) -> Resultado:
    return executar_conciliacao(pasta, armazenamento, configuracao_itau(), parecerista or PareceristaFalso())


def _cobranca_e_em_analise(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF NA LISTA", valor="1.000,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF NA FILA", valor="500,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF SEM CREDITO", valor="730,00", vencimento="02/06/2025"),
            TituloProtheus(documento="NF A VENCER", valor="90,00", vencimento="20/06/2025"),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("02/06/2025", "PIX RECEBIDO 10200300000140", "1.000,00"),
            LinhaItau("03/06/2025", "DEPOSITO EM DINHEIRO", "500,00"),
            LinhaItau("10/06/2025", "TARIFA BANCARIA", "1,00", tipo="D"),
        ],
    )


def _parecer_de_nome(pasta: Path) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 7", valor="3.200,00", cliente=NOME_B, cnpj=CNPJ_B)])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO CONF ALVORADA COM", "3.200,00")])


def _texto_da_celula(valor: object) -> str:
    return "" if valor is None else str(valor)


def _itens_com_decisao(secao: Elemento) -> list[str]:
    valores = [i.atributos["value"] for i in secao.todos("input") if i.atributos.get("name") == "item"]
    return list(dict.fromkeys(valores))


def _cartao(pagina: Elemento, item: str) -> Elemento:
    return pagina.por_id(f"item-{item}")


def test_sem_execucao_a_tela_diz_que_nao_ha_execucao_sem_inventar_dados(endereco: str, estado: Path) -> None:
    for caminho in ("/", "/cobranca", "/gestao"):
        pagina = _abrir(endereco, caminho)

        assert "Não há execução" in pagina.por_id("sem-execucao").texto
        assert not pagina.tem_id("resumo")
        assert pagina.todos("table") == []
    assert not estado.exists()


def test_armazenamento_sem_execucao_tambem_diz_que_nao_ha_execucao(endereco: str, armazenamento: Armazenamento) -> None:
    pagina = _abrir(endereco)

    assert "Não há execução" in pagina.por_id("sem-execucao").texto
    assert pagina.todos("table") == []


def test_navegacao_no_topo_leva_as_tres_visoes(endereco: str) -> None:
    for caminho in ("/", "/cobranca", "/gestao"):
        navegacao = _abrir(endereco, caminho).por_id("navegacao")

        assert [a.atributos["href"] for a in navegacao.todos("a")] == ["/", "/cobranca", "/gestao"]
        [ativa] = [a for a in navegacao.todos("a") if a.atributos.get("aria-current") == "page"]
        assert ativa.atributos["href"] == caminho


def test_topo_separa_os_dois_produtos_e_cada_um_mostra_so_as_suas_sub_abas(endereco: str, estado: Path) -> None:
    produtos = {
        "Conciliação Assistida": ["/", "/cobranca", "/gestao"],
        "Cobrança Endereçada": ["/enderecada/emissao", "/enderecada/migracao", "/enderecada/demonstracao"],
    }
    for produto, caminhos in produtos.items():
        for caminho in caminhos:
            pagina = _abrir(endereco, caminho)

            links = pagina.por_id("produtos").todos("a")
            assert [(a.texto, a.atributos["href"]) for a in links] == [
                ("Conciliação Assistida", "/"),
                ("Cobrança Endereçada", "/enderecada/emissao"),
            ]
            [ativo] = [a for a in links if a.atributos.get("aria-current") == "page"]
            assert ativo.texto == produto
            assert [a.atributos["href"] for a in pagina.por_id("navegacao").todos("a")] == caminhos
    assert [a.texto for a in _abrir(endereco, "/").por_id("navegacao").todos("a")] == ["Analista", "Cobrança", "Gestão"]
    assert [a.texto for a in _abrir(endereco, "/enderecada/emissao").por_id("navegacao").todos("a")] == [
        "Emissão",
        "Migração",
        "Demonstração",
    ]
    assert not _abrir(endereco, "/").tem_id("demonstracao")
    assert not estado.exists()


def test_resumo_completo_e_itens_da_tela_sao_os_do_workbook_da_mesma_execucao(
    endereco: str, pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _cobranca_e_em_analise(pasta)
    resultado = _executar(pasta, armazenamento)
    workbook_path = tmp_path / "conciliacao.xlsx"
    gerar_workbook(resultado, workbook_path)
    workbook = load_workbook(workbook_path)

    pagina = _abrir(endereco)

    resumo_do_workbook = [
        [_texto_da_celula(c) for c in linha] for linha in workbook["Resumo"].iter_rows(values_only=True)
    ]
    assert pagina.por_id("resumo").linhas() == resumo_do_workbook
    for aba, secao in (("Lista de baixa", "lista-de-baixa"), ("Fila de exceções", "fila-de-excecoes")):
        do_workbook = [linha[0] for linha in workbook[aba].iter_rows(min_row=2, values_only=True)]
        assert _itens_com_decisao(pagina.por_id(secao)) == do_workbook
        assert do_workbook


def test_resumo_simples_conta_o_que_falta_confirmar_e_resolver_e_a_situacao_da_ia(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _cobranca_e_em_analise(pasta)
    _executar(pasta, armazenamento)

    pagina = _abrir(endereco)

    assert pagina.por_id("pra-confirmar").texto.startswith("Pra confirmar 1 R$ 1.000,00")
    assert pagina.por_id("pra-resolver").texto == "Pra resolver 1 pagamento R$ 500,00 em pagamentos"
    ia = pagina.por_id("situacao-da-ia")
    assert ia.texto == "IA não precisou nenhum nome pra analisar"
    assert "numero" not in ia.classes and "aviso" not in ia.classes


def test_a_tela_le_a_ultima_execucao_persistida_sem_recalcular(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _cobranca_e_em_analise(pasta)
    _executar(pasta, armazenamento)
    resultado = _executar(pasta, armazenamento)
    for arquivo in pasta.iterdir():
        arquivo.unlink()

    pagina = _abrir(endereco)

    assert pagina.por_id("resumo").linhas()[0] == ["Execução", str(resultado.resumo.execucao)]
    assert resultado.resumo.execucao == 2
    assert len(pagina.por_id("itens-da-lista").linhas()) == 1 + len(resultado.lista_de_baixa)


def test_lista_de_baixa_mostra_franqueado_titulo_valor_e_por_que_com_a_explicacao_em_detalhes(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _cobranca_e_em_analise(pasta)
    resultado = _executar(pasta, armazenamento)
    [na_lista] = resultado.lista_de_baixa

    pagina = _abrir(endereco)

    tabela = pagina.por_id("itens-da-lista")
    cabecalho, linha = tabela.linhas()
    assert cabecalho == ["Franqueado", "Título", "Crédito", "Valor", "Por quê", "Decisão"]
    franqueado, titulo, credito, valor, por_que, _ = linha
    assert franqueado == na_lista.titulos[0].nome_do_franqueado
    assert titulo == "NF NA LISTA (Protheus)"
    assert credito == "02/06/2025 PIX RECEBIDO 10200300000140"
    assert valor == "R$ 1.000,00"
    assert na_lista.regra is not None and por_que.startswith(POR_QUE_DA_REGRA[na_lista.regra])
    [explicacao] = tabela.todos("dl")
    assert " ".join(na_lista.explicacao.split()) in explicacao.texto
    assert tabela.todos("details")[0].todos("summary")[0].texto == "ver detalhes"


def test_fila_vira_um_cartao_por_item_com_motivo_simples_titulo_sugerido_e_explicacao_em_detalhes(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _cobranca_e_em_analise(pasta)
    resultado = _executar(pasta, armazenamento)
    [na_fila] = resultado.fila_de_excecoes
    assert na_fila.motivo is not None

    cartao = _cartao(_abrir(endereco), na_fila.identificador)

    assert cartao.tag == "article"
    assert "DEPOSITO EM DINHEIRO" in cartao.texto and "R$ 500,00" in cartao.texto
    assert MOTIVO_EM_LINGUAGEM_SIMPLES[na_fila.motivo] in cartao.todos(classe="motivo")[0].texto
    assert "NF NA FILA" in cartao.todos(classe="sugerido")[0].texto
    botoes = {b.atributos.get("value"): b.texto for b in cartao.todos("button") if b.atributos.get("name") == "decisao"}
    assert set(botoes) == {"confirmar", "antecipação", "devolução", "erro", "possível fonte não integrada"}
    [explicacao] = cartao.todos("dl")
    assert " ".join(na_fila.explicacao.split()) in explicacao.texto


def test_parecer_da_ia_aparece_destacado_no_cartao(endereco: str, pasta: Path, armazenamento: Armazenamento) -> None:
    _parecer_de_nome(pasta)
    resultado = _executar(pasta, armazenamento, PareceristaFalso(sim_para=[("CONF ALVORADA COM", NOME_B)]))
    [na_fila] = resultado.fila_de_excecoes

    pagina = _abrir(endereco)

    cartao = _cartao(pagina, na_fila.identificador)
    assert "com-parecer" in cartao.classes
    [parecer] = cartao.todos(classe="parecer")
    assert "Parecer da IA" in parecer.texto and "abrevia" in parecer.texto and "sim" in parecer.texto
    assert cartao.todos(classe="ia-indisponivel") == []
    assert pagina.por_id("situacao-da-ia").texto == "IA funcionando 1 item com parecer"


def test_ia_indisponivel_aparece_marcada_no_cartao_e_no_resumo(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _parecer_de_nome(pasta)
    resultado = _executar(pasta, armazenamento, PareceristaFalso(indisponivel=True))
    [na_fila] = resultado.fila_de_excecoes

    pagina = _abrir(endereco)

    [marcacao] = _cartao(pagina, na_fila.identificador).todos(classe="ia-indisponivel")
    assert marcacao.texto == "IA indisponível"
    alerta = pagina.por_id("situacao-da-ia")
    assert {"aviso", "problema"} <= alerta.classes
    assert alerta.texto == "IA indisponível 1 item ficou sem parecer da IA. Vale chamar o TI."


NOSSO_NUMERO_DUPLICADO = "000077700011"
NOSSO_NUMERO_DA_FILA = "000077700012"


def _duplicidades_na_lista_e_na_fila(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(
                documento="NF 930001", valor="2.400,00", vencimento="10/03/2025", nosso_numero=NOSSO_NUMERO_DUPLICADO
            ),
            TituloProtheus(
                documento="NF 930002", valor="800,00", vencimento="10/03/2025", nosso_numero=NOSSO_NUMERO_DA_FILA
            ),
        ],
    )
    escrever_omie(
        pasta,
        [
            TituloOmie(documento="FAT-9301", valor="2400.00", vencimento="2025-03-10", cnpj="10200300000299"),
            TituloOmie(documento="FAT-9302", valor="800.00", vencimento="2025-03-10", cnpj="10200300000299"),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("10/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "2.400,00", documento=NOSSO_NUMERO_DUPLICADO),
            LinhaItau("11/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "790,00", documento=NOSSO_NUMERO_DA_FILA),
        ],
    )


def _duplicidade_de(resultado: Resultado, documento: str) -> str:
    [item] = [i for i in resultado.itens_de_titulo if documento in {t.documento for t in i.titulos}]
    return item.identificador


def test_titulo_da_lista_em_possivel_duplicidade_avisa_na_linha_qual_e_o_outro_titulo(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _duplicidades_na_lista_e_na_fila(pasta)
    resultado = _executar(pasta, armazenamento)
    [na_lista] = resultado.lista_de_baixa
    assert [t.documento for t in na_lista.titulos] == ["NF 930001"]

    [_, linha] = _abrir(endereco).por_id("itens-da-lista").todos("tr")

    [alerta] = linha.todos(classe="alerta-de-duplicidade")
    assert "NF 930001 (Protheus) pode ser a mesma dívida que FAT-9301 (Omie)" in alerta.texto
    [link] = alerta.todos("a")
    assert link.atributos["href"] == f"#item-{_duplicidade_de(resultado, 'NF 930001')}"


def test_cartao_da_fila_que_sugere_titulo_em_possivel_duplicidade_tambem_avisa(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _duplicidades_na_lista_e_na_fila(pasta)
    resultado = _executar(pasta, armazenamento)
    [na_fila] = resultado.fila_de_excecoes
    assert [t.documento for t in na_fila.titulos] == ["NF 930002"]

    pagina = _abrir(endereco)

    [alerta] = _cartao(pagina, na_fila.identificador).todos(classe="alerta-de-duplicidade")
    assert "NF 930002 (Protheus) pode ser a mesma dívida que FAT-9302 (Omie)" in alerta.texto
    assert alerta.todos("a")[0].atributos["href"] == f"#item-{_duplicidade_de(resultado, 'NF 930002')}"
    for item in resultado.itens_de_titulo:
        assert _cartao(pagina, item.identificador).todos(classe="alerta-de-duplicidade") == []


def test_pra_resolver_separa_pagamentos_de_duplicidades_com_o_valor_so_dos_pagamentos(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _duplicidades_na_lista_e_na_fila(pasta)
    _executar(pasta, armazenamento)

    pagina = _abrir(endereco)

    assert pagina.por_id("pra-resolver").texto == "Pra resolver 1 pagamento + 2 duplicidades R$ 790,00 em pagamentos"


def test_linha_da_lista_com_encargo_mostra_o_titulo_e_a_multa_e_os_juros_separados(
    endereco: str, pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    (tmp_path / "conciliacao.toml").write_text(
        CONFIGURACAO + "[encargo]\njuros_ao_mes_percentual = 1\n", encoding="utf-8"
    )
    escrever_protheus(
        pasta,
        [TituloProtheus(documento="NF 940001", valor="10.000,00", vencimento="10/03/2025", nosso_numero="000077700011")],
    )
    escrever_extrato_itau(
        pasta, [LinhaItau("20/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "10.233,33", documento="000077700011")]
    )
    executar_conciliacao(
        pasta,
        armazenamento,
        configuracao_itau(encargo=ParametrosDeEncargo(juros_ao_mes_percentual=Decimal("1"))),
        PareceristaFalso(),
    )

    [_, linha] = _abrir(endereco).por_id("itens-da-lista").linhas()

    assert linha[3] == "R$ 10.000,00 + multa R$ 200,00 + juros R$ 33,33 crédito R$ 10.233,33"


def test_linha_da_lista_so_com_multa_nao_mostra_juros(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [TituloProtheus(documento="NF 940001", valor="10.000,00", vencimento="10/03/2025", nosso_numero="000077700011")],
    )
    escrever_extrato_itau(
        pasta, [LinhaItau("20/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "10.200,00", documento="000077700011")]
    )
    _executar(pasta, armazenamento)

    [_, linha] = _abrir(endereco).por_id("itens-da-lista").linhas()

    assert linha[3] == "R$ 10.000,00 + multa R$ 200,00 crédito R$ 10.200,00"


def test_cobranca_mostra_quem_cobrar_com_dias_de_atraso_e_separa_os_titulos_em_analise(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _cobranca_e_em_analise(pasta)
    resultado = _executar(pasta, armazenamento)
    [na_fila] = resultado.fila_de_excecoes

    pagina = _abrir(endereco, "/cobranca")

    cabecalho, a_cobrar = pagina.por_id("titulos-a-cobrar").linhas()
    assert cabecalho == ["Franqueado", "Título", "Valor", "Vencimento", "Dias de atraso"]
    assert a_cobrar[1] == "NF SEM CREDITO (Protheus)"
    assert a_cobrar[2:] == ["R$ 730,00", "02/06/2025", "8"]
    cabecalho, em_analise = pagina.por_id("titulos-em-analise").linhas()
    assert cabecalho[-1] == "Por quê"
    assert em_analise[1] == "NF NA FILA (Protheus)"
    assert na_fila.motivo is not None
    assert em_analise[-1] == f"Pagamento de R$ 500,00 em 03/06 em análise ({na_fila.motivo.value})"
    assert pagina.todos("form") == []


def test_gestao_mostra_o_resolvido_sozinho_e_a_causa_raiz_em_grafico(
    endereco: str, pasta: Path, armazenamento: Armazenamento
) -> None:
    _cobranca_e_em_analise(pasta)
    resultado = _executar(pasta, armazenamento)
    causa_raiz = resultado.causa_raiz

    pagina = _abrir(endereco, "/gestao")

    assert pagina.por_id("resolvido-em-quantidade").texto == "Créditos resolvidos sozinhos 50,0% 1 de 2 Créditos"
    svg = pagina.por_id("grafico-causa-raiz")
    assert svg.tag == "svg"
    [com_credito] = [g for g in causa_raiz.por_motivo if g.quantidade]
    [barra] = svg.todos("g", classe="barra")
    textos = [t.texto for t in barra.todos("text")]
    assert textos == [com_credito.nome, f"1 Crédito(s), {formatar_reais(com_credito.valor_centavos)}"]
    assert float(barra.todos("path")[0].atributos["data-comprimento"]) > 0
    por_franqueado = pagina.por_id("causa-raiz-por-franqueado").linhas()
    assert [linha[:2] for linha in por_franqueado[1:]] == [
        [g.nome, str(g.quantidade)] for g in causa_raiz.por_franqueado
    ] + [["Total", str(causa_raiz.quantidade)]]


def test_listas_vazias_aparecem_como_vazias(endereco: str, pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="700,00", vencimento="20/06/2025")])
    escrever_extrato_itau(pasta, [LinhaItau("02/06/2025", "TARIFA BANCARIA", "1,00", tipo="D")])
    _executar(pasta, armazenamento)

    analista = _abrir(endereco)
    cobranca = _abrir(endereco, "/cobranca")

    for pagina, secao in (
        (analista, "lista-de-baixa"),
        (analista, "fila-de-excecoes"),
        (analista, "baixa-manual"),
        (cobranca, "lista-de-cobranca"),
        (cobranca, "em-analise"),
    ):
        assert pagina.por_id(secao).todos("table") == []
        assert pagina.por_id(secao).todos(classe="vazia")


def test_as_paginas_nao_usam_recursos_externos(endereco: str, pasta: Path, armazenamento: Armazenamento) -> None:
    _cobranca_e_em_analise(pasta)
    _executar(pasta, armazenamento)

    for caminho in (
        "/",
        "/cobranca",
        "/gestao",
        "/enderecada/emissao",
        "/enderecada/migracao",
        "/enderecada/demonstracao",
    ):
        html = _html(endereco, caminho)

        pagina = ler_pagina(html)
        assert pagina.todos("script") == [] and pagina.todos("link") == [] and pagina.todos("img") == []
        assert "url(" not in html and "@import" not in html and "src=" not in html


def test_rota_desconhecida_responde_404(endereco: str) -> None:
    with pytest.raises(urllib.error.HTTPError) as erro:
        urllib.request.urlopen(endereco + "/nada")

    assert erro.value.code == 404


def test_comando_sobe_a_tela_local_sobre_o_armazenamento(tmp_path: Path) -> None:
    with socket.socket() as livre:
        livre.bind(("127.0.0.1", 0))
        porta = livre.getsockname()[1]
    processo = subprocess.Popen(
        [
            sys.executable,
            "-m",
            "conciliacao.tela",
            "--estado",
            str(tmp_path / "estado.sqlite"),
            "--entrada",
            str(tmp_path / "entrada"),
            "--porta",
            str(porta),
            "--sem-navegador",
        ],
        stdout=subprocess.PIPE,
        text=True,
    )
    try:
        assert processo.stdout is not None
        linha = processo.stdout.readline()
        assert f"http://127.0.0.1:{porta}/" in linha
        pagina = _abrir(f"http://127.0.0.1:{porta}")
        assert pagina.tem_id("sem-execucao")
    finally:
        processo.terminate()
        processo.wait(timeout=10)


def _executar_e_fechar(pasta: Path, estado: Path, workbook: Path | None = None) -> Resultado:
    armazenamento = Armazenamento(estado)
    try:
        resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso(), workbook)
    finally:
        armazenamento.fechar()
    if workbook is not None:
        gerar_workbook(resultado, workbook)
    return resultado


def _decisoes(estado: Path) -> dict[str, DecisaoRegistrada]:
    armazenamento = Armazenamento(estado)
    try:
        return {decisao.item: decisao for decisao in armazenamento.decisoes()}
    finally:
        armazenamento.fechar()


def _responder(endereco: str, requisicao: urllib.request.Request) -> Elemento:
    with urllib.request.urlopen(requisicao) as resposta:
        assert resposta.status == 200
        assert urllib.parse.urlsplit(resposta.url).path == "/"
        return ler_pagina(resposta.read().decode("utf-8"))


def _postar(endereco: str, caminho: str, campos: dict[str, str] | None = None) -> Elemento:
    dados = urllib.parse.urlencode(campos or {}).encode("utf-8")
    return _responder(endereco, urllib.request.Request(endereco + caminho, data=dados, method="POST"))


def _postar_recusado(endereco: str, caminho: str, campos: dict[str, str]) -> tuple[int, str]:
    dados = urllib.parse.urlencode(campos).encode("utf-8")
    with pytest.raises(urllib.error.HTTPError) as erro:
        urllib.request.urlopen(urllib.request.Request(endereco + caminho, data=dados, method="POST"))
    return erro.value.code, erro.value.read().decode("utf-8")


def _enviar(endereco: str, arquivos: dict[str, bytes]) -> Elemento:
    fronteira = "fronteira-do-teste"
    partes = b"".join(
        f'--{fronteira}\r\nContent-Disposition: form-data; name="arquivos"; filename="{nome}"\r\n'
        "Content-Type: application/octet-stream\r\n\r\n".encode() + conteudo + b"\r\n"
        for nome, conteudo in arquivos.items()
    )
    requisicao = urllib.request.Request(
        endereco + "/enviar",
        data=partes + f"--{fronteira}--\r\n".encode(),
        method="POST",
        headers={"Content-Type": f"multipart/form-data; boundary={fronteira}"},
    )
    return _responder(endereco, requisicao)


def test_rejeitar_com_motivo_pela_tela_grava_pelo_mesmo_caminho_do_excel(
    endereco: str, pasta: Path, estado: Path
) -> None:
    _cobranca_e_em_analise(pasta)
    primeira = _executar_e_fechar(pasta, estado)
    [na_lista] = primeira.lista_de_baixa

    pagina = _postar(
        endereco,
        "/decidir",
        {"aba": "Lista de baixa", "item": na_lista.identificador, "decisao": "rejeitar", "motivo": "não reconheço"},
    )

    [linha] = pagina.por_id("itens-da-lista").linhas()[1:]
    assert "Rejeitado" in linha[-1] and "não reconheço" in linha[-1]
    assert pagina.por_id("pra-confirmar").texto.startswith("Pra confirmar 0")
    segunda = _executar_e_fechar(pasta, estado)
    decisao = _decisoes(estado)[na_lista.identificador]
    assert (decisao.tipo, decisao.motivo, decisao.origem) == (TipoDeDecisao.REJEITAR, "não reconheço", "Tela")
    assert decisao.autor == getpass.getuser()
    assert decisao.titulos == na_lista.titulos and decisao.explicacao == na_lista.explicacao
    assert segunda.lista_de_baixa == []
    assert na_lista.identificador in [item.identificador for item in segunda.fila_de_excecoes]
    assert segunda.resumo.decisoes.rejeitadas == 1


def test_confirmar_todos_pela_tela_confirma_a_lista_de_baixa(endereco: str, pasta: Path, estado: Path) -> None:
    _cobranca_e_em_analise(pasta)
    primeira = _executar_e_fechar(pasta, estado)
    [na_lista] = primeira.lista_de_baixa
    assert "Confirmar todos (1)" in _abrir(endereco).por_id("confirmar-todos").texto

    pagina = _postar(endereco, "/confirmar-lote")

    [linha] = pagina.por_id("itens-da-lista").linhas()[1:]
    assert "Confirmado" in linha[-1]
    assert "disabled" in pagina.por_id("confirmar-todos").todos("button")[0].atributos
    segunda = _executar_e_fechar(pasta, estado)
    decisao = _decisoes(estado)[na_lista.identificador]
    assert (decisao.tipo, decisao.origem) == (TipoDeDecisao.CONFIRMAR, "Tela")
    assert segunda.lista_de_baixa == []
    assert [baixa.titulo.documento for baixa in segunda.baixa_manual] == ["NF NA LISTA"]
    assert segunda.resumo.decisoes.confirmadas == 1


def test_resolver_item_da_fila_pela_tela(endereco: str, pasta: Path, estado: Path) -> None:
    _cobranca_e_em_analise(pasta)
    primeira = _executar_e_fechar(pasta, estado)
    [na_fila] = primeira.fila_de_excecoes

    pagina = _postar(
        endereco,
        "/decidir",
        {"aba": "Fila de exceções", "item": na_fila.identificador, "decisao": "confirmar"},
    )

    cartao = _cartao(pagina, na_fila.identificador)
    assert "decidido" in cartao.classes and "Confirmado" in cartao.todos(classe="decisao-registrada")[0].texto
    segunda = _executar_e_fechar(pasta, estado)
    decisao = _decisoes(estado)[na_fila.identificador]
    assert (decisao.tipo, decisao.origem) == (TipoDeDecisao.CONFIRMAR, "Tela")
    assert [titulo.documento for titulo in decisao.titulos] == ["NF NA FILA"]
    assert segunda.fila_de_excecoes == []


def test_a_tela_recusa_decisao_invalida_sem_gravar(endereco: str, pasta: Path, estado: Path) -> None:
    _cobranca_e_em_analise(pasta)
    primeira = _executar_e_fechar(pasta, estado)
    [na_lista] = primeira.lista_de_baixa

    status, corpo = _postar_recusado(
        endereco, "/decidir", {"aba": "Lista de baixa", "item": na_lista.identificador, "decisao": "rejeitar"}
    )
    outro_status, _ = _postar_recusado(
        endereco, "/decidir", {"aba": "Lista de baixa", "item": "não existe", "decisao": "confirmar"}
    )

    assert (status, outro_status) == (400, 400)
    assert "motivo" in ler_pagina(corpo).por_id("recusa").texto
    _executar_e_fechar(pasta, estado)
    assert _decisoes(estado) == {}


def test_sem_execucao_a_tela_nao_grava_decisao(endereco: str, estado: Path) -> None:
    status, corpo = _postar_recusado(endereco, "/confirmar-lote", {})

    assert status == 409
    assert "Não há execução" in ler_pagina(corpo).por_id("recusa").texto
    assert not estado.exists()


def test_conflito_entre_excel_e_tela_aparece_na_tela_com_os_mesmos_totais_do_workbook(
    endereco: str, pasta: Path, estado: Path, tmp_path: Path
) -> None:
    _cobranca_e_em_analise(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar_e_fechar(pasta, estado, workbook)
    [na_lista] = primeira.lista_de_baixa
    livro = load_workbook(workbook)
    livro.properties.lastModifiedBy = "Ana Analista"
    livro.save(workbook)
    _postar(endereco, "/confirmar-lote")

    _executar_e_fechar(pasta, estado, workbook)

    pagina = _abrir(endereco)
    [conflito] = pagina.por_id("conflitos").todos("li")
    assert na_lista.identificador in conflito.texto
    resumo_do_workbook = [
        [_texto_da_celula(c) for c in linha] for linha in load_workbook(workbook)["Resumo"].iter_rows(values_only=True)
    ]
    assert pagina.por_id("resumo").linhas() == resumo_do_workbook
    assert ["Conflitos entre Excel e tela", "1"] in resumo_do_workbook
    assert na_lista.identificador not in _decisoes(estado)


def test_classificar_item_da_fila_como_possivel_fonte_nao_integrada_pela_tela(
    endereco: str, pasta: Path, estado: Path
) -> None:
    _cobranca_e_em_analise(pasta)
    primeira = _executar_e_fechar(pasta, estado)
    [na_fila] = primeira.fila_de_excecoes
    assert "Fonte não integrada" in _cartao(_abrir(endereco), na_fila.identificador).texto

    _postar(
        endereco,
        "/decidir",
        {"aba": "Fila de exceções", "item": na_fila.identificador, "decisao": "possível fonte não integrada"},
    )

    segunda = _executar_e_fechar(pasta, estado)
    decisao = _decisoes(estado)[na_fila.identificador]
    assert (decisao.tipo, decisao.origem) == (TipoDeDecisao.POSSIVEL_FONTE_NAO_INTEGRADA, "Tela")
    [classificado] = segunda.fila_de_excecoes
    assert classificado.lancamento == na_fila.lancamento
    assert classificado.motivo is MotivoDeNaoIdentificacao.POSSIVEL_FONTE_NAO_INTEGRADA
    assert segunda.resumo.decisoes.fontes_nao_integradas == 1


def _arquivos_de_exemplo(tmp_path: Path) -> dict[str, bytes]:
    origem = tmp_path / "origem"
    origem.mkdir()
    _cobranca_e_em_analise(origem)
    (origem / "contas_a_pagar.txt").write_text("nada a ver\n", encoding="utf-8")
    return {arquivo.name: arquivo.read_bytes() for arquivo in sorted(origem.iterdir())}


def test_enviar_arquivos_e_conciliar_pela_tela_roda_a_execucao_do_comando_e_gera_o_excel(
    endereco: str, pasta: Path, estado: Path, saida: Path, tmp_path: Path
) -> None:
    arquivos = _arquivos_de_exemplo(tmp_path)

    pagina = _enviar(endereco, arquivos)

    assert sorted(p.name for p in pasta.iterdir()) == sorted(arquivos)
    linhas = pagina.por_id("arquivos-na-entrada").linhas()[1:]
    assert [linha[0] for linha in linhas] == sorted(arquivos)
    assert all("Novo" in linha[1] for linha in linhas)
    assert not estado.exists()

    pagina = _postar(endereco, "/conciliar")

    situacoes = {linha[0].split(" ")[0]: linha for linha in pagina.por_id("arquivos-na-entrada").linhas()[1:]}
    recusado = situacoes["contas_a_pagar.txt"]
    assert "Não reconhecido" in recusado[1] and "cabeçalho não corresponde" in recusado[0]
    assert "Lido na última conciliação" in situacoes["itau_cobranca.csv"][1]
    assert pagina.por_id("pra-confirmar").texto.startswith("Pra confirmar 1")
    assert pagina.por_id("resumo").linhas()[0] == ["Execução", "1"]
    assert load_workbook(saida)["Resumo"]["B1"].value == 1
    with urllib.request.urlopen(endereco + "/excel") as resposta:
        assert resposta.headers["Content-Disposition"] == 'attachment; filename="conciliacao.xlsx"'
        assert resposta.read() == saida.read_bytes()


def test_conciliar_pela_tela_aplica_as_decisoes_da_tela_e_lista_o_que_baixar_no_erp(
    endereco: str, pasta: Path, estado: Path, tmp_path: Path
) -> None:
    _enviar(endereco, _arquivos_de_exemplo(tmp_path))
    _postar(endereco, "/conciliar")
    _postar(endereco, "/confirmar-lote")
    assert "Aplica as 1 decisão(ões)" in _abrir(endereco).por_id("conciliar").texto

    pagina = _postar(endereco, "/conciliar")

    assert pagina.por_id("resumo").linhas()[0] == ["Execução", "2"]
    [baixa] = pagina.por_id("itens-da-baixa-manual").linhas()[1:]
    assert baixa[:2] == ["NF NA LISTA", "Protheus"] and baixa[3] == "R$ 1.000,00"
    assert pagina.por_id("pra-baixar").texto.startswith("Pra baixar no ERP 1")
    assert pagina.tem_id("baixar-excel")


def test_remover_arquivo_enviado_tira_ele_da_proxima_conciliacao(
    endereco: str, pasta: Path, tmp_path: Path
) -> None:
    _enviar(endereco, _arquivos_de_exemplo(tmp_path))

    pagina = _postar(endereco, "/remover-arquivo", {"nome": "contas_a_pagar.txt"})

    assert "contas_a_pagar.txt" not in [linha[0] for linha in pagina.por_id("arquivos-na-entrada").linhas()]
    assert not (pasta / "contas_a_pagar.txt").exists()
    status, _ = _postar_recusado(endereco, "/remover-arquivo", {"nome": "../estado.sqlite"})
    assert status == 400


def test_sem_arquivos_nao_concilia_e_nao_ha_excel_para_baixar(endereco: str, estado: Path) -> None:
    pagina = _abrir(endereco)
    assert "disabled" in pagina.por_id("conciliar").todos("button")[0].atributos
    assert not pagina.tem_id("baixar-excel")

    status, corpo = _postar_recusado(endereco, "/conciliar", {})

    assert status == 400 and "Envie" in ler_pagina(corpo).por_id("recusa").texto
    assert not estado.exists()
    with pytest.raises(urllib.error.HTTPError) as erro:
        urllib.request.urlopen(endereco + "/excel")
    assert erro.value.code == 404


def test_arquivo_com_nome_de_caminho_e_gravado_so_pelo_nome(endereco: str, pasta: Path, tmp_path: Path) -> None:
    _enviar(endereco, {"../../fora.csv": b"x"})

    assert [p.name for p in pasta.iterdir()] == ["fora.csv"]
    assert not (tmp_path / "fora.csv").exists()
