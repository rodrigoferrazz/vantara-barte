from pathlib import Path

from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.dominio import Item, MotivoDeNaoIdentificacao, TipoDeEvidencia
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import (
    LinhaItau,
    TituloOmie,
    TituloProtheus,
    configuracao_itau,
    escrever_extrato_itau,
    escrever_omie,
    escrever_protheus,
)
from tests.parecerista import PareceristaFalso

CNPJ_A = "10.200.300/0001-40"
CNPJ_B = "20.400.600/0001-80"


def _unico_da_fila(pasta: Path, armazenamento: Armazenamento) -> Item:
    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    assert resultado.lista_de_baixa == []
    [item] = resultado.fila_de_excecoes
    return item


def test_nome_truncado_que_e_prefixo_de_nome_do_franqueado_vai_para_a_fila_com_sugestao(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="3.200,00", cliente="Panificadora Aurora Ltda"),
            TituloProtheus(documento="NF 2", valor="900,00", cliente="Panificadora Aurora Ltda"),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO PANIFICADORA AURORA LT", "3.200,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert [t.documento for t in item.titulos] == ["NF 1"]
    assert item.evidencia == (TipoDeEvidencia.NOME_POR_PREFIXO,)
    assert "Panificadora Aurora Ltda" in item.explicacao
    assert "Pagador conhecido" in item.o_que_falta


def test_prefixo_vale_para_nome_do_franqueado_em_qualquer_sistema_de_origem(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="750,00", cliente="Aurora Panificadora SA")])
    escrever_omie(pasta, [TituloOmie(documento="FAT-1", valor="10.00", cliente="PANIFICADORA AURORA S/A")])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "TED RECEBIDA PANIFICADORA AURORA S", "750,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert [(t.documento, t.sistema_de_origem) for t in item.titulos] == [("NF 1", "Protheus")]


def test_prefixo_ignora_acento_caixa_e_simbolos(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="410,00", cliente="Pão & Café Comércio Ltda.")])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO PAO CAFE COMERC", "410,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert [t.documento for t in item.titulos] == ["NF 1"]


def test_nome_que_nao_e_prefixo_nao_conta_como_nome_por_prefixo(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="410,00", cliente="Panificadora Aurora Ltda")])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO AURORA PANIFICADORA", "410,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is not MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert TipoDeEvidencia.NOME_POR_PREFIXO not in item.evidencia


def test_prefixo_de_mais_de_um_franqueado_sem_desempate_pelo_valor_e_nome_ambiguo(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="500,00", cliente="Padaria Sol Nascente Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="500,00", cliente="Padaria Sol Poente ME", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO PADARIA SOL", "500,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.NOME_AMBIGUO
    assert item.titulos == ()
    assert "NF 1" in item.explicacao
    assert "NF 2" in item.explicacao
    assert item.evidencia == (TipoDeEvidencia.NOME_POR_PREFIXO,)


def test_prefixo_de_mais_de_um_franqueado_desempatado_pelo_valor_sugere_o_titulo(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="500,00", cliente="Padaria Sol Nascente Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="800,00", cliente="Padaria Sol Poente ME", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO PADARIA SOL", "800,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ
    assert [t.documento for t in item.titulos] == ["NF 2"]


def test_credito_sem_nome_nem_documento_que_so_bate_por_valor_sugere_o_titulo(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="5.670,00", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="1.000,00", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "5.670,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR
    assert [t.documento for t in item.titulos] == ["NF 1"]
    assert item.evidencia == (TipoDeEvidencia.SO_VALOR,)
    assert item.o_que_falta


def test_credito_sem_nome_com_valor_de_mais_de_um_titulo_nao_sugere(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="5.670,00", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="5.670,00", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "5.670,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.SEM_IDENTIFICACAO_DO_PAGADOR
    assert item.titulos == ()
    assert "NF 1" in item.explicacao
    assert "NF 2" in item.explicacao


def test_valor_nao_fecha_precede_pagador_provavel(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 1", valor="15.600,00", cliente="Cantina Serrana Ltda")])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO CANTINA SERRANA L", "14.820,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.VALOR_NAO_FECHA
    assert [t.documento for t in item.titulos] == ["NF 1"]
    assert "Diferença a menor" in item.explicacao


def test_valor_nao_fecha_precede_nome_ambiguo(pasta: Path, armazenamento: Armazenamento) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="500,00", cliente="Padaria Sol Nascente Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="800,00", cliente="Padaria Sol Poente ME", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO PADARIA SOL", "650,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.VALOR_NAO_FECHA
    assert item.titulos == ()


def test_nome_por_prefixo_de_franqueado_com_varios_titulos_que_nao_fecham_e_credito_sem_titulo_em_aberto(
    pasta: Path, armazenamento: Armazenamento
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="12.300,00", cliente="Padaria Lua Cheia Ltda"),
            TituloProtheus(documento="NF 2", valor="2.700,00", cliente="Padaria Lua Cheia Ltda"),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", "PIX RECEBIDO PADARIA LUA CHE", "1.150,00")])

    item = _unico_da_fila(pasta, armazenamento)

    assert item.motivo is MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO
    assert item.titulos == ()


def test_credito_sem_titulo_em_aberto_precede_nome_ambiguo_e_pagador_provavel(
    tmp_path: Path, armazenamento: Armazenamento
) -> None:
    julho = tmp_path / "julho"
    julho.mkdir()
    escrever_protheus(
        julho,
        [
            TituloProtheus(documento="NF 1", valor="500,00", cliente="Padaria Sol Nascente Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="800,00", cliente="Padaria Sol Poente ME", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(julho, [])
    executar_conciliacao(julho, armazenamento, configuracao_itau(), PareceristaFalso())
    agosto = tmp_path / "agosto"
    agosto.mkdir()
    escrever_protheus(agosto, [TituloProtheus(documento="NF 9", valor="1.850,00", cnpj="30.600.900/0001-10")])
    escrever_extrato_itau(
        agosto,
        [
            LinhaItau("05/08/2025", "PIX RECEBIDO PADARIA SOL", "1.850,00"),
            LinhaItau("06/08/2025", "PIX RECEBIDO PADARIA SOL NAS", "1.850,00"),
        ],
    )

    resultado = executar_conciliacao(agosto, armazenamento, configuracao_itau(), PareceristaFalso())

    assert resultado.lista_de_baixa == []
    assert [item.motivo for item in resultado.fila_de_excecoes] == [
        MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO,
        MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO,
    ]
    assert all(item.titulos == () for item in resultado.fila_de_excecoes)


def test_cada_item_da_fila_tem_um_motivo_titulo_sugerido_evidencia_e_o_que_falta_no_workbook(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="3.200,00", cliente="Panificadora Aurora Ltda", cnpj=CNPJ_A),
            TituloProtheus(documento="NF 2", valor="5.670,00", nosso_numero="0099", cnpj=CNPJ_B),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "PIX RECEBIDO PANIFICADORA AURORA LT", "3.200,00"),
            LinhaItau("05/03/2025", "DEPOSITO NAO IDENTIFICADO", "5.670,00"),
            LinhaItau("05/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "5.000,00", documento="0099"),
            LinhaItau("05/03/2025", "PIX RECEBIDO FULANO DE TAL", "12,34"),
        ],
    )

    resultado = executar_conciliacao(pasta, armazenamento, configuracao_itau(), PareceristaFalso())
    caminho = tmp_path / "saida.xlsx"
    gerar_workbook(resultado, caminho)

    assert len(resultado.fila_de_excecoes) == 4
    assert all(isinstance(item.motivo, MotivoDeNaoIdentificacao) for item in resultado.fila_de_excecoes)
    assert all(item.o_que_falta for item in resultado.fila_de_excecoes)
    linhas = list(load_workbook(caminho)["Fila de exceções"].iter_rows(values_only=True))
    fila = [dict(zip(linhas[0], linha)) for linha in linhas[1:]]
    por_historico = {linha["Histórico"]: linha for linha in fila}
    prefixo = por_historico["PIX RECEBIDO PANIFICADORA AURORA LT"]
    assert prefixo["Motivo"] == MotivoDeNaoIdentificacao.PAGADOR_PROVAVEL_PRIMEIRA_VEZ.value
    assert prefixo["Títulos sugeridos"] == "NF 1 (Protheus)"
    assert prefixo["Evidência"] == TipoDeEvidencia.NOME_POR_PREFIXO.value
    assert prefixo["O que falta"]
    so_valor = por_historico["DEPOSITO NAO IDENTIFICADO"]
    assert so_valor["Evidência"] == TipoDeEvidencia.SO_VALOR.value
    assert so_valor["Títulos sugeridos"] == "NF 2 (Protheus)"
    assert por_historico["LIQUIDACAO COBRANCA SIMPLES"]["Evidência"] == TipoDeEvidencia.NOSSO_NUMERO.value
    assert all(linha["Motivo"] and linha["O que falta"] for linha in fila)
