from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from conciliacao.armazenamento import Armazenamento
from conciliacao.cli import main
from conciliacao.dominio import MotivoDeNaoIdentificacao, Resultado, TipoDeDecisao, TipoDeEvidencia
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import gerar_workbook
from tests.arquivos import LinhaItau, TituloProtheus, configuracao_itau, escrever_extrato_itau, escrever_protheus
from tests.parecerista import PareceristaFalso

CNPJ = "10.200.300/0001-40"


def _executar(pasta: Path, armazenamento: Armazenamento, workbook: Path, ler_decisoes: bool = True) -> Resultado:
    resultado = executar_conciliacao(
        pasta,
        armazenamento,
        configuracao_itau(),
        PareceristaFalso(),
        workbook_anterior=workbook if ler_decisoes else None,
    )
    gerar_workbook(resultado, workbook)
    return resultado


def _linhas(aba: Worksheet) -> list[dict[str, object]]:
    cabecalho, *linhas = aba.iter_rows(values_only=True)
    return [dict(zip([str(c) for c in cabecalho], linha)) for linha in linhas]


def _resumo(workbook: Path) -> dict[object, object]:
    return {linha[0]: linha[1] for linha in load_workbook(workbook)["Resumo"].iter_rows(values_only=True)}


def _decidir(
    workbook: Path, decisoes: dict[str, tuple[str | None, str | None]] | None = None, autor: str | None = "Ana Analista"
) -> None:
    livro = load_workbook(workbook)
    aba = livro["Lista de baixa"]
    cabecalho = [str(c.value) for c in aba[1]]
    coluna_item = cabecalho.index("Item") + 1
    coluna_decisao = cabecalho.index("Decisão") + 1
    coluna_motivo = cabecalho.index("Motivo da rejeição") + 1
    for linha in range(2, aba.max_row + 1):
        item = aba.cell(row=linha, column=coluna_item).value
        if decisoes and item in decisoes:
            decisao, motivo = decisoes[str(item)]
            aba.cell(row=linha, column=coluna_decisao).value = decisao
            aba.cell(row=linha, column=coluna_motivo).value = motivo
    livro.properties.lastModifiedBy = autor
    livro.save(workbook)


def _dois_creditos_na_lista(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 970001", valor="1.000,00", nosso_numero="000097000011"),
            TituloProtheus(documento="NF 970002", valor="2.500,00", nosso_numero="000097000022"),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("10/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "1.020,00", documento="000097000011"),
            LinhaItau("05/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "2.500,00", documento="000097000022"),
        ],
    )


def test_lista_de_baixa_sai_no_workbook_com_identificador_e_confirmar_pre_preenchido(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"

    primeira = _executar(pasta, armazenamento, workbook)
    segunda = _executar(pasta, armazenamento, workbook, ler_decisoes=False)

    lista = _linhas(load_workbook(workbook)["Lista de baixa"])
    assert [linha["Item"] for linha in lista] == [item.identificador for item in primeira.lista_de_baixa]
    assert [item.identificador for item in segunda.lista_de_baixa] == [item.identificador for item in primeira.lista_de_baixa]
    assert [linha["Decisão"] for linha in lista] == ["confirmar", "confirmar"]
    assert [linha["Motivo da rejeição"] for linha in lista] == [None, None]


def test_confirmacao_gera_lista_para_baixa_manual_e_tira_o_credito_da_lista(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    _executar(pasta, armazenamento, workbook)
    _decidir(workbook)

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.lista_de_baixa == []
    assert resultado.fila_de_excecoes == []
    assert [(b.titulo.documento, b.principal_centavos, b.encargo_centavos) for b in resultado.baixa_manual] == [
        ("NF 970002", 250000, 0),
        ("NF 970001", 100000, 2000),
    ]
    baixa = _linhas(load_workbook(workbook)["Baixa manual no ERP"])
    assert [
        (linha["Título"], linha["Sistema de origem"], linha["Valor principal"], linha["Encargo"]) for linha in baixa
    ] == [("NF 970002", "Protheus", 2500, 0), ("NF 970001", "Protheus", 1000, 20)]
    assert _resumo(workbook)["Decisões lidas do Excel"] == "2 confirmada(s), 0 rejeitada(s), 0 não aplicada(s)"


def test_titulo_confirmado_nao_e_proposto_para_outro_credito(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(pasta, [TituloProtheus(documento="NF 970101", valor="700,00", cnpj=CNPJ)])
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", f"TED RECEBIDA CNPJ {CNPJ}", "700,00")])
    workbook = tmp_path / "conciliacao.xlsx"
    _executar(pasta, armazenamento, workbook)
    _decidir(workbook)
    escrever_extrato_itau(pasta, [LinhaItau("06/03/2025", f"TED RECEBIDA CNPJ {CNPJ}", "700,00")], nome="itau_cobranca_2.csv")

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.lista_de_baixa == []
    [item] = resultado.fila_de_excecoes
    assert item.lancamento.data.day == 6
    assert item.motivo is MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO
    assert resultado.cobranca == ()


def test_rejeicao_com_motivo_leva_o_credito_para_a_fila_e_o_par_nunca_volta(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    rejeitado = next(i for i in primeira.lista_de_baixa if i.titulos[0].documento == "NF 970001").identificador
    _decidir(workbook, {rejeitado: (None, "boleto pago em duplicidade pelo Franqueado")})

    segunda = _executar(pasta, armazenamento, workbook)
    _decidir(workbook)
    terceira = _executar(pasta, armazenamento, workbook)

    for resultado in (segunda, terceira):
        assert resultado.lista_de_baixa == []
        [item] = resultado.fila_de_excecoes
        assert item.identificador == rejeitado
        assert item.titulos == ()
        assert "rejeitada pela Analista" in item.explicacao
        assert "boleto pago em duplicidade pelo Franqueado" in item.explicacao
    assert [b.titulo.documento for b in segunda.baixa_manual] == ["NF 970002"]
    assert segunda.resumo.decisoes.rejeitadas == 1
    assert terceira.baixa_manual == ()


def test_rejeicao_escrita_tambem_vale_e_desmarcar_e_rejeitar(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    [um, outro] = [item.identificador for item in primeira.lista_de_baixa]
    _decidir(workbook, {um: ("Rejeitar", "não é deste Título"), outro: ("", "valor errado")})

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.resumo.decisoes.rejeitadas == 2
    assert sorted(item.identificador for item in resultado.fila_de_excecoes) == sorted([um, outro])


def test_credito_rejeitado_vai_para_a_fila_mesmo_que_outro_titulo_feche(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 970201", valor="900,00", nosso_numero="000097020011", cnpj=CNPJ),
            TituloProtheus(documento="NF 970202", valor="900,00", cnpj=CNPJ),
        ],
    )
    escrever_extrato_itau(pasta, [LinhaItau("05/03/2025", f"TED CNPJ {CNPJ}", "900,00", documento="000097020011")])
    workbook = tmp_path / "conciliacao.xlsx"
    [proposta] = _executar(pasta, armazenamento, workbook).lista_de_baixa
    assert [t.documento for t in proposta.titulos] == ["NF 970201"]
    _decidir(workbook, {proposta.identificador: (None, "Título errado")})

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.lista_de_baixa == []
    [item] = resultado.fila_de_excecoes
    assert [t.documento for t in item.titulos] == ["NF 970202"]
    assert item.motivo is MotivoDeNaoIdentificacao.MAIS_DE_UMA_COMBINACAO
    assert "Título errado" in item.explicacao


def test_rejeicao_sem_motivo_e_sinalizada_no_resumo_e_nao_e_aplicada(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    sem_motivo = primeira.lista_de_baixa[0].identificador
    _decidir(workbook, {sem_motivo: (None, None)})

    resultado = _executar(pasta, armazenamento, workbook)

    assert [item.identificador for item in resultado.lista_de_baixa] == [sem_motivo]
    assert resultado.resumo.decisoes.confirmadas == 1
    assert resultado.resumo.decisoes.rejeitadas == 0
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert sem_motivo in aviso and "sem motivo" in aviso
    resumo = _resumo(workbook)
    assert resumo["Decisões lidas do Excel"] == "1 confirmada(s), 0 rejeitada(s), 1 não aplicada(s)"
    assert any(sem_motivo in str(valor) and "sem motivo" in str(valor) for valor in resumo.values())
    assert [d.item for d in armazenamento.decisoes()] == [
        i.identificador for i in primeira.lista_de_baixa if i.identificador != sem_motivo
    ]


def test_decisao_desconhecida_nao_e_aplicada(pasta: Path, armazenamento: Armazenamento, tmp_path: Path) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    duvidoso = primeira.lista_de_baixa[0].identificador
    _decidir(workbook, {duvidoso: ("talvez", None)})

    resultado = _executar(pasta, armazenamento, workbook)

    assert [item.identificador for item in resultado.lista_de_baixa] == [duvidoso]
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert duvidoso in aviso and "talvez" in aviso


def test_toda_decisao_registra_quem_quando_evidencia_e_regra(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    [confirmado, rejeitado] = primeira.lista_de_baixa
    _decidir(workbook, {rejeitado.identificador: (None, "não reconheço")}, autor="Beatriz")
    salvo_em = load_workbook(workbook).properties.modified

    _executar(pasta, armazenamento, workbook)

    decisoes = {d.item: d for d in armazenamento.decisoes()}
    assert decisoes[confirmado.identificador].tipo is TipoDeDecisao.CONFIRMAR
    assert decisoes[rejeitado.identificador].tipo is TipoDeDecisao.REJEITAR
    assert decisoes[rejeitado.identificador].motivo == "não reconheço"
    for item in (confirmado, rejeitado):
        decisao = decisoes[item.identificador]
        assert decisao.autor == "Beatriz"
        assert decisao.momento == salvo_em
        assert isinstance(decisao.momento, datetime)
        assert decisao.origem == "Excel"
        assert decisao.execucao == primeira.resumo.execucao
        assert decisao.evidencia == (TipoDeEvidencia.NOSSO_NUMERO,)
        assert decisao.explicacao == item.explicacao
        assert decisao.regra is not None and decisao.regra == item.regra
        assert decisao.titulos == item.titulos


def test_workbook_nao_salvo_pela_analista_nao_aplica_nada(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    _executar(pasta, armazenamento, workbook)
    _decidir(workbook, autor=None)

    resultado = _executar(pasta, armazenamento, workbook)

    assert len(resultado.lista_de_baixa) == 2
    assert armazenamento.decisoes() == ()
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert "não foi salvo" in aviso


def test_workbook_de_execucao_anterior_a_ultima_nao_aplica_nada(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    antigo = tmp_path / "antigo.xlsx"
    atual = tmp_path / "conciliacao.xlsx"
    _executar(pasta, armazenamento, antigo)
    _executar(pasta, armazenamento, atual, ler_decisoes=False)
    _decidir(antigo)

    resultado = _executar(pasta, armazenamento, antigo)

    assert len(resultado.lista_de_baixa) == 2
    assert armazenamento.decisoes() == ()
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert "execução 1" in aviso and "execução 2" in aviso


def test_item_que_nao_esta_na_lista_da_execucao_e_ignorado_com_aviso(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_creditos_na_lista(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    _executar(pasta, armazenamento, workbook)
    livro = load_workbook(workbook)
    livro["Lista de baixa"].append(("L999",) + (None,) * 8 + ("confirmar", None))
    livro.properties.lastModifiedBy = "Ana Analista"
    livro.save(workbook)

    resultado = _executar(pasta, armazenamento, workbook)

    assert resultado.resumo.decisoes.confirmadas == 2
    [aviso] = resultado.resumo.decisoes.nao_aplicadas
    assert "L999" in aviso


def test_comando_le_as_decisoes_do_workbook_da_execucao_anterior(tmp_path: Path) -> None:
    pasta = tmp_path / "entrada"
    pasta.mkdir()
    _dois_creditos_na_lista(pasta)
    configuracao = tmp_path / "conciliacao.toml"
    configuracao.write_text('[[contas]]\npadrao = "itau_cobranca*.csv"\nbanco = "itau"\nconta = "1234-5"\n', encoding="utf-8")
    saida = tmp_path / "conciliacao.xlsx"
    argumentos = [str(pasta), "--configuracao", str(configuracao), "--estado", str(tmp_path / "e.sqlite"), "--saida", str(saida)]

    assert main(argumentos) == 0
    _decidir(saida)
    assert main(argumentos) == 0

    livro = load_workbook(saida)
    assert _linhas(livro["Lista de baixa"]) == []
    assert sorted(str(linha["Título"]) for linha in _linhas(livro["Baixa manual no ERP"])) == ["NF 970001", "NF 970002"]
    assert livro.sheetnames.index("Baixa manual no ERP") == livro.sheetnames.index("Lista de baixa") + 1
