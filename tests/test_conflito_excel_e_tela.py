from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from conciliacao.armazenamento import Armazenamento
from conciliacao.decisoes import confirmar_o_lote_pela_tela, registrar_decisao_da_tela
from conciliacao.dominio import LinhaDeDecisao, Resultado, TipoDeDecisao
from conciliacao.execucao import executar_conciliacao
from conciliacao.saidas import ABA_DA_FILA_DE_EXCECOES, ABA_DA_LISTA_DE_BAIXA, gerar_workbook
from tests.arquivos import LinhaItau, TituloProtheus, configuracao_itau, escrever_extrato_itau, escrever_protheus
from tests.parecerista import PareceristaFalso

NA_TELA = datetime(2025, 3, 12, 9, 30)


def _executar(pasta: Path, armazenamento: Armazenamento, workbook: Path, ler_decisoes: bool = True) -> Resultado:
    resultado = executar_conciliacao(
        pasta,
        armazenamento,
        configuracao_itau(),
        PareceristaFalso(),
        workbook_anterior=workbook if ler_decisoes else None,
    )
    gerar_workbook(resultado, workbook)
    return resultado


def _salvar_pela_analista(workbook: Path, fila: dict[str, str] | None = None) -> None:
    livro = load_workbook(workbook)
    aba = livro[ABA_DA_FILA_DE_EXCECOES]
    cabecalho = [str(c.value) for c in aba[1]]
    for linha in range(2, aba.max_row + 1):
        item = str(aba.cell(row=linha, column=cabecalho.index("Item") + 1).value)
        if fila and item in fila:
            aba.cell(row=linha, column=cabecalho.index("Decisão") + 1).value = fila[item]
    livro.properties.lastModifiedBy = "Ana Analista"
    livro.save(workbook)


def _resumo(workbook: Path) -> list[tuple[object, object]]:
    return [(linha[0], linha[1]) for linha in load_workbook(workbook)["Resumo"].iter_rows(values_only=True)]


def _dois_na_lista_e_um_na_fila(pasta: Path) -> None:
    escrever_protheus(
        pasta,
        [
            TituloProtheus(documento="NF 1", valor="1.000,00", nosso_numero="000097000011"),
            TituloProtheus(documento="NF 2", valor="2.500,00", nosso_numero="000097000022"),
        ],
    )
    escrever_extrato_itau(
        pasta,
        [
            LinhaItau("05/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "1.000,00", documento="000097000011"),
            LinhaItau("05/03/2025", "LIQUIDACAO COBRANCA SIMPLES", "2.500,00", documento="000097000022"),
            LinhaItau("06/03/2025", "DEPOSITO EM DINHEIRO", "77,00"),
        ],
    )


def test_item_decidido_no_excel_e_na_tela_com_a_mesma_decisao_e_conflito_e_nada_e_aplicado(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_na_lista_e_um_na_fila(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    disputado, so_no_excel = primeira.lista_de_baixa
    _salvar_pela_analista(workbook)
    registrar_decisao_da_tela(
        armazenamento, LinhaDeDecisao(ABA_DA_LISTA_DE_BAIXA, disputado.identificador, "confirmar"), "rodrigo", NA_TELA
    )

    segunda = _executar(pasta, armazenamento, workbook)

    assert [item.identificador for item in segunda.lista_de_baixa] == [disputado.identificador]
    decididos = {d.item: d for d in armazenamento.decisoes()}
    assert disputado.identificador not in decididos
    assert decididos[so_no_excel.identificador].origem == "Excel"
    [conflito] = segunda.resumo.decisoes.conflitos
    assert disputado.identificador in conflito and "Excel" in conflito and "tela" in conflito
    resumo = _resumo(workbook)
    assert ("Conflitos entre Excel e tela", 1) in resumo
    assert ("Conflito", conflito) in resumo


def test_decisoes_diferentes_no_excel_e_na_tela_tambem_sao_conflito(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_na_lista_e_um_na_fila(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    [na_fila] = primeira.fila_de_excecoes
    _salvar_pela_analista(workbook, {na_fila.identificador: "erro"})
    registrar_decisao_da_tela(
        armazenamento, LinhaDeDecisao(ABA_DA_FILA_DE_EXCECOES, na_fila.identificador, "devolução"), "rodrigo", NA_TELA
    )

    segunda = _executar(pasta, armazenamento, workbook)

    assert [item.identificador for item in segunda.fila_de_excecoes] == [na_fila.identificador]
    assert na_fila.identificador not in {d.item for d in armazenamento.decisoes()}
    assert len(segunda.resumo.decisoes.conflitos) == 1
    assert segunda.resumo.decisoes.confirmadas == 2


def test_decisao_da_tela_e_aplicada_pelo_mesmo_caminho_do_excel_uma_unica_vez(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_na_lista_e_um_na_fila(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    rejeitado, confirmado = primeira.lista_de_baixa
    [na_fila] = primeira.fila_de_excecoes
    registrar_decisao_da_tela(
        armazenamento,
        LinhaDeDecisao(ABA_DA_LISTA_DE_BAIXA, rejeitado.identificador, "rejeitar", motivo="não reconheço"),
        "rodrigo",
        NA_TELA,
    )
    registrar_decisao_da_tela(
        armazenamento,
        LinhaDeDecisao(
            ABA_DA_FILA_DE_EXCECOES,
            na_fila.identificador,
            "conciliado fora",
            franqueado="10.200.300/0001-40",
            referencia="acordo 12",
        ),
        "rodrigo",
        NA_TELA,
    )
    confirmar_o_lote_pela_tela(armazenamento, "rodrigo", NA_TELA)

    segunda = _executar(pasta, armazenamento, workbook)
    terceira = _executar(pasta, armazenamento, workbook, ler_decisoes=False)

    decisoes = {d.item: d for d in armazenamento.decisoes()}
    assert decisoes[rejeitado.identificador].tipo is TipoDeDecisao.REJEITAR
    assert decisoes[rejeitado.identificador].motivo == "não reconheço"
    assert decisoes[confirmado.identificador].tipo is TipoDeDecisao.CONFIRMAR
    assert decisoes[confirmado.identificador].titulos == confirmado.titulos
    assert decisoes[na_fila.identificador].tipo is TipoDeDecisao.CONCILIADO_FORA
    assert decisoes[na_fila.identificador].referencia == "acordo 12"
    for decisao in decisoes.values():
        assert (decisao.origem, decisao.autor, decisao.momento) == ("Tela", "rodrigo", NA_TELA)
    assert segunda.resumo.decisoes.conflitos == ()
    assert (segunda.resumo.decisoes.confirmadas, segunda.resumo.decisoes.rejeitadas) == (1, 1)
    assert segunda.lista_de_baixa == []
    assert [item.identificador for item in segunda.fila_de_excecoes] == [rejeitado.identificador]
    assert terceira.resumo.decisoes.confirmadas == terceira.resumo.decisoes.rejeitadas == 0
    assert len(armazenamento.decisoes()) == 3


def test_nova_decisao_na_tela_sobre_o_mesmo_item_substitui_a_anterior(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_na_lista_e_um_na_fila(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    item = primeira.lista_de_baixa[0].identificador
    registrar_decisao_da_tela(armazenamento, LinhaDeDecisao(ABA_DA_LISTA_DE_BAIXA, item, "confirmar"), "rodrigo", NA_TELA)
    registrar_decisao_da_tela(
        armazenamento, LinhaDeDecisao(ABA_DA_LISTA_DE_BAIXA, item, "rejeitar", motivo="valor errado"), "rodrigo", NA_TELA
    )

    _executar(pasta, armazenamento, workbook)

    [decisao] = armazenamento.decisoes()
    assert (decisao.item, decisao.tipo) == (item, TipoDeDecisao.REJEITAR)


def test_conflito_nao_volta_na_execucao_seguinte(pasta: Path, armazenamento: Armazenamento, tmp_path: Path) -> None:
    _dois_na_lista_e_um_na_fila(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    disputado = primeira.lista_de_baixa[0]
    _salvar_pela_analista(workbook)
    registrar_decisao_da_tela(
        armazenamento, LinhaDeDecisao(ABA_DA_LISTA_DE_BAIXA, disputado.identificador, "confirmar"), "rodrigo", NA_TELA
    )
    _executar(pasta, armazenamento, workbook)

    terceira = _executar(pasta, armazenamento, workbook, ler_decisoes=False)

    assert terceira.resumo.decisoes.conflitos == ()
    assert [item.identificador for item in terceira.lista_de_baixa] == [disputado.identificador]


def test_possivel_fonte_nao_integrada_no_excel_e_na_tela_e_conflito_e_nada_e_aplicado(
    pasta: Path, armazenamento: Armazenamento, tmp_path: Path
) -> None:
    _dois_na_lista_e_um_na_fila(pasta)
    workbook = tmp_path / "conciliacao.xlsx"
    primeira = _executar(pasta, armazenamento, workbook)
    [na_fila] = primeira.fila_de_excecoes
    _salvar_pela_analista(workbook, {na_fila.identificador: "possível fonte não integrada"})
    registrar_decisao_da_tela(
        armazenamento,
        LinhaDeDecisao(ABA_DA_FILA_DE_EXCECOES, na_fila.identificador, "possível fonte não integrada"),
        "rodrigo",
        NA_TELA,
    )

    segunda = _executar(pasta, armazenamento, workbook)

    [ainda_na_fila] = segunda.fila_de_excecoes
    assert ainda_na_fila.motivo is na_fila.motivo
    assert na_fila.identificador not in {d.item for d in armazenamento.decisoes()}
    assert len(segunda.resumo.decisoes.conflitos) == 1
    assert segunda.resumo.decisoes.fontes_nao_integradas == 0
