from decimal import Decimal
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from conciliacao.calendario import vencimento_efetivo
from conciliacao.dominio import (
    ArquivoRecusado,
    BaixaManual,
    CausaRaiz,
    Destino,
    GrupoDaCausaRaiz,
    Item,
    ItemDeTitulo,
    Metricas,
    ParteAutomatica,
    Resultado,
    Resumo,
    SituacaoNoErp,
    TipoDeDecisao,
    Titulo,
    TituloACancelar,
    TituloNaCobranca,
    TituloNoAcompanhamento,
    formatar_limite,
)
from conciliacao.normalizacao import formatar_percentual, formatar_reais

ABA_DA_LISTA_DE_BAIXA = Destino.LISTA_DE_BAIXA.value
ABA_DA_FILA_DE_EXCECOES = Destino.FILA_DE_EXCECOES.value
LINHA_DA_EXECUCAO_NO_RESUMO = "Execução"
COLUNA_DO_ITEM = "Item"
COLUNA_DA_DECISAO = "Decisão"
COLUNA_DO_MOTIVO_DA_REJEICAO = "Motivo da rejeição"
COLUNA_DOS_TITULOS_ESCOLHIDOS = "Títulos escolhidos"
COLUNA_DO_FRANQUEADO_DO_CONCILIADO_FORA = "Franqueado do Conciliado fora (CNPJ)"
COLUNA_DA_REFERENCIA_DO_CONCILIADO_FORA = "Referência do Conciliado fora"
ABA_DE_REVERSAO = "Reverter confirmação"
COLUNA_DO_MOTIVO_DA_REVERSAO = "Motivo da reversão"
_COLUNAS_DE_DECISAO_DA_FILA = (
    COLUNA_DA_DECISAO,
    COLUNA_DOS_TITULOS_ESCOLHIDOS,
    COLUNA_DO_FRANQUEADO_DO_CONCILIADO_FORA,
    COLUNA_DA_REFERENCIA_DO_CONCILIADO_FORA,
)

_FORMATO_DE_VALOR = "#,##0.00"
_FORMATO_DE_DATA = "DD/MM/YYYY"
_FORMATO_DE_PERCENTUAL = "0.0%"


def gerar_workbook(resultado: Resultado, caminho: Path) -> None:
    workbook = Workbook()
    resumo = workbook.active
    assert resumo is not None
    resumo.title = "Resumo"
    _preencher_resumo(resumo, resultado)
    _aba_de_itens(
        workbook.create_sheet(ABA_DA_LISTA_DE_BAIXA),
        resultado.lista_de_baixa,
        ("Títulos", "Explicação", COLUNA_DA_DECISAO, COLUNA_DO_MOTIVO_DA_REJEICAO),
    )
    _aba_de_baixa_manual(workbook.create_sheet("Baixa manual no ERP"), resultado.baixa_manual)
    _aba_de_cancelar_na_origem(workbook.create_sheet("Cancelar na origem"), resultado.cancelar_na_origem)
    _aba_de_acompanhamento_do_erp(workbook.create_sheet("Acompanhamento do ERP"), resultado.acompanhamento_do_erp)
    reversao = workbook.create_sheet(ABA_DE_REVERSAO)
    reversao.append((COLUNA_DO_ITEM, COLUNA_DO_MOTIVO_DA_REVERSAO))
    fila = workbook.create_sheet(ABA_DA_FILA_DE_EXCECOES)
    _aba_de_itens(
        fila,
        resultado.fila_de_excecoes,
        (
            "Motivo",
            "Títulos sugeridos",
            "Evidência",
            "O que falta",
            "Explicação",
            "Parecer da IA",
            *_COLUNAS_DE_DECISAO_DA_FILA,
        ),
    )
    _acrescentar_itens_de_titulo(fila, resultado.itens_de_titulo)
    _aba_de_cobranca(workbook.create_sheet("Lista de cobrança"), resultado)
    _aba_de_causa_raiz(workbook.create_sheet("Causa raiz"), resultado.causa_raiz)
    _aba_de_metricas(workbook.create_sheet("Métricas"), resultado.metricas)
    _aba_de_itens(workbook.create_sheet(Destino.FORA_DO_ESCOPO.value), resultado.fora_do_escopo, ("Explicação",))
    _aba_de_arquivos_recusados(workbook.create_sheet("Arquivos recusados"), resultado.arquivos_recusados)
    workbook.save(caminho)


def _preencher_resumo(aba: Worksheet, resultado: Resultado) -> None:
    for linha in linhas_do_resumo(resultado):
        aba.append(linha)
    aba.column_dimensions["A"].width = 32
    aba.column_dimensions["B"].width = 60


def linhas_do_resumo(resultado: Resultado) -> list[tuple[str, str | int]]:
    resumo = resultado.resumo
    linhas: list[tuple[str, str | int]] = []
    linhas.append((LINHA_DA_EXECUCAO_NO_RESUMO, resumo.execucao))
    linhas.append(("Arquivos lidos", len(resumo.arquivos_lidos)))
    for lido in resumo.arquivos_lidos:
        linhas.append((f"Lido: {lido}", ""))
    linhas.append(("Arquivos recusados", len(resumo.arquivos_recusados)))
    for recusado in resumo.arquivos_recusados:
        linhas.append((f"Recusado: {recusado.nome}", recusado.motivo))
    if resumo.retornos_sem_credito:
        linhas.append(("Retornos de cobrança sem Crédito", len(resumo.retornos_sem_credito)))
    for sem_credito in resumo.retornos_sem_credito:
        retorno = sem_credito.retorno
        linhas.append((f"Sem Crédito: {retorno.tipo.value} {retorno.identificador}", sem_credito.descrever()))
    linhas.append(("Lançamentos novos no Extrato", resumo.lancamentos_novos))
    for destino, total in resumo.totais.items():
        linhas.append((destino.value, f"{total.quantidade} lançamento(s), {formatar_reais(total.valor_centavos)}"))
    linhas.append(("Resolvido sozinho", _resolvido_sozinho(resumo.parte_automatica)))
    linhas.append(("Alerta de queda da parte automática", _alerta_de_queda(resumo)))
    linhas.append(("Itens de Título na Fila", len(resultado.itens_de_titulo)))
    referencia = resumo.data_de_referencia
    linhas.append(("Data de referência", f"{referencia.data:%d/%m/%Y}" if referencia.data else _SEM_DATA_DE_REFERENCIA))
    linhas.append(
        (
            "Conta que segura a Data de referência",
            "; ".join(
                f"{c.banco} {c.conta} ("
                + (f"último lançamento em {c.ultima_data:%d/%m/%Y}" if c.ultima_data else "sem Extrato")
                + ")"
                for c in referencia.contas_que_seguram
            ),
        )
    )
    linhas.append(("Lista de cobrança", total_da_cobranca(resultado.lista_de_cobranca)))
    linhas.append(("Títulos em análise", total_da_cobranca(resultado.em_analise)))
    linhas.append(("IA (parecer de nome)", _situacao_da_ia(resumo)))
    decisoes = resumo.decisoes
    linhas.append(
        (
            "Decisões lidas do Excel",
            f"{decisoes.confirmadas} confirmada(s), {decisoes.rejeitadas} rejeitada(s), "
            + (f"{decisoes.revertidas} revertida(s), " if decisoes.revertidas else "")
            + f"{len(decisoes.nao_aplicadas)} não aplicada(s)",
        )
    )
    linhas.append(("Conflitos entre Excel e tela", len(decisoes.conflitos)))
    linhas.append(
        (
            "Fila resolvida pelo Excel",
            f"{decisoes.classificadas} Crédito(s) sem título classificado(s), "
            f"{decisoes.fontes_nao_integradas} classificado(s) como Possível fonte não integrada, "
            f"{decisoes.conciliados_fora} Conciliado(s) fora, "
            f"{decisoes.duplicidades} Possível(is) duplicidade(s) decidida(s)",
        )
    )
    linhas.append(
        (
            "Conciliado fora (na causa raiz)",
            f"{len(resultado.conciliados_fora)} Crédito(s), "
            f"{formatar_reais(sum(item.lancamento.valor_centavos for item in resultado.conciliados_fora))}",
        )
    )
    linhas.append(("Títulos a cancelar na origem", len(resultado.cancelar_na_origem)))
    for situacao, rotulo in _ROTULOS_DO_ACOMPANHAMENTO.items():
        linhas.append((rotulo, sum(1 for t in resultado.acompanhamento_do_erp if t.situacao is situacao)))
    for aviso in decisoes.nao_aplicadas:
        linhas.append(("Decisão não aplicada", aviso))
    for conflito in decisoes.conflitos:
        linhas.append(("Conflito", conflito))
    return linhas


_ROTULOS_DO_ACOMPANHAMENTO = {
    SituacaoNoErp.AGUARDANDO_BAIXA_NO_ERP: "Aguardando baixa no ERP",
    SituacaoNoErp.BAIXADO_POR_FORA: "Baixados por fora",
    SituacaoNoErp.DIVERGENCIA_DE_BAIXA: "Divergências de baixa",
    SituacaoNoErp.ESTORNO_NECESSARIO_NO_ERP: "Estorno necessário no ERP",
}

_SEM_DATA_DE_REFERENCIA = "sem Data de referência: há conta configurada sem Extrato; nenhum Título vai para cobrança"


def _resolvido_sozinho(parte: ParteAutomatica) -> str:
    quantidade, valor = parte.percentual_da_quantidade, parte.percentual_do_valor
    if quantidade is None or valor is None:
        return "nenhum Crédito para conciliar nesta execução"
    return (
        f"{parte.resolvidos.quantidade} de {parte.creditos.quantidade} Crédito(s) ({formatar_percentual(quantidade)}); "
        f"{formatar_reais(parte.resolvidos.valor_centavos)} de {formatar_reais(parte.creditos.valor_centavos)} "
        f"({formatar_percentual(valor)})"
    )


def _alerta_de_queda(resumo: Resumo) -> str:
    anterior = resumo.parte_automatica_anterior
    if anterior is None or resumo.limite_de_queda is None:
        return "sem execução anterior para comparar"
    if resumo.alerta_de_queda is not None:
        return resumo.alerta_de_queda
    quantidade, valor = anterior.percentual_da_quantidade, anterior.percentual_do_valor
    if quantidade is None or valor is None:
        return f"sem alerta: na execução {resumo.execucao_anterior} não havia Crédito para conciliar"
    return (
        f"sem alerta: na execução {resumo.execucao_anterior} eram {formatar_percentual(quantidade)} em quantidade "
        f"e {formatar_percentual(valor)} em valor; limite de queda de {formatar_limite(resumo.limite_de_queda)} "
        "pontos percentuais"
    )


def total_da_cobranca(titulos: tuple[TituloNaCobranca, ...]) -> str:
    return f"{len(titulos)} Título(s), {formatar_reais(sum(t.saldo_centavos for t in titulos))}"


def _situacao_da_ia(resumo: Resumo) -> str:
    if resumo.itens_ia_indisponivel:
        return (
            f"indisponível: {resumo.itens_ia_indisponivel} item(ns) da Fila sem parecer, marcados 'IA indisponível'; "
            f"{resumo.itens_com_parecer} item(ns) com parecer"
        )
    if resumo.itens_com_parecer:
        return f"funcionando: {resumo.itens_com_parecer} item(ns) da Fila com parecer"
    return "não foi necessária nesta execução"


def _aba_de_itens(aba: Worksheet, itens: list[Item], colunas_do_destino: tuple[str, ...]) -> None:
    aba.append((COLUNA_DO_ITEM, "Banco", "Conta", "Data", "Histórico", "Documento", "Valor", *colunas_do_destino))
    for item in itens:
        lancamento = item.lancamento
        valores_do_destino = {
            "Títulos": descrever_titulos(item.titulos),
            "Títulos sugeridos": descrever_titulos(item.titulos),
            "Motivo": item.motivo.value if item.motivo else "",
            "Evidência": "; ".join(tipo.value for tipo in item.evidencia),
            "O que falta": item.o_que_falta,
            "Explicação": item.explicacao,
            "Parecer da IA": "; ".join(parecer.descrever() for parecer in item.pareceres),
            COLUNA_DA_DECISAO: TipoDeDecisao.CONFIRMAR.value if item.destino is Destino.LISTA_DE_BAIXA else None,
            COLUNA_DO_MOTIVO_DA_REJEICAO: None,
            COLUNA_DOS_TITULOS_ESCOLHIDOS: None,
            COLUNA_DO_FRANQUEADO_DO_CONCILIADO_FORA: None,
            COLUNA_DA_REFERENCIA_DO_CONCILIADO_FORA: None,
        }
        aba.append(
            (
                item.identificador,
                lancamento.banco,
                lancamento.conta,
                lancamento.data,
                lancamento.historico,
                lancamento.documento,
                Decimal(lancamento.valor_centavos) / 100,
                *(valores_do_destino[coluna] for coluna in colunas_do_destino),
            )
        )
        aba.cell(row=aba.max_row, column=4).number_format = _FORMATO_DE_DATA
        aba.cell(row=aba.max_row, column=7).number_format = _FORMATO_DE_VALOR
    aba.freeze_panes = "A2"


def _aba_de_baixa_manual(aba: Worksheet, baixas: tuple[BaixaManual, ...]) -> None:
    aba.append(
        (
            "Título",
            "Sistema de origem",
            "Franqueado",
            "CNPJ",
            "Vencimento",
            "Valor principal",
            "Encargo",
            "Item",
            "Data do Crédito",
            "Banco",
            "Conta",
            "Valor do Crédito",
        )
    )
    for baixa in baixas:
        titulo, credito = baixa.titulo, baixa.credito
        aba.append(
            (
                titulo.documento,
                titulo.sistema_de_origem,
                titulo.nome_do_franqueado,
                titulo.cnpj,
                titulo.vencimento,
                Decimal(baixa.principal_centavos) / 100,
                Decimal(baixa.encargo_centavos) / 100,
                baixa.item,
                credito.data,
                credito.banco,
                credito.conta,
                Decimal(credito.valor_centavos) / 100,
            )
        )
        for coluna in (5, 9):
            aba.cell(row=aba.max_row, column=coluna).number_format = _FORMATO_DE_DATA
        for coluna in (6, 7, 12):
            aba.cell(row=aba.max_row, column=coluna).number_format = _FORMATO_DE_VALOR
    aba.freeze_panes = "A2"


def _aba_de_cancelar_na_origem(aba: Worksheet, titulos: tuple[TituloACancelar, ...]) -> None:
    aba.append(
        ("Título", "Sistema de origem", "Franqueado", "CNPJ", "Vencimento", "Valor", "Título oficial", "Item")
    )
    for cancelar in titulos:
        titulo = cancelar.titulo
        aba.append(
            (
                titulo.documento,
                titulo.sistema_de_origem,
                titulo.nome_do_franqueado,
                titulo.cnpj,
                titulo.vencimento,
                Decimal(titulo.valor_centavos) / 100,
                descrever_titulos((cancelar.oficial,)),
                cancelar.item,
            )
        )
        aba.cell(row=aba.max_row, column=5).number_format = _FORMATO_DE_DATA
        aba.cell(row=aba.max_row, column=6).number_format = _FORMATO_DE_VALOR
    aba.freeze_panes = "A2"


def _aba_de_acompanhamento_do_erp(aba: Worksheet, titulos: tuple[TituloNoAcompanhamento, ...]) -> None:
    aba.append(
        (
            "Situação",
            "Título",
            "Sistema de origem",
            "Franqueado",
            "CNPJ",
            "Vencimento",
            "Valor",
            "Saldo",
            COLUNA_DO_ITEM,
            "Explicação",
        )
    )
    for acompanhado in titulos:
        titulo = acompanhado.titulo
        aba.append(
            (
                acompanhado.situacao.value,
                titulo.documento,
                titulo.sistema_de_origem,
                titulo.nome_do_franqueado,
                titulo.cnpj,
                titulo.vencimento,
                Decimal(titulo.valor_centavos) / 100,
                None if acompanhado.saldo_centavos is None else Decimal(acompanhado.saldo_centavos) / 100,
                acompanhado.item,
                acompanhado.explicacao,
            )
        )
        aba.cell(row=aba.max_row, column=6).number_format = _FORMATO_DE_DATA
        for coluna in (7, 8):
            aba.cell(row=aba.max_row, column=coluna).number_format = _FORMATO_DE_VALOR
    aba.freeze_panes = "A2"


def _acrescentar_itens_de_titulo(aba: Worksheet, itens: tuple[ItemDeTitulo, ...]) -> None:
    for item in itens:
        sem_lancamento = (None,) * 6
        aba.append(
            (
                item.identificador,
                *sem_lancamento,
                item.motivo.value,
                descrever_titulos(item.titulos),
                None,
                item.motivo.o_que_falta,
                item.explicacao,
            )
        )


_COLUNAS_DA_COBRANCA = (
    "Título",
    "Sistema de origem",
    "Franqueado",
    "CNPJ",
    "Vencimento",
    "Vencimento efetivo",
    "Valor",
    "Saldo",
)


def _aba_de_cobranca(aba: Worksheet, resultado: Resultado) -> None:
    aba.append(_COLUNAS_DA_COBRANCA)
    for na_cobranca in resultado.lista_de_cobranca:
        _linha_da_cobranca(aba, na_cobranca)
    aba.append(())
    aba.append(("Em análise", "Há Crédito candidato na Fila de exceções: não cobrar ainda."))
    aba.append((*_COLUNAS_DA_COBRANCA, "Créditos candidatos"))
    for em_analise in resultado.em_analise:
        _linha_da_cobranca(
            aba,
            em_analise,
            "; ".join(resultado.item_do_lancamento(c).identificador for c in em_analise.creditos_candidatos),
        )
    aba.freeze_panes = "A2"


def _linha_da_cobranca(aba: Worksheet, na_cobranca: TituloNaCobranca, *extras: str) -> None:
    titulo = na_cobranca.titulo
    aba.append(
        (
            titulo.documento,
            titulo.sistema_de_origem,
            titulo.nome_do_franqueado,
            titulo.cnpj,
            titulo.vencimento,
            vencimento_efetivo(titulo.vencimento),
            Decimal(titulo.valor_centavos) / 100,
            Decimal(na_cobranca.saldo_centavos) / 100,
            *extras,
        )
    )
    for coluna in (5, 6):
        aba.cell(row=aba.max_row, column=coluna).number_format = _FORMATO_DE_DATA
    for coluna in (7, 8):
        aba.cell(row=aba.max_row, column=coluna).number_format = _FORMATO_DE_VALOR


def _aba_de_causa_raiz(aba: Worksheet, causa_raiz: CausaRaiz) -> None:
    agregacoes = {
        "Por Motivo de não identificação": causa_raiz.por_motivo,
        "Por Franqueado": causa_raiz.por_franqueado,
        "Por banco": causa_raiz.por_banco,
        "Por canal": causa_raiz.por_canal,
    }
    total = GrupoDaCausaRaiz("Total", causa_raiz.quantidade, causa_raiz.valor_centavos)
    for titulo, grupos in agregacoes.items():
        if aba.max_row > 1:
            aba.append(())
        aba.append((titulo,))
        aba.append(("Grupo", "Quantidade", "% quantidade", "Valor", "% valor"))
        for grupo in (*grupos, total):
            fracao_da_quantidade = causa_raiz.fracao_da_quantidade(grupo)
            fracao_do_valor = causa_raiz.fracao_do_valor(grupo)
            aba.append(
                (
                    grupo.nome,
                    grupo.quantidade,
                    float(fracao_da_quantidade) if fracao_da_quantidade is not None else None,
                    Decimal(grupo.valor_centavos) / 100,
                    float(fracao_do_valor) if fracao_do_valor is not None else None,
                )
            )
            aba.cell(row=aba.max_row, column=3).number_format = _FORMATO_DE_PERCENTUAL
            aba.cell(row=aba.max_row, column=4).number_format = _FORMATO_DE_VALOR
            aba.cell(row=aba.max_row, column=5).number_format = _FORMATO_DE_PERCENTUAL
    aba.column_dimensions["A"].width = 50


def _aba_de_metricas(aba: Worksheet, metricas: Metricas) -> None:
    for titulo, coluna, grupos in (
        ("Taxa de rejeição por tipo de evidência", "Tipo de evidência", metricas.por_evidencia),
        ("Taxa de rejeição por regra", "Regra", metricas.por_regra),
    ):
        aba.append((titulo,))
        aba.append((coluna, "Confirmadas", "Rejeitadas", "Taxa de rejeição", "Revertidas depois de confirmadas"))
        for grupo in grupos:
            aba.append((grupo.nome, grupo.confirmadas, grupo.rejeitadas, _fracao(grupo.taxa_de_rejeicao), grupo.revertidas))
            aba.cell(row=aba.max_row, column=4).number_format = _FORMATO_DE_PERCENTUAL
        aba.append(())
    aba.append(("Acerto da primeira sugestão da IA",))
    aba.append(("Modelo", "Versão do prompt", "Avaliados", "Acertos", "Taxa de acerto", "Aguardando decisão"))
    for acerto in metricas.acerto_da_ia:
        aba.append(
            (
                acerto.modelo,
                acerto.versao_do_prompt,
                acerto.avaliados,
                acerto.acertos,
                _fracao(acerto.taxa_de_acerto),
                acerto.aguardando_decisao,
            )
        )
        aba.cell(row=aba.max_row, column=5).number_format = _FORMATO_DE_PERCENTUAL
    aba.column_dimensions["A"].width = 48


def _fracao(fracao: Fraction | None) -> float | None:
    return float(fracao) if fracao is not None else None


def _aba_de_arquivos_recusados(aba: Worksheet, recusados: tuple[ArquivoRecusado, ...]) -> None:
    aba.append(("Arquivo", "Motivo"))
    for recusado in recusados:
        aba.append((recusado.nome, recusado.motivo))
    aba.column_dimensions["A"].width = 40
    aba.column_dimensions["B"].width = 100
    aba.freeze_panes = "A2"


def descrever_titulos(titulos: tuple[Titulo, ...]) -> str:
    return "; ".join(f"{titulo.documento} ({titulo.sistema_de_origem})" for titulo in titulos)
