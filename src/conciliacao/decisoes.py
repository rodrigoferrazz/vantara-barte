import re
import unicodedata
from collections import Counter
from collections.abc import Callable, Iterable, Mapping
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from conciliacao.adaptadores import nome_do_pagador
from conciliacao.armazenamento import Armazenamento
from conciliacao.dominio import (
    DecisoesLidas,
    Item,
    ItemDeTitulo,
    Lancamento,
    MotivoDeNaoIdentificacao,
    Resposta,
    TipoDeDecisao,
    TipoDeEvidencia,
    Titulo,
)
from conciliacao.encargo import ParametrosDeEncargo, enquadrar
from conciliacao.normalizacao import nome_normalizado
from conciliacao.saidas import (
    ABA_DA_FILA_DE_EXCECOES,
    ABA_DA_LISTA_DE_BAIXA,
    COLUNA_DA_DECISAO,
    COLUNA_DA_REFERENCIA_DO_CONCILIADO_FORA,
    COLUNA_DO_FRANQUEADO_DO_CONCILIADO_FORA,
    COLUNA_DO_ITEM,
    COLUNA_DO_MOTIVO_DA_REJEICAO,
    COLUNA_DOS_TITULOS_ESCOLHIDOS,
    LINHA_DA_EXECUCAO_NO_RESUMO,
)

ORIGEM_EXCEL = "Excel"

type ChaveDoTitulo = tuple[str, str]

_DECISOES_DE_CREDITO_DA_FILA = (
    TipoDeDecisao.CONFIRMAR,
    TipoDeDecisao.ANTECIPACAO,
    TipoDeDecisao.DEVOLUCAO,
    TipoDeDecisao.ERRO,
    TipoDeDecisao.CONCILIADO_FORA,
)
_DECISOES_DE_DUPLICIDADE = (TipoDeDecisao.MESMA_DIVIDA, TipoDeDecisao.DIVIDAS_DISTINTAS)
_TITULO_COM_SISTEMA = re.compile(r"(?P<documento>.+?)\s*\((?P<sistema>[^()]+)\)")
_CNPJ_OU_RAIZ = re.compile(r"[\d./\- ]+")


@dataclass(frozen=True)
class _LinhaDeDecisao:
    aba: str
    item: str
    decisao: str
    motivo: str = ""
    titulos: str = ""
    franqueado: str = ""
    referencia: str = ""


@dataclass(frozen=True)
class _WorkbookLido:
    execucao: object
    autor: str
    momento: datetime | None
    linhas: tuple[_LinhaDeDecisao, ...]


@dataclass(frozen=True)
class DecisoesAplicadas:
    lidas: DecisoesLidas
    registradas: tuple[int, ...]
    duplicidades: tuple[int, ...] = ()


class _NaoAplicada(Exception):
    pass


@dataclass(frozen=True)
class _Decisao:
    item: str
    tipo: TipoDeDecisao
    usa: frozenset[ChaveDoTitulo]
    registrar: Callable[[], int]


class _Interpretacao:
    def __init__(
        self,
        armazenamento: Armazenamento,
        execucao: int,
        encargo: ParametrosDeEncargo,
        autor: str,
        momento: datetime,
    ) -> None:
        self._armazenamento = armazenamento
        self._execucao = execucao
        self._encargo = encargo
        self._autor = autor
        self._momento = momento
        resultado = armazenamento.carregar_resultado(execucao)
        self._lista = {item.identificador: item for item in resultado.lista_de_baixa}
        self._fila = {item.identificador: item for item in resultado.fila_de_excecoes}
        self._duplicidades = {item.identificador: item for item in resultado.itens_de_titulo}
        fechados = armazenamento.titulos_com_baixa_confirmada() | armazenamento.titulos_cancelados_na_origem()
        self._abertos = {
            _chave(titulo): titulo for titulo in armazenamento.titulos_da_execucao(execucao) if _chave(titulo) not in fechados
        }
        self._nomes = {raiz: franqueado.todos_os_nomes for raiz, franqueado in armazenamento.franqueados().items()}

    def interpretar(self, linha: _LinhaDeDecisao) -> _Decisao:
        if linha.aba == ABA_DA_LISTA_DE_BAIXA:
            return self._da_lista(linha)
        if linha.item in self._fila:
            return self._do_credito_da_fila(linha, self._fila[linha.item])
        if linha.item in self._duplicidades:
            return self._da_duplicidade(linha, self._duplicidades[linha.item])
        raise _NaoAplicada(
            f"Item {linha.item}: não está na Fila de exceções da execução {self._execucao}; decisão ignorada."
        )

    def _da_lista(self, linha: _LinhaDeDecisao) -> _Decisao:
        item = self._lista.get(linha.item)
        if item is None:
            raise _NaoAplicada(
                f"Item {linha.item}: não está na Lista de baixa da execução {self._execucao}; decisão ignorada."
            )
        tipo = _tipo_da_lista(linha.decisao)
        if tipo is None:
            raise _NaoAplicada(
                f"Item {linha.item}: decisão '{linha.decisao}' não reconhecida "
                "(use 'confirmar', ou deixe vazio ou 'rejeitar' com motivo); não aplicada."
            )
        if tipo is TipoDeDecisao.REJEITAR:
            if not linha.motivo:
                raise _NaoAplicada(f"Item {linha.item}: rejeição sem motivo; não aplicada.")
            return _Decisao(
                item.identificador, tipo, frozenset(), lambda: self._registrar(item, tipo, item.titulos, linha.motivo)
            )
        return self._confirmacao(item, item.titulos)

    def _do_credito_da_fila(self, linha: _LinhaDeDecisao, item: Item) -> _Decisao:
        tipo = _tipo(linha.decisao, _DECISOES_DE_CREDITO_DA_FILA)
        if tipo is None:
            raise _NaoAplicada(_nao_reconhecida(linha, _DECISOES_DE_CREDITO_DA_FILA))
        if tipo is TipoDeDecisao.CONFIRMAR:
            titulos = self._escolhidos(linha, self._abertos) if linha.titulos else item.titulos
            if not titulos:
                raise _NaoAplicada(f"Item {linha.item}: confirmação sem Título escolhido nem sugerido; não aplicada.")
            fechados = [titulo for titulo in titulos if _chave(titulo) not in self._abertos]
            if fechados:
                raise _NaoAplicada(
                    f"Item {linha.item}: Título(s) {_descrever(fechados)} não está(ão) mais em aberto; não aplicada."
                )
            return self._confirmacao(item, titulos)
        if tipo is TipoDeDecisao.CONCILIADO_FORA:
            return self._conciliado_fora(linha, item)
        if item.motivo is not MotivoDeNaoIdentificacao.CREDITO_SEM_TITULO_EM_ABERTO:
            raise _NaoAplicada(
                f"Item {linha.item}: '{tipo.value}' só classifica Crédito sem título em aberto, e o Motivo do item é "
                f"'{item.motivo}'; não aplicada."
            )
        return _Decisao(item.identificador, tipo, frozenset(), lambda: self._registrar(item, tipo, ()))

    def _conciliado_fora(self, linha: _LinhaDeDecisao, item: Item) -> _Decisao:
        digitos = re.sub(r"\D", "", linha.franqueado)
        if not _CNPJ_OU_RAIZ.fullmatch(linha.franqueado) or len(digitos) not in (8, 14):
            raise _NaoAplicada(
                f"Item {linha.item}: Conciliado fora exige o CNPJ ou a raiz do CNPJ do Franqueado "
                f"(recebido '{linha.franqueado}'); não aplicada."
            )
        if not linha.referencia:
            raise _NaoAplicada(f"Item {linha.item}: Conciliado fora exige a referência; não aplicada.")

        def registrar() -> int:
            decisao = self._registrar(item, TipoDeDecisao.CONCILIADO_FORA, ())
            self._armazenamento.registrar_conciliado_fora(decisao, digitos[:8], linha.referencia)
            return decisao

        return _Decisao(item.identificador, TipoDeDecisao.CONCILIADO_FORA, frozenset(), registrar)

    def _da_duplicidade(self, linha: _LinhaDeDecisao, item: ItemDeTitulo) -> _Decisao:
        tipo = _tipo(linha.decisao, _DECISOES_DE_DUPLICIDADE)
        if tipo is None:
            raise _NaoAplicada(_nao_reconhecida(linha, _DECISOES_DE_DUPLICIDADE))
        oficial: Titulo | None = None
        if tipo is TipoDeDecisao.MESMA_DIVIDA:
            escolhidos = self._escolhidos(linha, {_chave(titulo): titulo for titulo in item.titulos})
            if len(escolhidos) != 1:
                raise _NaoAplicada(
                    f"Item {linha.item}: 'mesma dívida' exige escolher um só Título oficial do grupo; não aplicada."
                )
            [oficial] = escolhidos
        outros = frozenset(_chave(titulo) for titulo in item.titulos if titulo != oficial) if oficial else frozenset()

        def registrar() -> int:
            return self._armazenamento.registrar_decisao_de_duplicidade(
                execucao=self._execucao,
                item=item.identificador,
                tipo=tipo,
                oficial=_chave(oficial) if oficial else None,
                autor=self._autor,
                momento=self._momento,
                origem=ORIGEM_EXCEL,
            )

        return _Decisao(item.identificador, tipo, outros, registrar)

    def _confirmacao(self, item: Item, titulos: tuple[Titulo, ...]) -> _Decisao:
        def registrar() -> int:
            decisao = self._registrar(item, TipoDeDecisao.CONFIRMAR, titulos)
            self._vincular_pagador(decisao, item, titulos)
            return decisao

        return _Decisao(item.identificador, TipoDeDecisao.CONFIRMAR, frozenset(map(_chave, titulos)), registrar)

    def _registrar(self, item: Item, tipo: TipoDeDecisao, titulos: tuple[Titulo, ...], motivo: str | None = None) -> int:
        return self._armazenamento.registrar_decisao(
            execucao=self._execucao,
            item=item.identificador,
            tipo=tipo,
            motivo=motivo,
            autor=self._autor,
            momento=self._momento,
            origem=ORIGEM_EXCEL,
            encargos=_encargos(titulos, item.lancamento, self._encargo),
        )

    def _vincular_pagador(self, decisao: int, item: Item, titulos: tuple[Titulo, ...]) -> None:
        nome = nome_do_pagador(item.lancamento)
        normalizado = nome_normalizado(nome) if nome else ""
        if not normalizado:
            return
        franqueados = sorted({titulo.franqueado for titulo in titulos})
        for franqueado in franqueados:
            por_nome = len(franqueados) == 1 and self._identificou_por_nome(item, franqueado, normalizado)
            self._armazenamento.registrar_vinculo_de_pagador(decisao, normalizado, franqueado, por_nome)

    def _identificou_por_nome(self, item: Item, franqueado: str, nome: str) -> bool:
        evidencia = set(item.evidencia)
        return (
            (
                TipoDeEvidencia.NOME_POR_PREFIXO in evidencia
                and any(nome_normalizado(n).startswith(nome) for n in self._nomes.get(franqueado, ()))
            )
            or (
                TipoDeEvidencia.PARECER_DE_LLM in evidencia
                and any(p.franqueado == franqueado and p.parecer.resposta is Resposta.SIM for p in item.pareceres)
            )
            or (TipoDeEvidencia.PAGADOR_CONHECIDO in evidencia and item.franqueado == franqueado)
        )

    def _escolhidos(self, linha: _LinhaDeDecisao, entre: Mapping[ChaveDoTitulo, Titulo]) -> tuple[Titulo, ...]:
        escolhidos: list[Titulo] = []
        for texto in (parte.strip() for parte in linha.titulos.split(";")):
            if not texto:
                continue
            com_sistema = _TITULO_COM_SISTEMA.fullmatch(texto)
            if com_sistema:
                chave = (com_sistema["sistema"].strip(), com_sistema["documento"].strip())
                candidatos = [entre[chave]] if chave in entre else []
            else:
                candidatos = [titulo for (_, documento), titulo in entre.items() if documento == texto]
            if len(candidatos) != 1:
                situacao = "identifica mais de um Título" if candidatos else "não é Título em aberto que possa ser escolhido"
                raise _NaoAplicada(
                    f"Item {linha.item}: '{texto}' {situacao} na execução {self._execucao}; não aplicada."
                )
            escolhidos.extend(candidatos)
        return tuple(dict.fromkeys(escolhidos))


def aplicar_decisoes_do_excel(
    caminho: Path, armazenamento: Armazenamento, encargo: ParametrosDeEncargo, agora: datetime
) -> DecisoesAplicadas:
    try:
        lido = _ler(caminho)
    except (OSError, KeyError, ValueError, BadZipFile, InvalidFileException) as erro:
        return _nada_aplicado(f"O workbook {caminho.name} não pôde ser lido ({erro}); nenhuma decisão aplicada.")
    ultima = armazenamento.ultima_execucao()
    if not lido.autor:
        return _nada_aplicado(
            f"O workbook {caminho.name} não foi salvo pela Analista (sem autor da última alteração); "
            "nenhuma decisão aplicada."
        )
    if ultima is None or lido.execucao != ultima:
        return _nada_aplicado(
            f"O workbook {caminho.name} é da execução {lido.execucao}, mas a última execução é a execução {ultima}; "
            "nenhuma decisão aplicada."
        )
    interpretacao = _Interpretacao(armazenamento, ultima, encargo, lido.autor, lido.momento or agora)
    repetidos = {
        chave for chave, vezes in Counter((linha.aba, linha.item) for linha in lido.linhas).items() if vezes > 1
    }
    avisos: list[str] = []
    decisoes: list[_Decisao] = []
    for linha in lido.linhas:
        if (linha.aba, linha.item) in repetidos:
            avisos.append(f"Item {linha.item}: aparece mais de uma vez na aba {linha.aba}; decisão não aplicada.")
            continue
        try:
            decisoes.append(interpretacao.interpretar(linha))
        except _NaoAplicada as nao_aplicada:
            avisos.append(str(nao_aplicada))
    usos = Counter(chave for decisao in decisoes for chave in decisao.usa)
    registradas: list[int] = []
    duplicidades: list[int] = []
    tipos: list[TipoDeDecisao] = []
    for decisao in decisoes:
        disputados = sorted(chave for chave in decisao.usa if usos[chave] > 1)
        if disputados:
            avisos.append(
                f"Item {decisao.item}: Título(s) {_descrever_chaves(disputados)} usado(s) por mais de uma decisão "
                "do workbook; nenhuma delas aplicada."
            )
            continue
        (duplicidades if decisao.tipo in _DECISOES_DE_DUPLICIDADE else registradas).append(decisao.registrar())
        tipos.append(decisao.tipo)
    return DecisoesAplicadas(
        DecisoesLidas(
            confirmadas=tipos.count(TipoDeDecisao.CONFIRMAR),
            rejeitadas=tipos.count(TipoDeDecisao.REJEITAR),
            nao_aplicadas=tuple(avisos),
            classificadas=sum(1 for tipo in tipos if tipo.classifica_credito_sem_titulo),
            conciliados_fora=tipos.count(TipoDeDecisao.CONCILIADO_FORA),
            duplicidades=len(duplicidades),
        ),
        tuple(registradas),
        tuple(duplicidades),
    )


def _nada_aplicado(aviso: str) -> DecisoesAplicadas:
    return DecisoesAplicadas(DecisoesLidas(nao_aplicadas=(aviso,)), ())


def _tipo_da_lista(decisao: str) -> TipoDeDecisao | None:
    normalizada = decisao.strip().lower()
    if normalizada == TipoDeDecisao.CONFIRMAR.value:
        return TipoDeDecisao.CONFIRMAR
    if normalizada in ("", TipoDeDecisao.REJEITAR.value):
        return TipoDeDecisao.REJEITAR
    return None


def _tipo(decisao: str, aceitos: Iterable[TipoDeDecisao]) -> TipoDeDecisao | None:
    return next((tipo for tipo in aceitos if _sem_acento(tipo.value) == _sem_acento(decisao)), None)


def _sem_acento(texto: str) -> str:
    decomposto = unicodedata.normalize("NFKD", texto)
    return " ".join("".join(c for c in decomposto if not unicodedata.combining(c)).lower().split())


def _nao_reconhecida(linha: _LinhaDeDecisao, aceitos: Iterable[TipoDeDecisao]) -> str:
    opcoes = ", ".join(f"'{tipo.value}'" for tipo in aceitos)
    return f"Item {linha.item}: decisão '{linha.decisao}' não reconhecida (use {opcoes}); não aplicada."


def _chave(titulo: Titulo) -> ChaveDoTitulo:
    return (titulo.sistema_de_origem, titulo.documento)


def _descrever(titulos: Iterable[Titulo]) -> str:
    return _descrever_chaves(map(_chave, titulos))


def _descrever_chaves(chaves: Iterable[ChaveDoTitulo]) -> str:
    return ", ".join(f"{documento} ({sistema})" for sistema, documento in chaves)


def _encargos(titulos: tuple[Titulo, ...], credito: Lancamento, encargo: ParametrosDeEncargo) -> dict[ChaveDoTitulo, int]:
    return {
        _chave(titulo): enquadrar(
            titulo,
            credito.valor_centavos if len(titulos) == 1 else titulo.valor_centavos,
            credito.data,
            encargo,
        ).encargo_pago_centavos
        for titulo in titulos
    }


def _ler(caminho: Path) -> _WorkbookLido:
    livro = load_workbook(caminho, read_only=True)
    try:
        execucao = next(
            (
                linha[1]
                for linha in livro["Resumo"].iter_rows(values_only=True)
                if linha and linha[0] == LINHA_DA_EXECUCAO_NO_RESUMO
            ),
            None,
        )
        lista = _linhas_da_aba(
            livro[ABA_DA_LISTA_DE_BAIXA].iter_rows(values_only=True),
            ABA_DA_LISTA_DE_BAIXA,
            {"motivo": COLUNA_DO_MOTIVO_DA_REJEICAO},
        )
        fila = _linhas_da_aba(
            livro[ABA_DA_FILA_DE_EXCECOES].iter_rows(values_only=True),
            ABA_DA_FILA_DE_EXCECOES,
            {
                "titulos": COLUNA_DOS_TITULOS_ESCOLHIDOS,
                "franqueado": COLUNA_DO_FRANQUEADO_DO_CONCILIADO_FORA,
                "referencia": COLUNA_DA_REFERENCIA_DO_CONCILIADO_FORA,
            },
        )
        return _WorkbookLido(
            execucao=execucao,
            autor=(livro.properties.lastModifiedBy or "").strip(),
            momento=livro.properties.modified,
            linhas=(*lista, *(linha for linha in fila if linha.decisao)),
        )
    finally:
        livro.close()


def _linhas_da_aba(
    linhas: Iterable[tuple[object, ...]], aba: str, campos: Mapping[str, str]
) -> tuple[_LinhaDeDecisao, ...]:
    cabecalho, *dados = linhas
    colunas = [str(c) for c in cabecalho]
    item, decisao = colunas.index(COLUNA_DO_ITEM), colunas.index(COLUNA_DA_DECISAO)
    indices = {campo: colunas.index(coluna) for campo, coluna in campos.items()}
    return tuple(
        _LinhaDeDecisao(
            aba,
            _celula(linha, item),
            _celula(linha, decisao),
            **{campo: _celula(linha, indice) for campo, indice in indices.items()},
        )
        for linha in dados
        if _celula(linha, item)
    )


def _celula(linha: tuple[object, ...], coluna: int) -> str:
    valor = linha[coluna] if coluna < len(linha) else None
    return "" if valor is None else str(valor).strip()
