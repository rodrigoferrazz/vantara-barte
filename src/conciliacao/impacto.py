import argparse
import sys
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from conciliacao.armazenamento import Armazenamento
from conciliacao.conferencia import ConferenciaDeSomaFalhou
from conciliacao.configuracao import Configuracao, ConfiguracaoInvalida, carregar_configuracao
from conciliacao.dominio import MotivoDeNaoIdentificacao, RankingDeMigracao, Resultado
from conciliacao.execucao import executar_conciliacao
from conciliacao.parecerista import PareceristaDeNome, parecerista_da_configuracao

ABA_DO_RANKING = "Ranking de migração"
ABA_DE_PREMISSAS = "Premissas"

_FORMATO_DE_VALOR = "#,##0.00"


@dataclass(frozen=True)
class AnaliseDeImpacto:
    configuracao: Configuracao
    antes: Resultado
    ranking: RankingDeMigracao


def analisar_impacto(pasta: Path, configuracao: Configuracao, parecerista: PareceristaDeNome) -> AnaliseDeImpacto:
    with TemporaryDirectory(prefix="analise_de_impacto_") as temporaria:
        armazenamento = Armazenamento(Path(temporaria) / "antes.sqlite")
        try:
            antes = executar_conciliacao(pasta, armazenamento, configuracao, parecerista)
            nomes = {raiz: franqueado.todos_os_nomes for raiz, franqueado in armazenamento.franqueados().items()}
        finally:
            armazenamento.fechar()
    ranking = RankingDeMigracao.da_fila(antes.itens, nomes, configuracao.analise_de_impacto.tamanho_da_onda)
    return AnaliseDeImpacto(configuracao, antes, ranking)


def gerar_workbook_do_impacto(analise: AnaliseDeImpacto, caminho: Path) -> None:
    workbook = Workbook()
    ranking = workbook.active
    assert ranking is not None
    ranking.title = ABA_DO_RANKING
    _aba_do_ranking(ranking, analise.ranking)
    _aba_de_premissas(workbook.create_sheet(ABA_DE_PREMISSAS), analise.configuracao)
    workbook.save(caminho)


def _aba_do_ranking(aba: Worksheet, ranking: RankingDeMigracao) -> None:
    aba.append(("Onda", "Posição", "Franqueado", "Raiz do CNPJ", "Créditos na Fila", "Valor"))
    for franqueado in ranking.franqueados:
        aba.append(
            (
                franqueado.onda,
                franqueado.posicao,
                franqueado.nome,
                franqueado.raiz,
                franqueado.quantidade,
                Decimal(franqueado.valor_centavos) / 100,
            )
        )
        aba.cell(row=aba.max_row, column=6).number_format = _FORMATO_DE_VALOR
    aba.append(())
    sem_franqueado = ranking.sem_franqueado
    aba.append(
        (None, None, sem_franqueado.nome, None, sem_franqueado.quantidade, Decimal(sem_franqueado.valor_centavos) / 100)
    )
    aba.cell(row=aba.max_row, column=6).number_format = _FORMATO_DE_VALOR
    aba.append((None, None, "Créditos sem Franqueado não entram em nenhuma Onda de migração."))
    aba.column_dimensions["C"].width = 60
    aba.freeze_panes = "A2"


def _aba_de_premissas(aba: Worksheet, configuracao: Configuracao) -> None:
    encargo = configuracao.encargo
    analise = configuracao.analise_de_impacto
    aba.append(("Premissa", "Valor"))
    for premissa in (
        ("Semente", analise.semente),
        ("Multa por atraso (%)", encargo.multa_percentual),
        ("Juros ao mês (%)", encargo.juros_ao_mes_percentual),
        ("Tolerância do Encargo (R$)", Decimal(encargo.tolerancia_centavos) / 100),
        ("Tamanho da Onda de migração (Franqueados)", analise.tamanho_da_onda),
        (
            "Motivos do Ranking de migração",
            "; ".join(motivo.value for motivo in MotivoDeNaoIdentificacao if motivo in RankingDeMigracao.MOTIVOS),
        ),
        (
            "Franqueado do Crédito no ranking",
            "o identificado pela conciliação; sem ele, o dos Títulos sugeridos na Fila quando são de um só Franqueado",
        ),
        (
            "Conciliação do antes",
            "a mesma da execução diária, com armazenamento temporário (sem decisões da Analista) e o Parecerista "
            "de nome da configuração",
        ),
    ):
        aba.append(premissa)
    aba.column_dimensions["A"].width = 44
    aba.column_dimensions["B"].width = 100


def main(argumentos: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="analisar-impacto",
        description="Analisa o impacto da Cobrança identificada sobre uma pasta base e gera o Ranking de migração.",
    )
    parser.add_argument("pasta", type=Path, help="pasta base com as exportações e os Extratos, do jeito que vêm")
    parser.add_argument("--configuracao", type=Path, default=Path("conciliacao.toml"))
    parser.add_argument("--saida", type=Path, default=Path("analise_de_impacto.xlsx"))
    opcoes = parser.parse_args(argumentos)

    if not opcoes.pasta.is_dir():
        print(f"Pasta base não encontrada: {opcoes.pasta}", file=sys.stderr)
        return 2
    try:
        configuracao = carregar_configuracao(opcoes.configuracao)
    except FileNotFoundError:
        print(f"Configuração não encontrada: {opcoes.configuracao}", file=sys.stderr)
        return 2
    except ConfiguracaoInvalida as erro:
        print(f"Configuração inválida: {erro}", file=sys.stderr)
        return 2

    try:
        analise = analisar_impacto(opcoes.pasta, configuracao, parecerista_da_configuracao(configuracao.ia))
    except ConferenciaDeSomaFalhou as erro:
        print(str(erro), file=sys.stderr)
        return 1
    gerar_workbook_do_impacto(analise, opcoes.saida)
    print(f"Análise de impacto gerada em {opcoes.saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
